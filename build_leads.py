"""Build 'BBG Leads & Trials.xlsx' (Google Sheets compatible)."""
import sys, datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule
from openpyxl.utils import get_column_letter as L
OUT = sys.argv[1]
NAVY, GOLD, GOLD_L = "081C2B", "EAA959", "FBEBD3"
F = "Arial"
f_title = Font(name=F, size=16, bold=True, color=GOLD); f_sub = Font(name=F, size=10, color="FFFFFF")
f_head = Font(name=F, size=10, bold=True, color="FFFFFF"); f_lab = Font(name=F, size=10); f_bold = Font(name=F, size=10, bold=True)
f_in = Font(name=F, size=10, color="0000FF"); f_calc = Font(name=F, size=10); f_note = Font(name=F, size=9, italic=True, color="555555")
f_sec = Font(name=F, size=11, bold=True, color=NAVY)
fill_navy = PatternFill("solid", fgColor=NAVY); fill_in = PatternFill("solid", fgColor="FFFF00"); fill_goldl = PatternFill("solid", fgColor=GOLD_L)
thin = Side(style="thin", color="BFBFBF"); box = Border(top=thin, bottom=thin, left=thin, right=thin)
DATE = "mmm d, yyyy"; wrap = Alignment(wrap_text=True, vertical="top")

wb = Workbook()
def banner(ws, title, sub, width):
    for r in (1, 2):
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=width)
        for c in range(1, width + 1): ws.cell(r, c).fill = fill_navy
    ws["A1"] = title; ws["A1"].font = f_title; ws["A2"] = sub; ws["A2"].font = f_sub
    ws["A2"].alignment = Alignment(wrap_text=True, vertical="center"); ws.row_dimensions[1].height = 30; ws.row_dimensions[2].height = 30

# ---- Leads
ws = wb.active; ws.title = "Leads"
COLS = [("Date asked", 12), ("First name", 14), ("Who it's for", 13), ("Interested in", 18), ("How they reached out", 15),
        ("Replied?", 9), ("Free trial date", 13), ("Came to trial?", 10), ("Joined?", 9), ("Plan", 16),
        ("Next follow-up", 13), ("Follow-up due?", 10), ("Handled by", 11), ("Notes (no health or contact details)", 40)]
banner(ws, "BBG LEADS & FREE TRIALS",
       "One row per person who asks about classes. First names only: phone numbers and anything about health stay in Gymdesk.", len(COLS))
HR = 4
for i, (h, w) in enumerate(COLS, 1):
    c = ws.cell(HR, i, h); c.font = f_head; c.fill = fill_navy; c.border = box; c.alignment = Alignment(wrap_text=True, vertical="center")
    ws.column_dimensions[L(i)].width = w
ws.row_dimensions[HR].height = 32
MAXR = HR + 500
for r in range(HR + 1, MAXR + 1):
    for i in range(1, len(COLS) + 1):
        c = ws.cell(r, i); c.border = box
        if i != 12: c.font = f_in; c.fill = fill_in
    for i in (1, 7, 11): ws.cell(r, i).number_format = DATE
    c = ws.cell(r, 12, f'=IF(OR(B{r}="",I{r}="Yes",I{r}="No - not now",K{r}=""),"",IF(K{r}<=TODAY(),"DUE",""))')
    c.font = Font(name=F, size=10, bold=True, color="C00000")
def dv(items, rng):
    d = DataValidation(type="list", formula1='"' + ",".join(items) + '"', allow_blank=True); ws.add_data_validation(d); d.add(rng)
dv(["Adult", "Teen (13-17)", "Kid (6-12)"], f"C{HR+1}:C{MAXR}")
dv(["Boxing", "Kids Boxing", "Boxing 11-14", "Women's Boxing", "Kickboxing", "Wrestling", "Jiu-Jitsu", "MMA", "Competition", "Yoga",
    "Personal training", "Birthday party", "Fitness camp", "Not sure yet"], f"D{HR+1}:D{MAXR}")
dv(["Facebook", "Instagram", "Cody's page", "Phone", "Text", "Email", "Walk-in", "Website", "Friend"], f"E{HR+1}:E{MAXR}")
dv(["Yes", "No"], f"F{HR+1}:F{MAXR}"); dv(["Yes", "No", "Rebooked"], f"H{HR+1}:H{MAXR}")
dv(["Yes", "Not yet", "No - not now"], f"I{HR+1}:I{MAXR}")
dv(["Drop-in $20", "Week $25", "Kids Program $50/4wk", "Womens Boxing $50/4wk", "Boxing Only $75/mo", "Grappling Only $75/mo", "Single Sport $75/mo", "All Access $125/mo", "All Access 1 year $1000", "Personal training", "Birthday party", "Fitness camp"], f"J{HR+1}:J{MAXR}")
dv(["Cody", "Coach", "Front desk"], f"M{HR+1}:M{MAXR}")
dd = DataValidation(type="date", allow_blank=True); ws.add_data_validation(dd); dd.add(f"A{HR+1}:A{MAXR}"); dd.add(f"G{HR+1}:G{MAXR}"); dd.add(f"K{HR+1}:K{MAXR}")
ws.conditional_formatting.add(f"I{HR+1}:I{MAXR}", CellIsRule(operator="equal", formula=['"Yes"'], fill=PatternFill("solid", fgColor="C6EFCE")))
# example row
ex = [datetime.date(2026, 10, 14), "(Example) Sam", "Adult", "Boxing", "Facebook", "Yes", datetime.date(2026, 10, 16), "Yes", "Not yet",
      None, datetime.date(2026, 10, 18), None, "Cody", "Example row: wants fitness, liked the Wednesday class. Replace with a real lead."]
for i, v in enumerate(ex, 1):
    if v is not None: ws.cell(HR + 1, i, v)
ws.freeze_panes = "C5"; ws.auto_filter.ref = f"A{HR}:{L(len(COLS))}{MAXR}"
R = lambda col: f"Leads!${col}${HR+1}:${col}${MAXR}"

# ---- Summary
wsS = wb.create_sheet("Summary", 0)
banner(wsS, "GYM LEADS SUMMARY", "Updates from the Leads tab. Pick the month in yellow.", 5)
for c, w in zip("ABCDE", (38, 16, 4, 30, 16)): wsS.column_dimensions[c].width = w
wsS["A4"] = "Month to report on (any date in that month)"; wsS["A4"].font = f_lab
c = wsS["B4"]; c.value = datetime.date(2026, 10, 1); c.font = f_in; c.fill = fill_in; c.number_format = "mmmm yyyy"; c.border = box
wsS["A5"] = "Month starts"; wsS["B5"] = "=DATE(YEAR(B4),MONTH(B4),1)"; wsS["B5"].number_format = DATE
wsS["A6"] = "Month ends"; wsS["B6"] = "=EOMONTH(B4,0)"; wsS["B6"].number_format = DATE
for a in ("A5", "A6"): wsS[a].font = f_note
r = 8
wsS.cell(r, 1, "THIS MONTH").font = f_sec; r += 1
rows = [
    ("New inquiries", f'=COUNTIFS({R("A")},">="&B5,{R("A")},"<="&B6,{R("B")},"<>")', "0"),
    ("Replied to", f'=COUNTIFS({R("A")},">="&B5,{R("A")},"<="&B6,{R("F")},"Yes")', "0"),
    ("Free trials booked", f'=COUNTIFS({R("A")},">="&B5,{R("A")},"<="&B6,{R("G")},"<>")', "0"),
    ("Came to their trial", f'=COUNTIFS({R("A")},">="&B5,{R("A")},"<="&B6,{R("H")},"Yes")', "0"),
    ("Joined", f'=COUNTIFS({R("A")},">="&B5,{R("A")},"<="&B6,{R("I")},"Yes")', "0"),
]
rr = {}
for lab, fml, fmt in rows:
    wsS.cell(r, 1, lab).font = f_lab; wsS.cell(r, 1).border = box
    c = wsS.cell(r, 2, fml); c.font = Font(name=F, size=11, bold=True); c.number_format = fmt; c.border = box
    rr[lab] = r; r += 1
for lab, num, den in (("Inquiry → trial", "Free trials booked", "New inquiries"), ("Trial → member", "Joined", "Came to their trial"),
                      ("Inquiry → member", "Joined", "New inquiries")):
    wsS.cell(r, 1, lab).font = f_lab; wsS.cell(r, 1).border = box
    c = wsS.cell(r, 2, f"=IF(B{rr[den]}=0,0,B{rr[num]}/B{rr[den]})"); c.number_format = "0%"; c.border = box; c.font = Font(name=F, size=11, bold=True)
    r += 1
r += 1
wsS.cell(r, 1, "RIGHT NOW").font = f_sec; r += 1
for lab, fml in (("Follow-ups due today or overdue", f'=COUNTIF({R("L")},"DUE")'),
                 ("Inquiries not replied to yet", f'=COUNTIFS({R("B")},"<>",{R("F")},"No")'),
                 ("Trials booked for today or later", f'=COUNTIFS({R("G")},">="&TODAY())')):
    wsS.cell(r, 1, lab).font = f_lab; wsS.cell(r, 1).border = box
    c = wsS.cell(r, 2, fml); c.font = Font(name=F, size=11, bold=True); c.border = box
    if "Trials booked" not in lab:
        wsS.conditional_formatting.add(f"B{r}", CellIsRule(operator="greaterThan", formula=["0"], fill=PatternFill("solid", fgColor="F8CBAD")))
    r += 1
r += 1
n = wsS.cell(r, 1, "Claude's Front Desk skill gives a one-line lead note after every reply it drafts. Paste it in as a new row on Leads. "
                   "Members who've stopped coming are checked from the Gymdesk attendance export, not this sheet.")
n.font = f_note; n.alignment = wrap; wsS.merge_cells(start_row=r, start_column=1, end_row=r, end_column=4); wsS.row_dimensions[r].height = 30
wsS.sheet_view.showGridLines = False
for w in wb.worksheets: w.sheet_properties.tabColor = GOLD if w.title == "Summary" else NAVY
wb.active = 0
wb.save(OUT); print("saved", OUT)
