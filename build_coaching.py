"""Build 'BBG Coaching Library.xlsx': classes, a drill library, 4-week lesson plans that check time, age and contact, and the safety rules.
Google Sheets compatible (INDEX/MATCH, COUNTIF, SEARCH; inline dropdowns without commas).
usage: python3 build_coaching.py OUT [research/coaching.json]"""
import sys, json, os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule
from openpyxl.utils import get_column_letter as L
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from coaching_data import DRILLS, CLASSES, CONTACT, AGES, MIN_AGE, HOUSE, EQUIPMENT, NO_EQUIPMENT

OUT = sys.argv[1]
RES = json.load(open(sys.argv[2])) if len(sys.argv) > 2 and os.path.exists(sys.argv[2]) else None
NAVY, GOLD, GOLD_L = "1C1D2F", "D19A4B", "FBEBD3"   # BBG logo navy and gold
F = "Arial"
f_title = Font(name=F, size=16, bold=True, color=GOLD); f_sub = Font(name=F, size=10, color="FFFFFF")
f_head = Font(name=F, size=10, bold=True, color="FFFFFF"); f_lab = Font(name=F, size=10); f_bold = Font(name=F, size=10, bold=True)
f_in = Font(name=F, size=10, color="0000FF"); f_note = Font(name=F, size=9, italic=True, color="555555")
f_sec = Font(name=F, size=11, bold=True, color=NAVY); f_calc = Font(name=F, size=10, bold=True, color="000000"); f_link = Font(name=F, size=9, color="1F5FA8")
fill_navy = PatternFill("solid", fgColor=NAVY); fill_in = PatternFill("solid", fgColor="FFFF00"); fill_goldl = PatternFill("solid", fgColor=GOLD_L)
thin = Side(style="thin", color="BFBFBF"); box = Border(top=thin, bottom=thin, left=thin, right=thin)
wrap = Alignment(wrap_text=True, vertical="top"); center = Alignment(horizontal="center", vertical="top")

wb = Workbook()
def banner(ws, title, sub, width):
    for r in (1, 2):
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=width)
        for c in range(1, width + 1): ws.cell(r, c).fill = fill_navy
    ws["A1"] = title; ws["A1"].font = f_title; ws["A1"].alignment = Alignment(vertical="center")
    ws["A2"] = sub; ws["A2"].font = f_sub; ws["A2"].alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[1].height = 30; ws.row_dimensions[2].height = 32
def head(ws, r, cols):
    for i, (h, w) in enumerate(cols, 1):
        c = ws.cell(r, i, h); c.font = f_head; c.fill = fill_navy; c.border = box; c.alignment = Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[L(i)].width = w
    ws.row_dimensions[r].height = 30
def inp(c, v=None):
    c.value = v; c.font = f_in; c.fill = fill_in; c.border = box; c.alignment = wrap
def calc(c, v, fmt=None):
    c.value = v; c.font = f_calc; c.border = box; c.alignment = wrap
    if fmt: c.number_format = fmt
def dv_list(ws, items, rng):
    d = DataValidation(type="list", formula1='"' + ",".join(items) + '"', allow_blank=True); ws.add_data_validation(d); d.add(rng)
def dv_range(ws, ref, rng):
    d = DataValidation(type="list", formula1=ref, allow_blank=True); ws.add_data_validation(d); d.add(rng)

# ---------------- Start Here
ws0 = wb.active; ws0.title = "Start Here"
banner(ws0, "BBG COACHING LIBRARY", "Class profiles, a drill library and 4-week lesson plans for every class on the schedule, with the safety rules every coach follows.", 3)
ws0.column_dimensions["A"].width = 4; ws0.column_dimensions["B"].width = 26; ws0.column_dimensions["C"].width = 94
steps = [
 ("1", "Classes", "Every class on the Gymdesk schedule: length, ages, the contact allowed and the focus. Cody sets the contact rules; the lesson plans check against them."),
 ("2", "Drills", f"{len(DRILLS)} drills: how to run each one, coaching points, safety, ages, contact level, equipment and minutes. Add your own at the bottom."),
 ("3", "Lesson Plans", "A 4-week block for each class, six slots per class: warm-up, two skills, a pad or partner drill, a game or live round, and a finisher. Pick drills from the dropdowns; the Check column flags a plan that runs over time, uses a drill for older ages, or uses more contact than the class allows. In an all-ages class it says 'split by age' when some drills are for older athletes, so younger kids do their own version."),
 ("4", "Safety", "The rules every coach follows (Rowan's Law, the Rule of Two, youth contact limits, what to do after a hit to the head), with sources. Read it before coaching kids."),
 ("5", "With Claude", "The bbg-coach skill uses this sheet: \"plan Wednesday's kids boxing\", \"a 4-week wrestling block on takedowns\", \"a sub-coach sheet for Thursday\", \"a parent update for the kids class\"."),
 ("6", "Who decides", "Cody, the head coach at every class, decides what's taught and who spars. These plans are starting points, not rules from a sport body: check the Safety tab and your coaching certification."),
]
r = 4
for n, a, b in steps:
    ws0.cell(r, 1, n).font = Font(name=F, size=12, bold=True, color=GOLD); ws0.cell(r, 2, a).font = f_bold
    ws0.cell(r, 1).alignment = Alignment(vertical="top"); ws0.cell(r, 2).alignment = Alignment(vertical="top", wrap_text=True)
    c = ws0.cell(r, 3, b); c.font = f_lab; c.alignment = wrap; ws0.row_dimensions[r].height = 34; r += 1
r += 1
ws0.cell(r, 2, "Contact levels").font = f_sec; r += 1
for a, b in zip(CONTACT, ["Solo work: shadowboxing, footwork, movement", "Hitting pads, shields or bags held by a coach or partner",
                          "Cooperative partner work with no strikes: drilling a takedown, a sweep, glove-tap games", "Live wrestling or grappling, matched by size, coach watching",
                          "Light touches to the body only, coach watching every pair", "Sparring with head contact: Competition Class only, Cody's OK, full gear"]):
    ws0.cell(r, 2, a).font = f_bold; ws0.cell(r, 3, b).font = f_lab; r += 1
r += 1
ws0.cell(r, 2, "Colours").font = f_sec; r += 1
for lab, fill, txt in (("Yellow, blue text", fill_in, "You type or pick here"), ("Black, bold", None, "Calculated: don't type over it")):
    c = ws0.cell(r, 2, lab); c.font = f_in if fill else f_lab
    if fill: c.fill = fill
    ws0.cell(r, 3, txt).font = f_lab; r += 1

# ---------------- Lists (lookup)
wl = wb.create_sheet("Lists")
wl["A1"] = "Ages"; wl["B1"] = "Youngest"; wl["D1"] = "Contact"
for c in ("A1", "B1", "D1"): wl[c].font = f_bold
for i, a in enumerate(AGES): wl.cell(2 + i, 1, a); wl.cell(2 + i, 2, MIN_AGE[a])
for i, c in enumerate(CONTACT): wl.cell(2 + i, 4, c)
AG_N = 1 + len(AGES)

# ---------------- Classes
wc = wb.create_sheet("Classes")
CCOL = [("Class", 20), ("Day and time", 24), ("Minutes", 9), ("Ages", 22), ("Youngest age", 9), ("Contact allowed (separate with ;)", 40), ("Focus", 60), ("Coach (first name)", 16), ("Splits into age groups?", 12)]
banner(wc, "CLASSES", "From the Gymdesk schedule and Cody's answers (October 2026). Change the contact allowed, youngest age or age groups here and every lesson plan re-checks.", len(CCOL))
head(wc, 4, CCOL)
CMAX = 4 + len(CLASSES) + 10
for k in range(CMAX - 4):
    rr = 5 + k
    row = CLASSES[k] if k < len(CLASSES) else (None,) * 9
    for i, v in enumerate(row, 1): inp(wc.cell(rr, i), v)
    wc.cell(rr, 9).alignment = center
    wc.row_dimensions[rr].height = 44 if k < len(CLASSES) else 18
wc.cell(CMAX + 2, 1, "Yoga (Tue 6:15-7:00 PM) is run by its own instructor and isn't covered here.").font = f_note
wc.cell(CMAX + 3, 1, "Splits into age groups? Yes = an all-ages class. Drills above the youngest age are fine, but younger kids do a drill for their age at the same time (the Check says 'split by age').").font = f_note
wc.cell(CMAX + 4, 1, "Gym equipment (Cody, Oct 2026): " + ", ".join(EQUIPMENT) + ". Not at BBG: " + ", ".join(NO_EQUIPMENT).lower() + ". Plans only use this gear.").font = f_note
dv_list(wc, ["Yes", "No"], f"I5:I{CMAX}")
wc.freeze_panes = "B5"

# ---------------- Drills
wd = wb.create_sheet("Drills")
DCOL = [("Drill", 30), ("Discipline", 12), ("Ages", 14), ("Level", 11), ("Contact", 16), ("Equipment", 22), ("Minutes", 8), ("How to run it", 58), ("Coaching points", 38), ("Safety", 38), ("Youngest age", 9)]
banner(wd, "DRILLS", "The drill library. Add new drills at the bottom: pick Ages and Contact from the dropdowns so the lesson plans can check them.", len(DCOL))
head(wd, 4, DCOL)
DMAX = 4 + len(DRILLS) + 60
for k in range(DMAX - 4):
    rr = 5 + k
    row = DRILLS[k] if k < len(DRILLS) else None
    if row:
        for i, v in enumerate(row, 1):
            inp(wd.cell(rr, i), v)
        wd.cell(rr, 1).font = Font(name=F, size=10, bold=True, color="0000FF")
        wd.row_dimensions[rr].height = 66
    else:
        for i in range(1, 11): inp(wd.cell(rr, i))
    calc(wd.cell(rr, 11), f'=IF(C{rr}="","",IFERROR(INDEX(Lists!$B$2:$B${AG_N},MATCH(C{rr},Lists!$A$2:$A${AG_N},0)),""))'); wd.cell(rr, 11).alignment = center
dv_list(wd, AGES, f"C5:C{DMAX}"); dv_list(wd, CONTACT, f"E5:E{DMAX}"); dv_list(wd, ["Beginner", "Intermediate", "Advanced"], f"D5:D{DMAX}")
dv_list(wd, ["Warm-up", "Boxing", "Kickboxing", "Wrestling", "Jiu-jitsu", "MMA", "Kids games", "Conditioning", "Cool-down"], f"B5:B{DMAX}")
wd.freeze_panes = "B5"; wd.auto_filter.ref = f"A4:K{DMAX}"

# ---------------- Lesson Plans
SLOTS = ["Warm-up", "Skill 1", "Skill 2", "Pads or partner", "Game or live", "Finish / cool-down"]
P = {
 "Kids Boxing": [
  ("Stance and guard", [("Animal walks", 8), ("Stance and guard", 8), ("Step-drag footwork", 10), ("Pad target game", 10), ("Follow the leader footwork", 8), ("Cool-down stretch and breathing", 6)]),
  ("The jab", [("Dynamic warm-up", 10), ("Jab on pads", 10), ("Footwork ladder and cones", 8), ("Heavy bag rounds", 10), ("Relay races", 8), ("Cool-down stretch and breathing", 6)]),
  ("Jab and cross", [("Jump rope", 6), ("1-2 on pads", 10), ("Step-drag footwork", 8), ("Heavy bag rounds", 10), ("Pad target game", 10), ("Cool-down stretch and breathing", 6)]),
  ("Slips and review", [("Animal walks", 8), ("Slip line", 8), ("1-2 on pads", 10), ("Heavy bag rounds", 10), ("Follow the leader footwork", 8), ("Cool-down stretch and breathing", 6)])],
 "Wrestling": [
  ("Stance and safe falling", [("Grappling movement warm-up", 10), ("Breakfalls", 6), ("Stance and motion", 6), ("Sprawl drill", 5), ("Knee tag", 6), ("Cool-down stretch and breathing", 5)]),
  ("Penetration step", [("Breakfalls", 6), ("Penetration step", 8), ("Sprawl drill", 5), ("Sumo push", 6), ("Live wrestling from the knees", 10), ("Cool-down stretch and breathing", 5)]),
  ("Pins", [("Animal walks", 6), ("Half nelson turn", 8), ("Turtle flip", 6), ("Sprawl drill", 5), ("Live wrestling from the knees", 10), ("Cool-down stretch and breathing", 5)]),
  ("Review and games", [("Breakfalls", 6), ("Penetration step", 8), ("Knee tag", 6), ("Sprawl drill", 5), ("Live wrestling from the knees", 10), ("Cool-down stretch and breathing", 5)])],
 "Boxing 11-14": [
  ("Fundamentals", [("Jump rope", 6), ("Stance and guard", 8), ("Step-drag footwork", 8), ("1-2 on pads", 10), ("Heavy bag rounds", 12), ("Cool-down stretch and breathing", 6)]),
  ("Combinations", [("Dynamic warm-up", 10), ("1-2-3 on pads", 10), ("Coach-called combinations", 8), ("Heavy bag rounds", 12), ("Bob and weave under the rope", 6), ("Cool-down stretch and breathing", 6)]),
  ("Defence", [("Jump rope", 6), ("Slip line", 8), ("Bob and weave under the rope", 8), ("Heavy bag rounds", 12), ("Body-only partner touch", 8), ("Cool-down stretch and breathing", 6)]),
  ("Putting it together", [("Dynamic warm-up", 10), ("Footwork ladder and cones", 8), ("1-2-3 on pads", 10), ("Heavy bag rounds", 12), ("Body-only partner touch", 8), ("Cool-down stretch and breathing", 6)])],
 "Boxing": [
  ("Fundamentals", [("Jump rope", 6), ("Stance and guard", 6), ("Step-drag footwork", 8), ("1-2 on pads", 10), ("Heavy bag rounds", 12), ("Body-weight intervals", 8)]),
  ("Hooks and uppercuts", [("Shadowboxing round", 6), ("1-2-3 on pads", 8), ("Uppercuts on pads", 8), ("Heavy bag rounds", 12), ("Bag and body-weight circuit", 12), ("Cool-down stretch and breathing", 5)]),
  ("Defence", [("Jump rope", 6), ("Slip line", 8), ("Pad defense feeds", 10), ("Heavy bag rounds", 12), ("Body-only partner touch", 8), ("Cool-down stretch and breathing", 5)]),
  ("Combinations and conditioning", [("Dynamic warm-up", 10), ("Coach-called combinations", 8), ("Uppercuts on pads", 8), ("Bag and body-weight circuit", 14), ("Body-only partner touch", 8), ("Cool-down stretch and breathing", 5)])],
 "Women's Boxing": [
  ("Fundamentals", [("Jump rope", 5), ("Stance and guard", 6), ("Step-drag footwork", 8), ("1-2 on pads", 8), ("Heavy bag rounds", 10), ("Cool-down stretch and breathing", 5)]),
  ("Combinations", [("Jump rope", 5), ("1-2-3 on pads", 8), ("Slip line", 6), ("Bag and body-weight circuit", 12), ("Coach-called combinations", 6), ("Cool-down stretch and breathing", 5)]),
  ("Defence", [("Shadowboxing round", 6), ("Slip line", 6), ("Pad defense feeds", 8), ("Heavy bag rounds", 10), ("Body-weight intervals", 8), ("Cool-down stretch and breathing", 5)]),
  ("Power and confidence", [("Jump rope", 5), ("Uppercuts on pads", 6), ("Coach-called combinations", 6), ("Bag and body-weight circuit", 12), ("Heavy bag rounds", 10), ("Cool-down stretch and breathing", 5)])],
 "Kickboxing": [
  ("Kicks 101", [("Jump rope", 6), ("Teep on the shield", 8), ("Roundhouse to the body shield", 10), ("1-2 on pads", 8), ("Heavy bag rounds", 12), ("Cool-down stretch and breathing", 5)]),
  ("Combinations", [("Dynamic warm-up", 10), ("Switch kick", 8), ("Jab-cross-kick on pads", 10), ("Heavy bag rounds", 12), ("Battle rope intervals", 8), ("Cool-down stretch and breathing", 5)]),
  ("Defence", [("Shadowboxing round", 6), ("Low-kick checks (slow)", 8), ("Jab-cross-kick on pads", 10), ("Slip line", 6), ("Heavy bag rounds", 12), ("Cool-down stretch and breathing", 5)]),
  ("Fight fit", [("Jump rope", 6), ("Teep on the shield", 8), ("Roundhouse to the body shield", 8), ("Bag and body-weight circuit", 14), ("Body-only partner touch", 8), ("Cool-down stretch and breathing", 5)])],
 "Jiu-Jitsu (no-gi)": [
  ("Mount escapes", [("Grappling movement warm-up", 10), ("Tap practice", 5), ("Mount escape: bridge and roll", 20), ("Mount escape: elbow-knee", 20), ("Positional sparring: guard vs pass", 30), ("Cool-down stretch and breathing", 10)]),
  ("Guard", [("Grappling movement warm-up", 10), ("Shrimp chains", 5), ("Guard retention frames", 20), ("Scissor sweep", 20), ("Positional sparring: guard vs pass", 30), ("Cool-down stretch and breathing", 10)]),
  ("Passing", [("Grappling movement warm-up", 10), ("Technical stand-up", 5), ("Toreando pass", 20), ("Guard retention frames", 15), ("Positional sparring: guard vs pass", 35), ("Cool-down stretch and breathing", 10)]),
  ("Submissions", [("Grappling movement warm-up", 10), ("Tap practice", 5), ("Armbar from mount", 20), ("Rear naked choke from the back", 20), ("Positional sparring: guard vs pass", 30), ("Cool-down stretch and breathing", 10)])],
 "MMA": [
  ("Strike to takedown", [("Dynamic warm-up", 10), ("Sprawl and strike shadow", 6), ("Strike to level change on pads", 10), ("Pummeling for underhooks", 8), ("Clinch positional sparring", 12), ("Cool-down stretch and breathing", 5)]),
  ("Takedowns", [("Grappling movement warm-up", 10), ("Penetration step", 8), ("Double-leg takedown (drill)", 12), ("1-2 on pads", 8), ("Sprawl drill", 6), ("Cool-down stretch and breathing", 5)]),
  ("Against the wall", [("Jump rope", 6), ("Wall walk get-up", 10), ("Pummeling for underhooks", 8), ("Clinch positional sparring", 12), ("Body-weight intervals", 8), ("Cool-down stretch and breathing", 5)]),
  ("Putting it together", [("Dynamic warm-up", 10), ("Strike to level change on pads", 10), ("Jab-cross-kick on pads", 8), ("Live wrestling from the knees", 12), ("Sprawl and strike shadow", 6), ("Cool-down stretch and breathing", 5)])],
 "Competition Class": [
  ("Defence under pressure", [("Shadowboxing round", 6), ("Pad defense feeds", 10), ("Heavy bag rounds", 10), ("Controlled technical sparring", 15), ("Body-weight intervals", 8), ("Cool-down stretch and breathing", 5)]),
  ("Takedown defence", [("Jump rope", 6), ("1-2-3 on pads", 8), ("Strike to level change on pads", 8), ("Clinch positional sparring", 12), ("Controlled technical sparring", 15), ("Cool-down stretch and breathing", 5)]),
  ("Ground game", [("Dynamic warm-up", 10), ("Uppercuts on pads", 6), ("Jab-cross-kick on pads", 8), ("Positional sparring: guard vs pass", 12), ("Controlled technical sparring", 15), ("Cool-down stretch and breathing", 5)]),
  ("Fight week", [("Shadowboxing round", 6), ("Coach-called combinations", 8), ("Heavy bag rounds", 10), ("Controlled technical sparring", 12), ("Bag and body-weight circuit", 12), ("Cool-down stretch and breathing", 5)])],
}
names = {d[0] for d in DRILLS}
for cl, weeks in P.items():
    for th, slots in weeks:
        for dn, _ in slots: assert dn in names, (cl, dn)

wp = wb.create_sheet("Lesson Plans")
PC = [("Class", 18), ("Week", 6), ("Theme", 20)]
for s_ in SLOTS: PC += [(s_, 22), ("Min", 5)]
PC += [("Total", 7), ("Class min", 7), ("Check", 22), ("Notes", 30)]
banner(wp, "LESSON PLANS", "Pick a drill for each slot. Minutes fill in from the drill (type over to change). Check flags plans that run long, use a drill meant for older athletes, or use more contact than the class allows. Amber split by age means an all-ages class runs that drill in age groups.", len(PC))
head(wp, 4, PC)
# hidden helpers after Notes: class youngest age, class allowed contact, 6 slot codes
HC = len(PC) + 1
for i, h in enumerate(["Youngest", "Allowed"] + [f"Code {k+1}" for k in range(6)] + ["Splits"]):
    c = wp.cell(4, HC + i, h); c.font = f_note
rows_out = []
for cl, weeks in P.items():
    for w, (th, slots) in enumerate(weeks, 1): rows_out.append((cl, w, th, slots))
PMAX = 4 + len(rows_out) + 60
DR = f"Drills!$A$5:$A${DMAX}"
for k in range(PMAX - 4):
    rr = 5 + k
    row = rows_out[k] if k < len(rows_out) else None
    inp(wp.cell(rr, 1), row[0] if row else None); inp(wp.cell(rr, 2), row[1] if row else None); inp(wp.cell(rr, 3), row[2] if row else None)
    for j in range(6):
        dc, mc = 4 + 2 * j, 5 + 2 * j
        dname = row[3][j][0] if row else None
        inp(wp.cell(rr, dc), dname)
        mref = f"{L(dc)}{rr}"
        if row: inp(wp.cell(rr, mc), row[3][j][1])
        else:
            inp(wp.cell(rr, mc), f'=IF({mref}="","",IFERROR(INDEX(Drills!$G$5:$G${DMAX},MATCH({mref},{DR},0)),""))')
        wp.cell(rr, mc).alignment = center
    tc = 4 + 12
    calc(wp.cell(rr, tc), f'=IF(A{rr}="","",SUM(E{rr},G{rr},I{rr},K{rr},M{rr},O{rr}))'); wp.cell(rr, tc).alignment = center
    calc(wp.cell(rr, tc + 1), f'=IF(A{rr}="","",IFERROR(INDEX(Classes!$C$5:$C${CMAX},MATCH(A{rr},Classes!$A$5:$A${CMAX},0)),""))'); wp.cell(rr, tc + 1).alignment = center
    # helpers
    yc, ac = L(HC), L(HC + 1)
    wp.cell(rr, HC, f'=IF(A{rr}="","",IFERROR(INDEX(Classes!$E$5:$E${CMAX},MATCH(A{rr},Classes!$A$5:$A${CMAX},0)),""))').font = f_note
    wp.cell(rr, HC + 1, f'=IF(A{rr}="","",IFERROR(INDEX(Classes!$F$5:$F${CMAX},MATCH(A{rr},Classes!$A$5:$A${CMAX},0)),""))').font = f_note
    for j in range(6):
        d = f"{L(4 + 2 * j)}{rr}"
        f = (f'=IF(OR({d}="",$A{rr}=""),"",IF(ISNA(MATCH({d},{DR},0)),"?",'
             f'IF(N(INDEX(Drills!$K$5:$K${DMAX},MATCH({d},{DR},0)))>N(${yc}{rr}),"A","")&'
             f'IF(ISNUMBER(SEARCH(INDEX(Drills!$E$5:$E${DMAX},MATCH({d},{DR},0)),${ac}{rr})),"","C")))')
        wp.cell(rr, HC + 2 + j, f).font = f_note
    codes = f"{L(HC + 2)}{rr}:{L(HC + 7)}{rr}"
    sc = f"{L(HC + 8)}{rr}"
    wp.cell(rr, HC + 8, f'=IF(A{rr}="","",IFERROR(INDEX(Classes!$I$5:$I${CMAX},MATCH(A{rr},Classes!$A$5:$A${CMAX},0)),""))').font = f_note
    issues = (f'IF(AND(N(Q{rr})>0,N(P{rr})>N(Q{rr})),"Over by "&(P{rr}-Q{rr})&" min. ","")'
              f'&IF(AND(COUNTIF({codes},"*A*")>0,{sc}<>"Yes"),"Age check. ","")&IF(COUNTIF({codes},"*C*")>0,"Contact check. ","")'
              f'&IF(COUNTIF({codes},"~?")>0,"Drill not in library. ","")')
    split = f'IF(AND(COUNTIF({codes},"*A*")>0,{sc}="Yes")," · split by age","")'
    chk = (f'=IF(A{rr}="","",IF(TRIM({issues})="","OK"&{split}&IF(N(Q{rr})>N(P{rr})," ("&(Q{rr}-P{rr})&" min spare)",""),'
           f'TRIM({issues})&{split}))')
    calc(wp.cell(rr, tc + 2), chk)
    inp(wp.cell(rr, tc + 3))
    wp.row_dimensions[rr].height = 44 if row else 18
dv_range(wp, f"=Classes!$A$5:$A${CMAX}", f"A5:A{PMAX}")
for j in range(6): dv_range(wp, f"={DR}", f"{L(4 + 2 * j)}5:{L(4 + 2 * j)}{PMAX}")
R_ = f"R5:R{PMAX}"
wp.conditional_formatting.add(R_, FormulaRule(formula=['AND(LEFT(R5,2)="OK",ISNUMBER(SEARCH("split",R5)))'], fill=PatternFill("solid", fgColor="FFEB9C"), font=Font(name=F, color="9C5700", bold=True), stopIfTrue=True))
wp.conditional_formatting.add(R_, FormulaRule(formula=['LEFT(R5,2)="OK"'], fill=PatternFill("solid", fgColor="C6EFCE"), font=Font(name=F, color="006100", bold=True)))
wp.conditional_formatting.add(R_, FormulaRule(formula=['AND(R5<>"",LEFT(R5,2)<>"OK")'], fill=PatternFill("solid", fgColor="FFC7CE"), font=Font(name=F, color="9C0006", bold=True)))
for i in range(HC, HC + 9): wp.column_dimensions[L(i)].hidden = True
wp.freeze_panes = "D5"; wp.auto_filter.ref = f"A4:{L(len(PC))}{PMAX}"

# ---------------- Safety
wsf = wb.create_sheet("Safety")
SC = [("Topic", 22), ("Who it's for", 16), ("The rule or guidance", 80), ("Source", 34)]
banner(wsf, "SAFETY", "What every BBG coach follows. From official sources (October 2026). Where a source isn't clear, Cody decides; this isn't legal or medical advice.", len(SC))
head(wsf, 4, SC)
r = 5
for row in HOUSE:
    for i, v in enumerate(row, 1):
        c = wsf.cell(r, i, v); c.font = f_bold if i == 1 else f_lab; c.border = box; c.alignment = wrap
    wsf.row_dimensions[r].height = 46; r += 1
if RES:
    r += 1
    wsf.cell(r, 1, "From the sport bodies and Ontario").font = f_sec; r += 1
    for it in RES.get("items", []):
        if it.get("confidence") == "low": continue
        vals = (it.get("topic", ""), it.get("age_group", ""), it.get("guidance", ""), (it.get("source_name", "") + " " + it.get("source_url", "")).strip())
        for i, v in enumerate(vals, 1):
            c = wsf.cell(r, i, v); c.font = f_link if i == 4 else (f_bold if i == 1 else f_lab); c.border = box; c.alignment = wrap
        wsf.row_dimensions[r].height = max(30, 14 * (len(vals[2]) // 95 + 1) + 6); r += 1
    if RES.get("open_questions"):
        r += 1; wsf.cell(r, 1, "For Cody to confirm").font = f_sec; r += 1
        for q in RES["open_questions"]:
            c = wsf.cell(r, 3, q); c.font = f_lab; c.alignment = wrap; wsf.cell(r, 1, "Open question").font = f_note
            wsf.row_dimensions[r].height = max(18, 14 * (len(q) // 95 + 1) + 4); r += 1
    if RES.get("answered_oct20"):
        r += 1; wsf.cell(r, 1, "Answered by Cody (Oct 20, 2026)").font = f_sec; r += 1
        for q in RES["answered_oct20"]:
            c = wsf.cell(r, 3, q); c.font = f_lab; c.alignment = wrap; wsf.cell(r, 1, "Answered").font = f_note
            wsf.row_dimensions[r].height = max(18, 14 * (len(q) // 95 + 1) + 4); r += 1
else:
    wsf.cell(r + 1, 1, "Sport-body guidance: to be added.").font = f_note
wsf.freeze_panes = "A5"

for w in wb.worksheets:
    w.page_setup.orientation = "landscape"; w.sheet_properties.pageSetUpPr.fitToPage = True; w.page_setup.fitToWidth = 1; w.page_setup.fitToHeight = 0
wb._sheets = [wb["Start Here"], wb["Classes"], wb["Lesson Plans"], wb["Drills"], wb["Safety"], wb["Lists"]]
wb["Lists"].sheet_state = "hidden"
wb.active = 0
wb.save(OUT); print("saved", OUT, len(DRILLS), "drills", len(rows_out), "plans")
