"""Build 'BBG Skills Passport.xlsx': four levels of skills for each kids' and youth class, a roster, a progress log and a
progress card per athlete for parents. First names only. Google Sheets compatible (SUMPRODUCT, COUNTIFS, INDEX/MATCH).
usage: python3 build_passport.py OUT"""
import sys, os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule, CellIsRule
from openpyxl.utils import get_column_letter as L
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from coaching_data import DRILLS

OUT = sys.argv[1]
NAVY, GOLD, GOLD_L = "1C1D2F", "D19A4B", "FBEBD3"
F = "Arial"
f_title = Font(name=F, size=16, bold=True, color=GOLD); f_sub = Font(name=F, size=10, color="FFFFFF")
f_head = Font(name=F, size=10, bold=True, color="FFFFFF"); f_lab = Font(name=F, size=10); f_bold = Font(name=F, size=10, bold=True)
f_in = Font(name=F, size=10, color="0000FF"); f_note = Font(name=F, size=9, italic=True, color="555555")
f_sec = Font(name=F, size=11, bold=True, color=NAVY); f_calc = Font(name=F, size=10, bold=True, color="000000")
fill_navy = PatternFill("solid", fgColor=NAVY); fill_in = PatternFill("solid", fgColor="FFFF00"); fill_goldl = PatternFill("solid", fgColor=GOLD_L)
thin = Side(style="thin", color="BFBFBF"); box = Border(top=thin, bottom=thin, left=thin, right=thin)
wrap = Alignment(wrap_text=True, vertical="top"); center = Alignment(horizontal="center", vertical="center", wrap_text=True); DATE = "mmm d, yyyy"
LEVELS = ["Rookie", "Contender", "Challenger", "Champion"]

# program -> 4 levels -> [(skill, how it's checked, drill)]
P = {
 "Kids Boxing": [
  [("Stance and guard", "Holds a good stance and guard for 20 seconds while the coach checks", "Stance and guard"),
   ("Step-drag footwork", "Steps forward, back, left and right without crossing the feet", "Step-drag footwork"),
   ("Jab on pads", "10 jabs on the pads, each one back to the guard", "Jab on pads"),
   ("Rope-less skipping", "30 seconds of jump-rope footwork (no rope)", "Jump rope"),
   ("Gym rules", "Stops on the whistle, takes water breaks, puts gloves on and off safely", "")],
  [("1-2 on pads", "Jab-cross on the pads with the back foot pivoting", "1-2 on pads"),
   ("Footwork ladder", "One ladder pattern start to finish without stopping", "Footwork ladder and cones"),
   ("Slips", "Slips left and right along the slip line", "Slip line"),
   ("Bag round", "One 1-minute heavy bag round with the guard up", "Heavy bag rounds"),
   ("Glove tag", "Stays in stance for a whole round of glove tag", "Glove tag")],
  [("Jab-step-jab", "Jab, step in, jab on the pads", "Jab on pads"),
   ("Slip and counter", "Slips, then jabs straight back, on the slip line", "Slip line"),
   ("Bob and weave", "Ducks under the low rope with the knees, guard up", "Bob and weave under the rope"),
   ("Pad target speed", "Hits the right pad on the coach's call, quickly and cleanly", "Pad target game"),
   ("Two bag rounds", "Two 1-minute bag rounds with good stance", "Heavy bag rounds")],
  [("Combination", "1-2, slip, 1-2 on the pads", "1-2 on pads"),
   ("Pivots", "Pivots and changes direction on the coach's call", "Follow the leader footwork"),
   ("Three bag rounds", "Three 1-minute bag rounds", "Heavy bag rounds"),
   ("Helps a newer kid", "Shows a newer kid the stance and guard", "Stance and guard"),
   ("Leads the warm-up", "Leads part of the warm-up for the class", "Animal walks")]],
 "Wrestling": [
  [("Stance and motion", "Moves forward, back and in circles in a good stance", "Stance and motion"),
   ("Back breakfall", "Back breakfall from squatting: chin tucked, slaps the mat", "Breakfalls"),
   ("Side breakfall", "Side breakfall on both sides", "Breakfalls"),
   ("Sprawl", "Sprawls on the whistle and gets back to stance", "Sprawl drill"),
   ("Shoulder roll", "Forward and back shoulder rolls over the shoulder, not the neck", "Grappling movement warm-up")],
  [("Penetration step (knees)", "Penetration step from the knees, back straight, head up", "Penetration step"),
   ("Knee tag", "Protects the knees in stance for a whole round", "Knee tag"),
   ("Sumo push", "Stays low and drives with the legs", "Sumo push"),
   ("Technical stand-up", "Stands up in base with a hand protecting the face", "Technical stand-up"),
   ("The tap", "Taps clearly and lets go at once when a partner taps", "Tap practice")],
  [("Penetration step (standing)", "Level change and penetration step from standing", "Penetration step"),
   ("Half nelson turn", "Turns a partner with pressure on the shoulder, never the neck", "Half nelson turn"),
   ("Turtle defence", "Keeps elbows tight so a partner can't flip them", "Turtle flip"),
   ("Hip escape", "Shrimps the length of the mat both ways", "Shrimp chains"),
   ("Live go", "Uses a move from class in live wrestling from the knees", "Live wrestling from the knees")],
  [("Sprawl and go behind", "Sprawls on a shot and comes around behind", "Sprawl drill"),
   ("Pin with control", "Holds a pin for 3 seconds with control", "Half nelson turn"),
   ("Sportsmanship", "Shakes hands, follows the coach's calls, wins and loses well", ""),
   ("Teaches breakfalls", "Shows a newer kid the back breakfall", "Breakfalls"),
   ("Conditioning", "10 sprawls back to stance in 30 seconds", "Sprawl drill")]],
 "Boxing 11-14": [
  [("Stance and guard", "Good stance and guard, weight on the balls of the feet", "Stance and guard"),
   ("Footwork", "Step-drag in all four directions in stance", "Step-drag footwork"),
   ("Jab and 1-2", "Jab and 1-2 on the pads with the hands back to the face", "1-2 on pads"),
   ("Rope-less skipping", "Two minutes of jump-rope footwork (no rope)", "Jump rope"),
   ("Wraps and rules", "Wraps own hands correctly and knows the gym's safety rules", "")],
  [("1-2-3", "Jab, cross, lead hook on the pads", "1-2-3 on pads"),
   ("Called combinations", "Throws the coach's called combinations without dropping the hands", "Coach-called combinations"),
   ("Slip and counter", "Slips both sides on the line with a counter", "Slip line"),
   ("Bag rounds", "Two 2-minute bag rounds", "Heavy bag rounds"),
   ("Footwork ladder", "Ladder patterns at speed, guard up", "Footwork ladder and cones")],
  [("Bob and weave", "Bob and weave under the rope with the knees", "Bob and weave under the rope"),
   ("Body touch control", "Body-only touch drill: stays light and in control", "Body-only partner touch"),
   ("Step out", "Finishes every combination with a step out, hands up", "Coach-called combinations"),
   ("Three bag rounds", "Three 2-minute bag rounds with good form", "Heavy bag rounds"),
   ("Angles", "Pivots to an angle after a combination", "Footwork ladder and cones")],
  [("Leads the warm-up", "Leads the class warm-up", "Dynamic warm-up"),
   ("Holds pads safely", "Holds the pads safely for a partner's jab and 1-2", "Jab on pads"),
   ("Three light rounds", "Three body-only touch rounds, staying light", "Body-only partner touch"),
   ("Helps newer athletes", "Helps a newer athlete with stance and footwork", ""),
   ("Safety leader", "Can explain the gym's safety rules, including what to do after a hit to the head", "")]],
}
drills = {d[0] for d in DRILLS}
for prog, levels in P.items():
    for lv in levels:
        for s, h, d in lv: assert d == "" or d in drills, (prog, d)

wb = Workbook()
def banner(ws, title, sub, width):
    for r in (1, 2):
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=width)
        for c in range(1, width + 1): ws.cell(r, c).fill = fill_navy
    ws["A1"] = title; ws["A1"].font = f_title; ws["A1"].alignment = Alignment(vertical="center")
    ws["A2"] = sub; ws["A2"].font = f_sub; ws["A2"].alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[1].height = 30; ws.row_dimensions[2].height = 32
def head(ws, r, cols):
    for i, (t, w) in enumerate(cols, 1):
        c = ws.cell(r, i, t); c.font = f_head; c.fill = fill_navy; c.border = box; c.alignment = center
        ws.column_dimensions[L(i)].width = w
    ws.row_dimensions[r].height = 30
def inp(c, v=None, fmt=None):
    c.value = v; c.font = f_in; c.fill = fill_in; c.border = box; c.alignment = wrap
    if fmt: c.number_format = fmt
def calc(c, v, fmt=None, al=None):
    c.value = v; c.font = f_calc; c.border = box; c.alignment = al or Alignment(vertical="top", wrap_text=True)
    if fmt: c.number_format = fmt
def dv_list(ws, items, rng):
    d = DataValidation(type="list", formula1='"' + ",".join(items) + '"', allow_blank=True); ws.add_data_validation(d); d.add(rng)
def dv_range(ws, ref, rng):
    d = DataValidation(type="list", formula1=ref, allow_blank=True); ws.add_data_validation(d); d.add(rng)

# ---------------- Start Here
ws0 = wb.active; ws0.title = "Start Here"
banner(ws0, "BBG SKILLS PASSPORT", "Four levels for each kids' and youth class (Rookie, Contender, Challenger, Champion). Coaches tick off skills; parents get a progress card.", 3)
ws0.column_dimensions["A"].width = 4; ws0.column_dimensions["B"].width = 26; ws0.column_dimensions["C"].width = 96
steps = [
 ("1", "Levels", "Kids Boxing, Wrestling and Boxing 11-14 each have four levels of five skills. Most skills match a drill in the Coaching Library, so they're taught before they're checked. A few (gym rules, sportsmanship, helping newer athletes) are learned in every class. None of them involve hitting to the head."),
 ("2", "Roster", "One row per athlete: first name only (add a number if two kids share a name, e.g. Liam 2), their class and the date they started. No last names, birthdays, phone numbers or health details."),
 ("3", "Log a skill", "When a coach sees a skill done well, add a row on the Progress tab: date, athlete, skill, coach. Log each skill once. The Roster shows each athlete's level and how many skills are left."),
 ("4", "Progress card", "On the Card tab, pick an athlete: it shows their level, what they've passed and what's next. Screenshot it or ask Claude for a parent message to go with it. Share it with the parent, never post it publicly."),
 ("5", "Levelling up", "When all five skills of a level are passed, the athlete moves up. A small moment in class (a shout-out, a sticker) is plenty. No prizes or money."),
]
r = 4
for n, a, b in steps:
    ws0.cell(r, 1, n).font = Font(name=F, size=12, bold=True, color=GOLD); ws0.cell(r, 2, a).font = f_bold
    ws0.cell(r, 1).alignment = Alignment(vertical="top"); ws0.cell(r, 2).alignment = Alignment(vertical="top", wrap_text=True)
    c = ws0.cell(r, 3, b); c.font = f_lab; c.alignment = wrap; ws0.row_dimensions[r].height = 34; r += 1
r += 1
ws0.cell(r, 2, "Colours").font = f_sec; r += 1
for lab, fill, t in (("Yellow, blue text", fill_in, "You type or pick here"), ("Black, bold", None, "Calculated: don't type over it")):
    c = ws0.cell(r, 2, lab); c.font = f_in if fill else f_lab
    if fill: c.fill = fill
    ws0.cell(r, 3, t).font = f_lab; r += 1

# ---------------- Levels
wl = wb.create_sheet("Levels")
LC = [("Class", 16), ("Level", 7), ("Level name", 12), ("Skill", 24), ("How it's checked", 60), ("Drill in the Coaching Library", 30), ("Key", 6)]
banner(wl, "LEVELS", "The skills for each level. Change the wording if you like, but keep each class's skills together and in level order.", len(LC))
head(wl, 4, LC)
rows = []
for prog, levels in P.items():
    for k, lv in enumerate(levels, 1):
        for s, h, d in lv: rows.append((prog, k, LEVELS[k - 1], s, h, d))
LMAX = 4 + len(rows) + 30
for k in range(LMAX - 4):
    rr = 5 + k
    row = rows[k] if k < len(rows) else (None,) * 6
    for i, v in enumerate(row, 1): inp(wl.cell(rr, i), v)
    calc(wl.cell(rr, 7), f'=IF(D{rr}="","",A{rr}&" | "&D{rr})')
    wl.cell(rr, 7).font = f_note
    if k < len(rows): wl.row_dimensions[rr].height = 28
dv_list(wl, LEVELS, f"C5:C{LMAX}")
wl.column_dimensions["G"].hidden = True
wl.freeze_panes = "A5"
LA = f"Levels!$A$5:$A${LMAX}"; LB = f"Levels!$B$5:$B${LMAX}"; LD = f"Levels!$D$5:$D${LMAX}"; LG = f"Levels!$G$5:$G${LMAX}"; LE = f"Levels!$E$5:$E${LMAX}"
PROGS = list(P.keys())

# ---------------- Roster
wr = wb.create_sheet("Roster")
RC = [("Athlete (first name)", 18), ("Class", 16), ("Started", 12), ("Coach", 11), ("Rookie", 9), ("Contender", 10), ("Challenger", 10), ("Champion", 10),
      ("Current level", 14), ("Skills left in level", 10), ("Last skill passed", 13), ("Notes (no health details)", 36)]
banner(wr, "ROSTER", "First names only. Level counts and the current level fill in from the Progress tab.", len(RC))
head(wr, 4, RC)
RMAX = 4 + 150
for rr in range(5, RMAX + 1):
    inp(wr.cell(rr, 1)); inp(wr.cell(rr, 2)); inp(wr.cell(rr, 3), None, DATE); inp(wr.cell(rr, 4))
    for k in range(4):
        f = (f'=IF(OR($A{rr}="",$B{rr}=""),"",SUMPRODUCT(({LA}=$B{rr})*({LB}={k+1})*(COUNTIFS(Progress!$B$5:$B$3000,$A{rr},Progress!$G$5:$G$3000,{LG})>0))'
             f'&" / "&COUNTIFS({LA},$B{rr},{LB},{k+1}))')
        calc(wr.cell(rr, 5 + k), f, None, center)
    # numbers for current level
    def done(k): return f'SUMPRODUCT(({LA}=$B{rr})*({LB}={k})*(COUNTIFS(Progress!$B$5:$B$3000,$A{rr},Progress!$G$5:$G$3000,{LG})>0))'
    def tot(k): return f'COUNTIFS({LA},$B{rr},{LB},{k})'
    cur = (f'=IF(OR($A{rr}="",$B{rr}=""),"",IF({done(1)}<{tot(1)},"1 Rookie",IF({done(2)}<{tot(2)},"2 Contender",'
           f'IF({done(3)}<{tot(3)},"3 Challenger",IF({done(4)}<{tot(4)},"4 Champion","All done")))))')
    calc(wr.cell(rr, 9), cur, None, center)
    left = (f'=IF(OR(I{rr}="",I{rr}="All done"),"",'
            f'{tot(1)}*(LEFT(I{rr},1)="1")+{tot(2)}*(LEFT(I{rr},1)="2")+{tot(3)}*(LEFT(I{rr},1)="3")+{tot(4)}*(LEFT(I{rr},1)="4")'
            f'-({done(1)}*(LEFT(I{rr},1)="1")+{done(2)}*(LEFT(I{rr},1)="2")+{done(3)}*(LEFT(I{rr},1)="3")+{done(4)}*(LEFT(I{rr},1)="4")))')
    calc(wr.cell(rr, 10), left, "0", center)
    calc(wr.cell(rr, 11), f'=IF($A{rr}="","",IF(COUNTIFS(Progress!$B$5:$B$3000,$A{rr})=0,"",SUMPRODUCT(MAX((Progress!$B$5:$B$3000=$A{rr})*Progress!$A$5:$A$3000))))', DATE, center)
    inp(wr.cell(rr, 12))
dv_list(wr, PROGS, f"B5:B{RMAX}"); dv_list(wr, ["Cody", "Sean", "Jordan"], f"D5:D{RMAX}")
wr.conditional_formatting.add(f"I5:I{RMAX}", CellIsRule(operator="equal", formula=['"All done"'], fill=PatternFill("solid", fgColor="C6EFCE"), font=Font(name=F, color="006100", bold=True)))
wr.conditional_formatting.add(f"J5:J{RMAX}", CellIsRule(operator="equal", formula=["1"], fill=PatternFill("solid", fgColor="FFEB9C"), font=Font(name=F, color="9C5700", bold=True)))
wr.freeze_panes = "B5"

# ---------------- Progress
wp = wb.create_sheet("Progress")
PC = [("Date passed", 12), ("Athlete", 18), ("Skill", 26), ("Coach", 11), ("Note", 36), ("Class", 16), ("Key", 6), ("Level", 10)]
banner(wp, "PROGRESS", "One row each time a coach sees a skill done well. Pick the athlete and the skill; class and level fill in.", len(PC))
head(wp, 4, PC)
for rr in range(5, 3001):
    inp(wp.cell(rr, 1), None, DATE); inp(wp.cell(rr, 2)); inp(wp.cell(rr, 3)); inp(wp.cell(rr, 4)); inp(wp.cell(rr, 5))
    calc(wp.cell(rr, 6), f'=IF(B{rr}="","",IFERROR(INDEX(Roster!$B$5:$B${RMAX},MATCH(B{rr},Roster!$A$5:$A${RMAX},0)),"Not on roster"))')
    calc(wp.cell(rr, 7), f'=IF(OR(B{rr}="",C{rr}=""),"",F{rr}&" | "&C{rr})'); wp.cell(rr, 7).font = f_note
    calc(wp.cell(rr, 8), f'=IF(G{rr}="","",IFERROR(INDEX(Levels!$C$5:$C${LMAX},MATCH(G{rr},{LG},0)),"Not this class\'s skill"))')
dv_range(wp, f"=Roster!$A$5:$A${RMAX}", "B5:B3000"); dv_range(wp, f"={LD}", "C5:C3000"); dv_list(wp, ["Cody", "Sean", "Jordan"], "D5:D3000")
d = DataValidation(type="date", allow_blank=True); wp.add_data_validation(d); d.add("A5:A3000")
wp.conditional_formatting.add("F5:H3000", FormulaRule(formula=['OR(F5="Not on roster",H5="Not this class\'s skill")'], fill=PatternFill("solid", fgColor="FFC7CE"), font=Font(name=F, color="9C0006", bold=True)))
wp.column_dimensions["G"].hidden = True
wp.freeze_panes = "C5"

# ---------------- Card
wc = wb.create_sheet("Card")
for c, w in zip("ABCDE", (4, 24, 58, 14, 4)): wc.column_dimensions[c].width = w
for r in range(1, 40):
    for c in range(1, 6): wc.cell(r, c).fill = PatternFill("solid", fgColor="FFFFFF")
for r in range(1, 5):
    for c in range(1, 6): wc.cell(r, c).fill = fill_navy
wc["B2"] = "BBG SKILLS PASSPORT"; wc["B2"].font = Font(name=F, size=18, bold=True, color=GOLD)
wc["B3"] = "Bancroft Boxing & Grappling"; wc["B3"].font = Font(name=F, size=10, color="FFFFFF")
wc.row_dimensions[2].height = 30
wc["B6"] = "Athlete"; wc["B6"].font = f_bold
inp(wc["C6"], None); wc["C6"].font = Font(name=F, size=14, bold=True, color="0000FF")
dv_range(wc, f"=Roster!$A$5:$A${RMAX}", "C6")
wc["D6"] = "← pick a name"; wc["D6"].font = f_note
wc["B7"] = "Class"; wc["B7"].font = f_bold
calc(wc["C7"], f'=IF($C$6="","",IFERROR(INDEX(Roster!$B$5:$B${RMAX},MATCH($C$6,Roster!$A$5:$A${RMAX},0)),""))')
wc["B8"] = "Level now"; wc["B8"].font = f_bold
calc(wc["C8"], f'=IF($C$6="","",IFERROR(INDEX(Roster!$I$5:$I${RMAX},MATCH($C$6,Roster!$A$5:$A${RMAX},0)),""))')
wc["C8"].font = Font(name=F, size=14, bold=True, color=NAVY)
for c in ("C7", "C8"): wc[c].border = Border()
# hidden helpers: first row of the class in Levels and how many skills it has
wc["G1"] = "first"; wc["H1"] = "count"
for c in ("G1", "H1", "G2", "H2"): wc[c].font = Font(name=F, size=8, color="FFFFFF")
wc["G2"] = f'=IF($C$7="","",MATCH($C$7,{LA},0))'; wc["H2"] = f'=IF($C$7="","",COUNTIF({LA},$C$7))'
wc.column_dimensions["G"].hidden = True; wc.column_dimensions["H"].hidden = True
hr = 10
for i, h in enumerate(["Level", "Skill", "Done"], 2):
    c = wc.cell(hr, i, h); c.font = f_head; c.fill = fill_navy; c.alignment = center; c.border = box
for k in range(20):
    r = hr + 1 + k
    idx = f"($G$2+{k})"
    wc.cell(r, 2, f'=IF(OR($C$7="",{k}>=N($H$2)),"",INDEX(Levels!$B$5:$B${LMAX},{idx})&" "&INDEX(Levels!$C$5:$C${LMAX},{idx}))')
    wc.cell(r, 3, f'=IF(OR($C$7="",{k}>=N($H$2)),"",INDEX({LD},{idx})&": "&INDEX({LE},{idx}))')
    wc.cell(r, 4, f'=IF(OR($C$7="",{k}>=N($H$2)),"",IF(COUNTIFS(Progress!$B$5:$B$3000,$C$6,Progress!$G$5:$G$3000,INDEX({LG},{idx}))>0,"✓ Passed",IF(IFERROR(INDEX(Levels!$B$5:$B${LMAX},{idx})=VALUE(LEFT($C$8,1)),FALSE),"Next up","Later")))')
    for c in (2, 3, 4):
        cc = wc.cell(r, c); cc.border = box; cc.font = f_lab if c != 4 else f_bold; cc.alignment = Alignment(vertical="center", wrap_text=True)
    wc.row_dimensions[r].height = 30
wc.conditional_formatting.add(f"D{hr+1}:D{hr+20}", CellIsRule(operator="equal", formula=['"✓ Passed"'], fill=PatternFill("solid", fgColor="C6EFCE"), font=Font(name=F, color="006100", bold=True)))
wc.conditional_formatting.add(f"B{hr+1}:D{hr+20}", FormulaRule(formula=[f'$D{hr+1}="Next up"'], fill=fill_goldl))
wc.cell(hr + 22, 2, "Share with the parent only. First names only; never post a child's card publicly.").font = f_note
wc.sheet_view.showGridLines = False

for w in wb.worksheets:
    w.page_setup.orientation = "portrait" if w.title == "Card" else "landscape"
    w.sheet_properties.pageSetUpPr.fitToPage = True; w.page_setup.fitToWidth = 1; w.page_setup.fitToHeight = 0
wb._sheets = [wb["Start Here"], wb["Card"], wb["Roster"], wb["Progress"], wb["Levels"]]
wb.active = 0
wb.save(OUT); print("saved", OUT, len(rows), "skills")

# ---------------- skill reference
D = os.path.join(os.path.dirname(os.path.abspath(__file__)), "skills/bbg-coach/references"); os.makedirs(D, exist_ok=True)
Lm = ["# BBG Skills Passport levels", "", "The same as the Levels tab of the BBG Skills Passport sheet. Five skills per level; an athlete moves up when all five are passed.", ""]
for prog, levels in P.items():
    Lm += [f"## {prog}", ""]
    for k, lv in enumerate(levels, 1):
        Lm.append(f"**Level {k}: {LEVELS[k-1]}**")
        Lm += [f"- {s}: {h}" + (f" (drill: {d})" if d else "") for s, h, d in lv]
        Lm.append("")
open(os.path.join(D, "passport.md"), "w").write("\n".join(Lm))
print("wrote passport.md")
