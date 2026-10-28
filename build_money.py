"""Build 'BBG & CCCL Money Summary.xlsx': paste bank / Eventbrite / Gymdesk exports, keyword rules sort each line by business and category,
and the Monthly and Categories tabs total it up. Google Sheets compatible (SUMIFS, COUNTIFS, SUMPRODUCT, INDEX; inline dropdowns without commas).
usage: python3 build_money.py OUT"""
import sys, datetime as dt
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.utils import get_column_letter as L

OUT = sys.argv[1]
NAVY, GOLD, GOLD_L = "081C2B", "EAA959", "FBEBD3"
F = "Arial"
f_title = Font(name=F, size=16, bold=True, color=GOLD); f_sub = Font(name=F, size=10, color="FFFFFF")
f_head = Font(name=F, size=10, bold=True, color="FFFFFF"); f_lab = Font(name=F, size=10); f_bold = Font(name=F, size=10, bold=True)
f_in = Font(name=F, size=10, color="0000FF"); f_note = Font(name=F, size=9, italic=True, color="555555")
f_sec = Font(name=F, size=11, bold=True, color=NAVY); f_calc = Font(name=F, size=10, bold=True, color="000000")
fill_navy = PatternFill("solid", fgColor=NAVY); fill_in = PatternFill("solid", fgColor="FFFF00"); fill_goldl = PatternFill("solid", fgColor=GOLD_L)
fill_grey = PatternFill("solid", fgColor="F2F2F2")
thin = Side(style="thin", color="BFBFBF"); box = Border(top=thin, bottom=thin, left=thin, right=thin)
wrap = Alignment(wrap_text=True, vertical="top"); MONEY = '$#,##0.00;[Red]-$#,##0.00;"-"'; MONEY0 = '$#,##0;[Red]-$#,##0;"-"'

BUSINESSES = ["BBG", "CCCL", "Shared", "Exclude", "Check"]
INCOME = ["Memberships and classes", "Drop-ins and trials", "Personal training", "Parties and camps", "Ticket sales", "Sponsorship", "Merch", "Livestream", "Grants", "Other income"]
EXPENSE = ["Rent", "Utilities", "Phone and internet", "Insurance", "Equipment", "Software and subscriptions", "Marketing and ads", "Venue", "Officials and medical",
           "Permits and fees", "Event production", "Coach and staff pay", "Supplies", "Travel", "Bank and payment fees", "Other expense"]
OTHER = ["Transfer", "Owner money in or out", "Personal"]
CATS = INCOME + EXPENSE + OTHER
SOURCES = ["Bank", "Credit card", "Eventbrite", "Gymdesk", "Cash", "Other"]
# keyword, business, category, note  (starting rules: edit them to match the bank's wording)
RULES = [
 ("EVENTBRITE", "CCCL", "Ticket sales", "Eventbrite payouts (already after Eventbrite's fees)"),
 ("GYMDESK", "BBG", "Memberships and classes", "Gymdesk payouts; check the wording on your statement"),
 ("ONTARIO ATHLETICS", "CCCL", "Permits and fees", "Athletics Commission permit"),
 ("FACEBK", "Check", "Marketing and ads", "Meta ads: pick BBG or CCCL for each one"),
 ("META ", "Check", "Marketing and ads", "Meta ads: pick BBG or CCCL"),
 ("CANVA", "Shared", "Software and subscriptions", ""),
 ("GOOGLE", "Shared", "Software and subscriptions", "Google Workspace or storage"),
 ("ANTHROPIC", "Shared", "Software and subscriptions", "Claude subscription"),
 ("CLAUDE.AI", "Shared", "Software and subscriptions", "Claude subscription"),
 ("HYDRO", "BBG", "Utilities", "Gym hydro"),
 ("ENBRIDGE", "BBG", "Utilities", "Gas"),
 ("BELL ", "Shared", "Phone and internet", ""),
 ("ROGERS", "Shared", "Phone and internet", ""),
 ("INSURANCE", "Check", "Insurance", "Gym or event insurance: pick the business"),
 ("SERVICE CHARGE", "Shared", "Bank and payment fees", ""),
 ("MONTHLY FEE", "Shared", "Bank and payment fees", ""),
 ("E-TRANSFER", "Check", "Other income", "Interac e-Transfers: could be a member, a sponsor or a transfer. Pick each one"),
 ("ETRANSFER", "Check", "Other income", "Same as above"),
 ("TRANSFER TO", "Exclude", "Transfer", "Moving money between your own accounts isn't income or spending"),
 ("TRANSFER FROM", "Exclude", "Transfer", ""),
 ("AMAZON", "Check", "Supplies", "Pick the business"),
 ("CANADIAN TIRE", "Check", "Supplies", ""),
 ("HOME DEPOT", "BBG", "Supplies", ""),
 ("STAPLES", "Shared", "Supplies", "Printing, office"),
 ("ESSO", "Shared", "Travel", "Fuel"),
 ("PETRO", "Shared", "Travel", "Fuel"),
 ("SHELL", "Shared", "Travel", "Fuel"),
]

wb = Workbook()
def banner(ws, title, sub, width):
    for r in (1, 2):
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=width)
        for c in range(1, width + 1): ws.cell(r, c).fill = fill_navy
    ws["A1"] = title; ws["A1"].font = f_title; ws["A1"].alignment = Alignment(vertical="center")
    ws["A2"] = sub; ws["A2"].font = f_sub; ws["A2"].alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[1].height = 30; ws.row_dimensions[2].height = 32
def head(ws, r, cols, h=30):
    for i, (t, w) in enumerate(cols, 1):
        c = ws.cell(r, i, t); c.font = f_head; c.fill = fill_navy; c.border = box; c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        ws.column_dimensions[L(i)].width = w
    ws.row_dimensions[r].height = h
def inp(c, v=None, fmt=None):
    c.value = v; c.font = f_in; c.fill = fill_in; c.border = box; c.alignment = wrap
    if fmt: c.number_format = fmt
def calc(c, v, fmt=None, bold=True):
    c.value = v; c.font = f_calc if bold else f_lab; c.border = box; c.alignment = Alignment(vertical="top")
    if fmt: c.number_format = fmt
def dv_list(ws, items, rng):
    d = DataValidation(type="list", formula1='"' + ",".join(items) + '"', allow_blank=True); ws.add_data_validation(d); d.add(rng)
def dv_range(ws, ref, rng):
    d = DataValidation(type="list", formula1=ref, allow_blank=True); ws.add_data_validation(d); d.add(rng)

# ---------------- Start Here
ws0 = wb.active; ws0.title = "Start Here"
banner(ws0, "BBG & CCCL MONEY SUMMARY", "Paste what the bank, Eventbrite and Gymdesk export. Each line sorts itself by business and category, and the Monthly tab adds it up.", 3)
ws0.column_dimensions["A"].width = 4; ws0.column_dimensions["B"].width = 26; ws0.column_dimensions["C"].width = 96
steps = [
 ("1", "Export", "Once a month, download the transactions as a CSV from the business bank account (and the credit card), plus the Eventbrite payout report and the Gymdesk payments report if you want the detail."),
 ("2", "Delete account numbers first", "Bank CSVs often include the account or card number in a column. Delete that column before pasting. Never paste account numbers, card numbers or anyone's banking details into this sheet."),
 ("3", "Paste", "On the Transactions tab, paste the date, description and amount into the yellow columns (money in is positive, money out is negative). Pick where it came from."),
 ("4", "Let the rules sort it", "Business and Category fill in from the Rules tab (the first keyword found in the description wins). Anything marked Check needs a quick decision: type it in the Override columns."),
 ("5", "Read the Monthly tab", "Income, spending and what's left for BBG and CCCL each month, plus shared costs. The Categories tab shows where the money goes over the year."),
 ("6", "Keep the rules growing", "When the same kind of line keeps showing up as Check, add a keyword for it on the Rules tab. After a couple of months almost everything sorts itself."),
 ("7", "Second Chance", "Second Chance is a separate nonprofit: give it its own bank account and its own copy of this sheet. Never mix its money with BBG or CCCL."),
 ("8", "What this is (and isn't)", "A clear monthly picture, built from what actually went through the bank (so only money that was paid counts). It isn't bookkeeping or tax advice: your accountant still does the books and HST."),
]
r = 4
for n, a, b in steps:
    ws0.cell(r, 1, n).font = Font(name=F, size=12, bold=True, color=GOLD); ws0.cell(r, 2, a).font = f_bold
    ws0.cell(r, 1).alignment = Alignment(vertical="top"); ws0.cell(r, 2).alignment = Alignment(vertical="top", wrap_text=True)
    c = ws0.cell(r, 3, b); c.font = f_lab; c.alignment = wrap; ws0.row_dimensions[r].height = 34; r += 1
r += 1
ws0.cell(r, 2, "Business choices").font = f_sec; r += 1
for a, b in (("BBG", "The gym"), ("CCCL", "The fight league and its shows"), ("Shared", "Costs that serve both (phone, software, fuel, bank fees). Shown separately on the Monthly tab."),
             ("Exclude", "Not income or spending: moving money between your own accounts, owner money in or out, personal"), ("Check", "The rules couldn't tell: decide and type it in Override")):
    ws0.cell(r, 2, a).font = f_bold; c = ws0.cell(r, 3, b); c.font = f_lab; c.alignment = wrap; r += 1
r += 1
ws0.cell(r, 2, "Colours").font = f_sec; r += 1
for lab, fill, txt in (("Yellow, blue text", fill_in, "You type or paste here"), ("Black, bold", None, "Calculated: don't type over it")):
    c = ws0.cell(r, 2, lab); c.font = f_in if fill else f_lab
    if fill: c.fill = fill
    ws0.cell(r, 3, txt).font = f_lab; r += 1

# ---------------- Lists (hidden)
wl = wb.create_sheet("Lists")
wl["A1"] = "Category"; wl["B1"] = "Type"; wl["D1"] = "Business"
for c in ("A1", "B1", "D1"): wl[c].font = f_bold
for i, c in enumerate(CATS):
    wl.cell(2 + i, 1, c); wl.cell(2 + i, 2, "Income" if c in INCOME else ("Expense" if c in EXPENSE else "Other"))
for i, b in enumerate(BUSINESSES): wl.cell(2 + i, 4, b)
CN = 1 + len(CATS)

# ---------------- Rules
wr = wb.create_sheet("Rules")
RC = [("Keyword in the description", 30), ("Business", 12), ("Category", 26), ("Notes", 60)]
banner(wr, "RULES", "If a transaction's description contains the keyword (capitals don't matter), it gets this business and category. The first rule that matches wins, so put specific keywords above general ones.", len(RC))
head(wr, 4, RC)
RMAX = 4 + 200
for k in range(RMAX - 4):
    rr = 5 + k
    row = RULES[k] if k < len(RULES) else (None, None, None, None)
    for i, v in enumerate(row, 1): inp(wr.cell(rr, i), v)
dv_list(wr, BUSINESSES, f"B5:B{RMAX}"); dv_range(wr, f"=Lists!$A$2:$A${CN}", f"C5:C{RMAX}")
wr.freeze_panes = "A5"
RK = f"Rules!$A$5:$A${RMAX}"

# ---------------- Transactions
wt = wb.create_sheet("Transactions")
TC = [("Date", 11), ("Description (from the bank)", 44), ("Amount (+ in, - out)", 13), ("From", 11), ("Rule business", 11), ("Rule category", 22),
      ("Override business", 11), ("Override category", 22), ("Business", 11), ("Category", 22), ("Month", 10), ("Type", 9), ("Note", 30), ("Rule row", 6)]
banner(wt, "TRANSACTIONS", "Paste date, description and amount into the yellow columns. Business and Category come from the Rules tab; type in Override to change one.", len(TC))
head(wt, 4, TC)
TMAX = 4 + 3000
for rr in range(5, TMAX + 1):
    inp(wt.cell(rr, 1), None, "yyyy-mm-dd"); inp(wt.cell(rr, 2)); inp(wt.cell(rr, 3), None, MONEY); inp(wt.cell(rr, 4))
    # first matching rule: 1000-row trick returns the highest value for the earliest row
    calc(wt.cell(rr, 14), f'=IF(B{rr}="","",IFERROR(1000-SUMPRODUCT(MAX(ISNUMBER(SEARCH({RK},B{rr}))*({RK}<>"")*(1000-ROW({RK})))),""))', "0", bold=False)
    calc(wt.cell(rr, 5), f'=IF(OR(B{rr}="",N(N{rr})=0,N{rr}=1000),IF(B{rr}="","","Check"),INDEX(Rules!$B:$B,N{rr}))', None, bold=False)
    calc(wt.cell(rr, 6), f'=IF(OR(B{rr}="",N(N{rr})=0,N{rr}=1000),"",INDEX(Rules!$C:$C,N{rr}))', None, bold=False)
    inp(wt.cell(rr, 7)); inp(wt.cell(rr, 8))
    calc(wt.cell(rr, 9), f'=IF(B{rr}="","",IF(G{rr}<>"",G{rr},E{rr}))')
    calc(wt.cell(rr, 10), f'=IF(B{rr}="","",IF(H{rr}<>"",H{rr},F{rr}))', None, bold=False)
    calc(wt.cell(rr, 11), f'=IF(A{rr}="","",DATE(YEAR(A{rr}),MONTH(A{rr}),1))', "mmm yyyy", bold=False)
    calc(wt.cell(rr, 12), f'=IF(J{rr}="",IF(B{rr}="","",IF(N(C{rr})>=0,"Income","Expense")),IFERROR(INDEX(Lists!$B$2:$B${CN},MATCH(J{rr},Lists!$A$2:$A${CN},0)),""))', None, bold=False)
    inp(wt.cell(rr, 13))
dv_list(wt, SOURCES, f"D5:D{TMAX}"); dv_list(wt, BUSINESSES, f"G5:G{TMAX}"); dv_range(wt, f"=Lists!$A$2:$A${CN}", f"H5:H{TMAX}")
wt.conditional_formatting.add(f"I5:I{TMAX}", CellIsRule(operator="equal", formula=['"Check"'], fill=PatternFill("solid", fgColor="FFC7CE"), font=Font(name=F, color="9C0006", bold=True)))
wt.conditional_formatting.add(f"I5:I{TMAX}", CellIsRule(operator="equal", formula=['"Exclude"'], fill=fill_grey, font=Font(name=F, color="808080")))
wt.column_dimensions["N"].hidden = True
wt.freeze_panes = "C5"; wt.auto_filter.ref = f"A4:M{TMAX}"
T = lambda col: f"Transactions!${col}$5:${col}${TMAX}"

# ---------------- Monthly
wm = wb.create_sheet("Monthly")
MC = [("Month", 12), ("BBG in", 12), ("BBG out", 12), ("BBG left", 12), ("CCCL in", 12), ("CCCL out", 12), ("CCCL left", 12), ("Shared costs (net)", 12), ("All left", 13), ("Lines to check", 10)]
banner(wm, "MONTHLY", "Money in, money out and what's left, by business. Excludes transfers, owner money and personal. Shared costs (phone, software, fuel, bank fees) are shown on their own.", len(MC))
wm["A3"] = "Year"; wm["A3"].font = f_bold
inp(wm["B3"], 2026, "0"); wm["C3"] = "Change the year to see a different one."; wm["C3"].font = f_note
head(wm, 4, MC)
for i in range(12):
    r = 5 + i
    calc(wm.cell(r, 1), f"=DATE($B$3,{i+1},1)", "mmm yyyy")
    def s(bus, sign):
        return f'=SUMIFS({T("C")},{T("I")},"{bus}",{T("K")},$A{r},{T("C")},"{sign}0")'
    calc(wm.cell(r, 2), s("BBG", ">"), MONEY0); calc(wm.cell(r, 3), f'=-SUMIFS({T("C")},{T("I")},"BBG",{T("K")},$A{r},{T("C")},"<0")', MONEY0)
    calc(wm.cell(r, 4), f"=B{r}-C{r}", MONEY0)
    calc(wm.cell(r, 5), s("CCCL", ">"), MONEY0); calc(wm.cell(r, 6), f'=-SUMIFS({T("C")},{T("I")},"CCCL",{T("K")},$A{r},{T("C")},"<0")', MONEY0)
    calc(wm.cell(r, 7), f"=E{r}-F{r}", MONEY0)
    calc(wm.cell(r, 8), f'=-SUMIFS({T("C")},{T("I")},"Shared",{T("K")},$A{r})', MONEY0)
    calc(wm.cell(r, 9), f"=D{r}+G{r}-H{r}", MONEY0)
    calc(wm.cell(r, 10), f'=COUNTIFS({T("I")},"Check",{T("K")},$A{r})', "0")
tr = 17
wm.cell(tr, 1, "Year").font = f_bold
for i in range(2, 11):
    c = wm.cell(tr, i, f"=SUM({L(i)}5:{L(i)}16)"); c.font = f_bold; c.border = box; c.number_format = MONEY0 if i < 10 else "0"
for i in range(1, 11): wm.cell(tr, i).fill = fill_goldl
wm.conditional_formatting.add("J5:J17", CellIsRule(operator="greaterThan", formula=["0"], fill=PatternFill("solid", fgColor="FFC7CE"), font=Font(name=F, color="9C0006", bold=True)))
r = tr + 2
wm.cell(r, 1, "Good to know").font = f_sec; r += 1
for t in ("'Left' is money in minus money out for the month. It isn't profit for tax purposes (no HST, depreciation or owner pay adjustments).",
          "A show's money lands in the months it moved: sponsor payments when they paid, Eventbrite when it paid out, costs when they were paid.",
          "Lines to check: rows still marked Check. Sort them on the Transactions tab and the totals update."):
    c = wm.cell(r, 1, t); c.font = f_note; wm.merge_cells(start_row=r, start_column=1, end_row=r, end_column=10); r += 1
wm.freeze_panes = "B5"

# ---------------- Categories
wk = wb.create_sheet("Categories")
KC = [("Category", 26), ("Type", 9), ("BBG", 13), ("CCCL", 13), ("Shared", 13), ("Total", 13)]
banner(wk, "CATEGORIES", "Where the money comes from and goes, for the year on the Monthly tab.", len(KC))
head(wk, 4, KC)
YR_START = "DATE(Monthly!$B$3,1,1)"; YR_END = "DATE(Monthly!$B$3,12,1)"
for i, cat in enumerate(INCOME + EXPENSE):
    r = 5 + i
    calc(wk.cell(r, 1), cat, None, bold=True); calc(wk.cell(r, 2), "Income" if cat in INCOME else "Expense", None, bold=False)
    for j, bus in enumerate(("BBG", "CCCL", "Shared")):
        calc(wk.cell(r, 3 + j), f'=SUMIFS({T("C")},{T("J")},$A{r},{T("I")},"{bus}",{T("K")},">="&{YR_START},{T("K")},"<="&{YR_END})', MONEY0)
    calc(wk.cell(r, 6), f"=SUM(C{r}:E{r})", MONEY0)
    if cat == INCOME[-1]:
        pass
kr = 5 + len(INCOME + EXPENSE)
wk.cell(kr, 1, "Income total").font = f_bold
wk.cell(kr + 1, 1, "Spending total").font = f_bold
for j in range(3, 7):
    col = L(j)
    c = wk.cell(kr, j, f'=SUMIFS({col}5:{col}{kr-1},$B$5:$B${kr-1},"Income")'); c.font = f_bold; c.number_format = MONEY0; c.border = box
    c = wk.cell(kr + 1, j, f'=SUMIFS({col}5:{col}{kr-1},$B$5:$B${kr-1},"Expense")'); c.font = f_bold; c.number_format = MONEY0; c.border = box
for rr in (kr, kr + 1):
    for j in range(1, 7): wk.cell(rr, j).fill = fill_goldl
wk.cell(kr + 3, 1, "Spending shows as negative numbers. Uncategorized lines (Check) and Exclude lines aren't counted here.").font = f_note
wk.freeze_panes = "C5"

for w in wb.worksheets:
    w.page_setup.orientation = "landscape"; w.sheet_properties.pageSetUpPr.fitToPage = True; w.page_setup.fitToWidth = 1; w.page_setup.fitToHeight = 0
wb._sheets = [wb["Start Here"], wb["Monthly"], wb["Categories"], wb["Transactions"], wb["Rules"], wb["Lists"]]
wb["Lists"].sheet_state = "hidden"
wb.active = 0
wb.save(OUT); print("saved", OUT, len(RULES), "rules")
