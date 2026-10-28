"""Build 'CCCL Event Profit Predictor.xlsx' (Google Sheets compatible: no post-2007 functions)."""
import sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.comments import Comment
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule, FormulaRule, ColorScaleRule
from openpyxl.utils import get_column_letter as L

OUT = sys.argv[1]
TOWNS = sys.argv[2] if len(sys.argv) > 2 else None
VENUES27 = sys.argv[3] if len(sys.argv) > 3 else None
NAVY, GOLD, GOLD_L, GREY = "081C2B", "EAA959", "FBEBD3", "F2F2F2"
F = "Arial"
f_title = Font(name=F, size=16, bold=True, color=GOLD)
f_sub = Font(name=F, size=10, color="FFFFFF")
f_head = Font(name=F, size=11, bold=True, color="FFFFFF")
f_sec = Font(name=F, size=11, bold=True, color=NAVY)
f_lab = Font(name=F, size=10, color="000000")
f_bold = Font(name=F, size=10, bold=True)
f_in = Font(name=F, size=10, color="0000FF")          # typed inputs
f_link = Font(name=F, size=10, color="008000")        # pulled from another tab
f_calc = Font(name=F, size=10, color="000000")
f_note = Font(name=F, size=9, italic=True, color="555555")
fill_navy = PatternFill("solid", fgColor=NAVY)
fill_gold = PatternFill("solid", fgColor=GOLD)
fill_goldl = PatternFill("solid", fgColor=GOLD_L)
fill_in = PatternFill("solid", fgColor="FFFF00")      # edit me
fill_grey = PatternFill("solid", fgColor=GREY)
thin = Side(style="thin", color="BFBFBF")
box = Border(top=thin, bottom=thin, left=thin, right=thin)
USD = '$#,##0;($#,##0);"-"'
NUM = '#,##0;(#,##0);"-"'
PCT = '0%;(0%);"-"'
wrap = Alignment(wrap_text=True, vertical="top")
center = Alignment(horizontal="center", vertical="center", wrap_text=True)

wb = Workbook()

def banner(ws, title, sub, width):
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=width)
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=width)
    ws["A1"] = title; ws["A1"].font = f_title
    ws["A2"] = sub; ws["A2"].font = f_sub
    for r in (1, 2):
        for c in range(1, width + 1):
            ws.cell(r, c).fill = fill_navy
    ws.row_dimensions[1].height = 30
    ws.row_dimensions[2].height = 30
    ws["A2"].alignment = Alignment(wrap_text=True, vertical="center")
    ws["A1"].alignment = Alignment(vertical="center")

def section(ws, r, text, width):
    ws.cell(r, 1, text).font = f_head
    for c in range(1, width + 1):
        ws.cell(r, c).fill = fill_navy if c == 1 else fill_navy
    ws.row_dimensions[r].height = 20

def inp(cell, value, fmt=None, key=True):
    cell.value = value
    cell.font = f_in
    if key: cell.fill = fill_in
    if fmt: cell.number_format = fmt
    cell.border = box

def note(ws, r, c, text):
    ws.cell(r, c, text).font = f_note
    ws.cell(r, c).alignment = wrap

# ------------------------------------------------------------------ Inputs
wsI = wb.active; wsI.title = "Inputs"
NV = 6                      # venue columns on Inputs/Compare (B..G)
NC = NV + 2                 # notes column on Inputs
W = NC
banner(wsI, "CCCL EVENT PROFIT PREDICTOR  ·  INPUTS",
       "Type only in YELLOW cells (blue text). Everything on the Compare tab updates by itself. "
       "Numbers marked ESTIMATE are placeholders: replace them with real quotes as they come in.", W)
wsI.column_dimensions["A"].width = 46
for i in range(NV):
    wsI.column_dimensions[L(2 + i)].width = 19
wsI.column_dimensions[L(NC)].width = 62

section(wsI, 4, "GLOBAL ASSUMPTIONS  (same for every venue)", W)
g = {}
rows = [
    ("low", "Low crowd: % of sellable capacity that shows up", 0.35, PCT,
     "ESTIMATE. A bad night: weather, a clash with another event, weak card."),
    ("exp", "Expected crowd: % of sellable capacity", 0.50, PCT,
     "Set from the first show: about 500 people (Cody said 400-600) in a building of about 1,000. "
     "Tune it on the 'Haliburton 2026 Check' tab after each show."),
    ("high", "High crowd: % of sellable capacity", 0.80, PCT, "ESTIMATE. A big card or a strong pre-sale."),
    ("comp", "People who didn't buy a ticket: % of the crowd", 0.20, PCT,
     "From the first show: 400 tickets sold, about 500 people in the building, so about 1 in 5 came free "
     "(sponsor guests, fighters' teams, volunteers, officials). Tune it after each show."),
    ("minc", "Smallest crowd worth doing a show for (people)", 1000, NUM,
     "Cody, Oct 2026: 1,000+. A venue whose expected crowd is below this gets flagged on Compare."),
]
r = 5
for key, lab, val, fmt, nt in rows:
    wsI.cell(r, 1, lab).font = f_lab; wsI.cell(r, 1).alignment = Alignment(vertical="top")
    inp(wsI.cell(r, 2), val, fmt); note(wsI, r, NC, nt); g[key] = f"Inputs!$B${r}"; r += 1

# ticket tiers
r += 1
for c, h in zip(range(1, 4), ("Ticket tier", "Price per person ($)", "Share of paid tickets")):
    x = wsI.cell(r, c, h); x.font = f_bold; x.fill = fill_goldl; x.border = box
note(wsI, r, NC, "Prices are what CCCL charged in May 2026 (Cody, Oct 2026). A VIP table is $1,500 and seats 8-10 depending on the venue's tables; 8 seats = $187.50 a seat. "
                 "Shares: 48 of the 400 tickets were table seats (12%); the GA / VIP split is an ESTIMATE. Shares must add to 100%.")
tier_start = r + 1
for name, price, share in (("General admission", 60, 0.78), ("VIP single seat", 100, 0.10),
                           ("VIP table seat ($1,500 table of 8, per seat)", 187.5, 0.12)):
    r += 1
    wsI.cell(r, 1, name).font = f_lab
    inp(wsI.cell(r, 2), price, USD); inp(wsI.cell(r, 3), share, PCT)
tier_end = r
r += 1
wsI.cell(r, 1, "Shares add up to (must be 100%)").font = f_lab
x = wsI.cell(r, 3, f"=SUM(C{tier_start}:C{tier_end})"); x.font = f_calc; x.number_format = PCT; x.border = box
share_chk = f"C{r}"
wsI.conditional_formatting.add(share_chk, CellIsRule(operator="notEqual", formula=["1"],
                               fill=PatternFill("solid", fgColor="F8CBAD")))
r += 1
wsI.cell(r, 1, "Average price per paid ticket").font = f_bold
x = wsI.cell(r, 2, f"=IF(SUM(C{tier_start}:C{tier_end})=0,0,SUMPRODUCT(B{tier_start}:B{tier_end},C{tier_start}:C{tier_end})/SUM(C{tier_start}:C{tier_end}))")
x.font = f_calc; x.number_format = '$#,##0.00'; x.border = box
g["avg"] = f"Inputs!$B${r}"
r += 2
rows = [
    ("fee", "Ticketing fees CCCL pays (% of ticket sales)", 0.079, '0.0%',
     "CCCL paid the Eventbrite fees in 2026 (Cody). Eventbrite Canada (help centre, Oct 2026): 3.5% + $1.29 per ticket, plus 2.9% processing. "
     "That's 8.6% on a $60 ticket, 7.7% on $100 and 6.5% on a $1,500 table, about 7.9% across the 2026 mix. Passing fees to buyers would make this 0%."),
    ("tax", "HST inside the ticket price that goes to the government (%)", 0, '0.0%',
     "0% unless CCCL is registered for HST and the ticket prices included tax. Then use 11.5% (the 13% HST inside a "
     "tax-included price). Ask the bookkeeper."),
    ("bar", "Bar & concession profit per person in the building ($)", 0, '$#,##0.00',
     "0: in 2026 the arena or a local group ran the bar and kept it (Cody). If CCCL ever runs or shares the bar, put the profit per person here."),
    ("merch", "Merch profit per person ($)", 0, '$#,##0.00', "0 until CCCL sells shirts or hats at shows. $1-3 a person is typical once it does (ESTIMATE)."),
    ("stream", "Livestream / pay-per-view profit per event ($)", 0, USD,
     "0: the 2026 livestream didn't bring in money (Cody). The Marketing Kit's Livestream tab works out what a paid stream could bring in."),
    ("cont", "Contingency buffer (% of fixed costs)", 0.05, PCT,
     "Cushion for surprises. 5-10% is normal for a first or second show."),
    ("sdrop", "Stress test: crowd drop vs Expected", 0.30, PCT, "What if 30% fewer people show up?"),
    ("spdrop", "Stress test: sponsorship that falls through", 0.50, PCT, "What if half the sponsor money never comes?"),
]
for key, lab, val, fmt, nt in rows:
    wsI.cell(r, 1, lab).font = f_lab; wsI.cell(r, 1).alignment = Alignment(vertical="top")
    inp(wsI.cell(r, 2), val, fmt); note(wsI, r, NC, nt); g[key] = f"Inputs!$B${r}"; r += 1

# venues
r += 1
section(wsI, r, "VENUE OPTIONS  (one column per option)", W)
r += 1
BIG = dict(security=3500, production=3500, seating=2000, athletes=2500, hosp=1200)   # city arenas: ESTIMATES
VEN = [
    dict(name="A.J. LaRue Arena", town="Haliburton", contact="Dysart et al arena 705-457-1740 x635",
         date="May 2027 (2nd annual)", cap=1000, rate=1500, days=1, hrs=0, staff=30,
         spons=18000, mkt=2000),
    dict(name="North Hastings Community Centre", town="Bancroft", contact="Town of Bancroft 613-332-2062, communitycentre@bancroft.ca",
         date="June 2027 (dry floor; the Gemboree fills it on the Civic Holiday weekend)", cap=800, rate=1000, days=2, hrs=16, staff=30,
         spons=12000, mkt=1500),
    dict(name="Dr. Hamilton Crawford Memorial Arena", town="Marmora", contact="Marmora & Lake Parks & Rec 613-472-2410, arena@marmoraandlake.ca",
         date="Aug 2027 (dry floor May to Aug)", cap=1000, rate=600, days=2, hrs=16, staff=30,
         spons=6000, mkt=2500),
    dict(name="Peterborough Memorial Centre", town="Peterborough", contact="Jose Garcia-Perez, Event Services, 705-743-3561 x2504",
         date="Feb or Nov 2027, around Petes home games (holds winter floor events)", cap=2000, rate=6000, days=1, hrs=60, staff=40,
         spons=10000, mkt=5000, **BIG),
    dict(name="Slush Puppie Place", town="Kingston", contact="ASM Global 613-650-5000 ('Private Venue Bookings' on slushpuppieplace.com)",
         date="Nov 2027, around Frontenacs home games (held BTC MMA in March 2023 and 2024)", cap=2000, rate=8000, days=1, hrs=80, staff=40,
         spons=6000, mkt=6000, **BIG),
    dict(name="Tribute Communities Centre", town="Oshawa", contact="Oak View Group 905-438-8881 x116 ('Request a Proposal' online)",
         date="Winter 2027 (renovation starts in 2027; PFC plans Oshawa in Nov)", cap=2000, rate=8000, days=1, hrs=80, staff=40,
         spons=6000, mkt=6000, **BIG),
]
# Cody's real 2026 costs (Oct 2026): cage 3,000, officials 3,000, medical 2,000, insurance 2,000, security 1,500,
# lights/sound/stream 1,000, everything else 15,000 (split below: permit, marketing, athletes, seating, food, other).
COMMON = dict(permit=600, cage=3000, officials=3000, medical=2000, insurance=2000, security=1500,
              production=1000, stream=0, athletes=1500, seating=1000, hosp=600, other=9300)
v = {}  # row numbers
spec = [
    ("name", "Venue (building)", None, "text"),
    ("town", "Town", None, "text"),
    ("contact", "Who to call for a quote", None, "text"),
    ("date", "Date idea", None, "text"),
    ("cap", "Sellable capacity (people)", NUM, "num"),
    ("rate", "Venue rental per day ($)", USD, "num"),
    ("days", "Days booked (setup + show)", NUM, "num"),
    ("hrs", "Venue staff hours (setup + teardown)", NUM, "num"),
    ("staff", "Venue staff rate per hour ($)", USD, "num"),
    ("permit", "Ontario Athletics Commission event permit ($)", USD, "num"),
    ("cage", "Cage rental & delivery ($)", USD, "num"),
    ("officials", "Officials: referees, judges, timekeeper ($)", USD, "num"),
    ("medical", "Medical: ringside doctor, paramedics ($)", USD, "num"),
    ("insurance", "Event insurance ($)", USD, "num"),
    ("security", "Security ($)", USD, "num"),
    ("production", "Lighting, sound & staging ($)", USD, "num"),
    ("stream", "Livestream production ($)", USD, "num"),
    ("mkt", "Marketing, posters & ads ($)", USD, "num"),
    ("athletes", "Athlete costs: medals, belts, travel, officials' hotel ($)", USD, "num"),
    ("seating", "Chairs, tables & VIP setup ($)", USD, "num"),
    ("hosp", "Crew & volunteer food ($)", USD, "num"),
    ("other", "Anything else ($)", USD, "num"),
    ("spons", "Sponsorship expected: signed AND paid ($)", USD, "num"),
]
NOTES = {
    "name": "", "town": "", "contact": "Phone numbers from each municipality's website (Oct 2026).",
    "date": "Cody's months for 2027: Jan, Feb, Apr, May, Jun, Aug, Nov. Arenas have ice in winter, so Jan, Feb and Nov need a hall, event centre or casino.",
    "cap": "Set it to about twice the crowd you expect (the Expected crowd % is 50%). Haliburton ~1,000 and Bancroft 800 are ESTIMATES (not published). "
           "Marmora 1,000 is published. City arenas are set to 2,000 (ESTIMATE): Peterborough's published setups run 1,722 to 5,501, "
           "Kingston lists 5,700 for cold-season events, Oshawa 1,000 to 7,600. Cody's minimum crowd is 1,000.",
    "rate": "Haliburton: $1,500 was the real 2026 arena bill (Cody), entered as one day. Marmora $600/day is published. Bancroft is an ESTIMATE. "
            "Peterborough, Kingston and Oshawa don't publish rates: ESTIMATES until they quote (rent, staff, ushers, box office). "
            "For scale, Peterborough's smaller Healthy Planet Arena charges $3,700.75 a day for a floor pad.",
    "days": "Setup day + show day, unless the venue quotes the whole event.",
    "hrs": "ESTIMATE. 0 where the quote already includes staff. City arenas need ushers, box office and a bigger setup crew.",
    "staff": "Haliburton lists $30/hr for setup and teardown with staff. Others are ESTIMATES.",
    "permit": "Ontario Athletics Commission: $600 permit per show. Under the regulation a 'contest' is one bout or several in a row, so one $600 permit "
              "should cover the whole card (confirm with the Commission). The $3,000 permit covers up to 8 shows in a year. Assumed to be inside Cody's $15,000 'everything else'.",
    "cage": "Cage, officials, medical, insurance, security and lights/sound/stream are Cody's real 2026 costs. For the three city arenas, security, "
            "lights/sound, seating, athlete costs and crew food are raised (ESTIMATES) for a 2,000-seat room and the longer drive. The cage must sit on a platform 3-4 ft high (Ontario rule): make sure the quote includes it.",
    "stream": "0: in 2026 the stream was part of the $1,000 lights/sound/stream bill.",
    "mkt": "ESTIMATE, part of Cody's $15,000 'everything else'. Bigger towns need more.",
    "athletes": "ESTIMATE, part of Cody's $15,000 'everything else'.",
    "other": "The rest of Cody's $15,000 'everything else' (it's $15,000 minus the permit, marketing, athletes, seating and food lines). "
             "Replace with the real invoices when Cody finds them.",
    "spons": "Haliburton 2026: 15 sponsors, $18,000 (Cody). With 4+ shows in 2027, the same local sponsors can't give that to every show, so count only "
             "what's signed AND paid for THIS show. Other venues: ESTIMATES (Bancroft is home; the cities are new markets).",
}
for key, lab, fmt, kind in spec:
    wsI.cell(r, 1, lab).font = f_lab
    wsI.cell(r, 1).alignment = Alignment(vertical="top", wrap_text=True)
    for i, ven in enumerate(VEN):
        val = ven.get(key, COMMON.get(key))
        if key in COMMON and ven["cap"] is None:
            val = None
        c = wsI.cell(r, 2 + i)
        if kind == "text":
            inp(c, val if val else None, None, key=(key == "name"))
            c.alignment = Alignment(wrap_text=True, vertical="top")
            if key == "name":
                c.font = Font(name=F, size=10, bold=True, color="0000FF")
        else:
            inp(c, val, fmt)
    if NOTES.get(key): note(wsI, r, NC, NOTES[key])
    v[key] = r
    r += 1
import math
for k in v:
    lines = math.ceil(len(NOTES.get(k, "")) / 88)
    if k in ("contact", "date", "name"): lines = max(lines, 3)
    if lines > 1: wsI.row_dimensions[v[k]].height = 13 * lines + 2
wsI.freeze_panes = "B4"

# ------------------------------------------------------------------ Compare
wsC = wb.create_sheet("Compare")
W = NV + 1
LASTC = L(NV + 1)
banner(wsC, "CCCL EVENT PROFIT PREDICTOR  ·  COMPARE VENUES",
       "Nothing to type here. Every number comes from the Inputs tab. Green = pulled from Inputs, black = calculated.", W)
wsC.column_dimensions["A"].width = 52
cols = [L(2 + i) for i in range(NV)]
for c in cols:
    wsC.column_dimensions[c].width = 21
R = {}
r = 4

def row(key, label, formula_fn, fmt=USD, font=f_calc, bold=False, fill=None):
    global r
    wsC.cell(r, 1, label).font = f_bold if bold else f_lab
    for i, col in enumerate(cols):
        ic = L(2 + i)  # matching Inputs column
        c = wsC.cell(r, 2 + i, formula_fn(col, ic))
        c.font = Font(name=F, size=10, bold=bold, color=font.color)
        c.number_format = fmt; c.border = box
        if fill: c.fill = fill
    R[key] = r; r += 1

def has(col):  # guard: venue column filled in?
    return f"{col}${R['cap']}>0"

# header row with venue names
wsC.cell(r, 1, "Venue").font = f_bold
for i, col in enumerate(cols):
    ic = L(2 + i)
    c = wsC.cell(r, 2 + i, f'=IF(Inputs!{ic}{v["name"]}="","(add another city)",Inputs!{ic}{v["name"]}&" · "&Inputs!{ic}{v["town"]})')
    c.font = Font(name=F, size=10, bold=True, color="FFFFFF"); c.fill = fill_navy; c.alignment = center
wsC.row_dimensions[r].height = 48
R["hdr"] = r; r += 1
row("cap", "Sellable capacity (people)", lambda c, ic: f"=Inputs!{ic}{v['cap']}", NUM, f_link)

r += 1; section(wsC, r, "COSTS THAT DON'T DEPEND ON THE CROWD", W); r += 1
row("venue", "Venue rental (rate × days)", lambda c, ic: f"=Inputs!{ic}{v['rate']}*Inputs!{ic}{v['days']}")
row("vstaff", "Venue staff (hours × rate)", lambda c, ic: f"=Inputs!{ic}{v['hrs']}*Inputs!{ic}{v['staff']}")
row("oth", "All other costs (permit, cage, officials, medical, etc.)",
    lambda c, ic: f"=SUM(Inputs!{ic}{v['permit']}:Inputs!{ic}{v['other']})")
row("cont", "Contingency buffer", lambda c, ic: f"=({c}{R['venue']}+{c}{R['vstaff']}+{c}{R['oth']})*{g['cont']}")
row("fixed", "TOTAL EVENT COST", lambda c, ic: f"=SUM({c}{R['venue']}:{c}{R['cont']})", USD, f_calc, True, fill_goldl)

r += 1; section(wsC, r, "WHAT EACH PAID TICKET BRINGS IN", W); r += 1
row("net", "Ticket money after fees and HST (per paid ticket)", lambda c, ic: f"=IF({has(c)},{g['avg']}*(1-{g['fee']}-{g['tax']}),0)", '$#,##0.00;($#,##0.00);"-"')
row("extra", "Bar + merch (per paid ticket, comps included)",
    lambda c, ic: f"=IF({has(c)},({g['bar']}+{g['merch']})/(1-{g['comp']}),0)", '$#,##0.00;($#,##0.00);"-"')
row("contrib", "Total per paid ticket", lambda c, ic: f"={c}{R['net']}+{c}{R['extra']}", '$#,##0.00', f_calc, True)

r += 1; section(wsC, r, "BREAK-EVEN: WHAT IT TAKES NOT TO LOSE MONEY", W); r += 1
row("be_t", "Paid tickets needed (with expected sponsors)",
    lambda c, ic: f"=IF(AND({has(c)},{c}{R['contrib']}>0),MAX(0,ROUNDUP(({c}{R['fixed']}-Inputs!{ic}{v['spons']}-{g['stream']})/{c}{R['contrib']},0)),0)",
    NUM, f_calc, True, fill_goldl)
row("be_pct", "... as % of capacity",
    lambda c, ic: f"=IF({has(c)},{c}{R['be_t']}/(1-{g['comp']})/{c}{R['cap']},0)", PCT)
row("be_t0", "Paid tickets needed with NO sponsors",
    lambda c, ic: f"=IF(AND({has(c)},{c}{R['contrib']}>0),MAX(0,ROUNDUP(({c}{R['fixed']}-{g['stream']})/{c}{R['contrib']},0)),0)", NUM)
row("be_sp", "Sponsorship needed to break even on a LOW night",
    lambda c, ic: f"=IF({has(c)},MAX(0,{c}{R['fixed']}-{g['stream']}-ROUND({c}{R['cap']}*{g['low']}*(1-{g['comp']}),0)*{c}{R['contrib']}),0)",
    USD, f_calc, True)
row("sp_have", "Sponsorship you expect (from Inputs)", lambda c, ic: f"=Inputs!{ic}{v['spons']}", USD, f_link)

def scenario(tag, label, pct_ref, spons_mult="1"):
    global r
    r += 1; section(wsC, r, label, W); r += 1
    row(f"{tag}_att", "Crowd (people in the building)",
        lambda c, ic: f"=ROUND({c}{R['cap']}*{pct_ref},0)", NUM)
    row(f"{tag}_paid", "Paid tickets", lambda c, ic: f"=ROUND({c}{R[tag + '_att']}*(1-{g['comp']}),0)", NUM)
    row(f"{tag}_tix", "Ticket money after fees", lambda c, ic: f"={c}{R[tag + '_paid']}*{c}{R['net']}")
    row(f"{tag}_bar", "Bar + merch", lambda c, ic: f"={c}{R[tag + '_att']}*({g['bar']}+{g['merch']})")
    row(f"{tag}_sp", "Sponsorship + livestream",
        lambda c, ic: f"=IF({has(c)},Inputs!{ic}{v['spons']}*{spons_mult}+{g['stream']},0)")
    row(f"{tag}_rev", "Total money in", lambda c, ic: f"=SUM({c}{R[tag + '_tix']}:{c}{R[tag + '_sp']})", USD, f_calc, True)
    row(f"{tag}_cost", "Total event cost", lambda c, ic: f"={c}{R['fixed']}")
    row(f"{tag}_pl", "PROFIT / (LOSS)", lambda c, ic: f"={c}{R[tag + '_rev']}-{c}{R[tag + '_cost']}",
        USD, f_calc, True, fill_goldl)

scenario("lo", "SCENARIO: LOW CROWD", g["low"])
scenario("ex", "SCENARIO: EXPECTED CROWD", g["exp"])
scenario("hi", "SCENARIO: HIGH CROWD", g["high"])
scenario("st", "STRESS TEST: expected crowd drops AND sponsors pull out", f"{g['exp']}*(1-{g['sdrop']})",
         f"(1-{g['spdrop']})")

r += 1; section(wsC, r, "BOTTOM LINE", W); r += 1
row("verdict", "Verdict",
    lambda c, ic: (f'=IF(NOT({has(c)}),"",IF({c}{R["lo_pl"]}>=0,"SAFE: profits even on a low night",'
                   f'IF({c}{R["ex_pl"]}>=0,"OK: profits if the expected crowd shows","RISKY: loses money at the expected crowd")))'
                   f'&IF(AND({has(c)},{c}{R["ex_att"]}<{g["minc"]}),"  ·  BELOW MINIMUM: expected crowd is under Cody\'s minimum","")'),
    "@", f_calc, True)
wsC.row_dimensions[R["verdict"]].height = 58
for i in range(NV):
    wsC.cell(R["verdict"], 2 + i).alignment = center
wsC.cell(r, 1, "Best option at the expected crowd").font = f_bold
pl = f"B{R['ex_pl']}:{LASTC}{R['ex_pl']}"
c = wsC.cell(r, 2, f'=IF(MAX({pl})=0,"",INDEX(B{R["hdr"]}:{LASTC}{R["hdr"]},MATCH(MAX({pl}),{pl},0)))')
c.font = Font(name=F, size=11, bold=True, color=NAVY); c.fill = fill_gold
wsC.merge_cells(start_row=r, start_column=2, end_row=r, end_column=NV + 1)
R["best"] = r; r += 2
note(wsC, r, 1, "A budget tool, not a promise. Every ESTIMATE on the Inputs tab needs a real quote before money is committed. "
                "Sponsorship only counts once it's signed and paid.")
wsC.merge_cells(start_row=r, start_column=1, end_row=r, end_column=NV + 1)
wsC.row_dimensions[r].height = 30

# colour profit cells
red = Font(name=F, size=10, bold=True, color="C00000")
for k in ("lo_pl", "ex_pl", "hi_pl", "st_pl"):
    rng = f"B{R[k]}:{LASTC}{R[k]}"
    wsC.conditional_formatting.add(rng, CellIsRule(operator="lessThan", formula=["0"], font=red,
                                   fill=PatternFill("solid", fgColor="F8CBAD")))
    wsC.conditional_formatting.add(rng, CellIsRule(operator="greaterThan", formula=["0"],
                                   fill=PatternFill("solid", fgColor="C6EFCE")))
vr = f"B{R['verdict']}:{LASTC}{R['verdict']}"
wsC.conditional_formatting.add(vr, FormulaRule(formula=[f'LEFT(B{R["verdict"]},4)="SAFE"'], fill=PatternFill("solid", fgColor="C6EFCE")))
wsC.conditional_formatting.add(vr, FormulaRule(formula=[f'LEFT(B{R["verdict"]},2)="OK"'], fill=PatternFill("solid", fgColor="FFEB9C")))
wsC.conditional_formatting.add(vr, FormulaRule(formula=[f'LEFT(B{R["verdict"]},5)="RISKY"'], fill=PatternFill("solid", fgColor="F8CBAD")))
wsC.freeze_panes = "B5"

# ------------------------------------------------------------------ Haliburton check
wsH = wb.create_sheet("Haliburton 2026 Check")
W = 5
banner(wsH, "HALIBURTON MAY 2026: WHAT REALLY HAPPENED",
       "Part 1 adds up Cody's numbers (Oct 2026): about $18,000 profit. That profit went into the gym and other projects (Kyle, Oct 2026), "
       "which is why it felt like breaking even. Part 2 records where the money went: type the real amounts once Cody has the Eventbrite payout report. Part 3 checks the model against the real result.", W)
for c, w in zip("ABCDE", (50, 17, 17, 17, 62)):
    wsH.column_dimensions[c].width = w
H = {}
def hrow(key, label, value, fmt=USD, kind="in", tip="", bold=False, fill=None):
    """kind: in = yellow input, calc = black formula, link = green formula"""
    global r
    wsH.cell(r, 1, label).font = f_bold if bold else f_lab
    wsH.cell(r, 1).alignment = Alignment(vertical="top", wrap_text=True)
    c = wsH.cell(r, 2)
    if kind == "in":
        inp(c, value, fmt)
    else:
        c.value = value; c.number_format = fmt; c.border = box
        c.font = Font(name=F, size=10, bold=bold, color=(f_link if kind == "link" else f_calc).color)
        if fill: c.fill = fill
    if tip:
        note(wsH, r, 3, tip); wsH.merge_cells(start_row=r, start_column=3, end_row=r, end_column=5)
        if len(tip) > 95: wsH.row_dimensions[r].height = 15 * (1 + len(tip) // 95)
    H[key] = f"B{r}"; r += 1

r = 4
section(wsH, r, "PART 1: THE NUMBERS CODY GAVE (Oct 2026)", W); r += 1
hrow("ga_p", "General admission price", 60, USD, tip="From Cody.")
hrow("vip_p", "VIP single seat price", 100, USD, tip="From Cody.")
hrow("tab_p", "VIP table price", 1500, USD, tip="From Cody. Tables seat 8-10 depending on the venue's tables (Kyle).")
hrow("tab_s", "Seats per table", 8, NUM)
hrow("tab_n", "VIP tables sold", 6, NUM)
hrow("sold", "Tickets sold, all types", 400, NUM, tip="From Cody. Assumed to include the 48 table seats.")
hrow("vip_n", "VIP single seats sold", 40, NUM, tip="ESTIMATE. Cody didn't split GA and VIP. Fix it from the Eventbrite sales report.")
hrow("tab_seats", "Table seats (tables × seats)", f"={H['tab_n']}*{H['tab_s']}", NUM, "calc")
hrow("ga_n", "General admission tickets (what's left)", f"=MAX(0,{H['sold']}-{H['tab_seats']}-{H['vip_n']})", NUM, "calc")
hrow("gross", "Ticket money before fees", f"={H['ga_n']}*{H['ga_p']}+{H['vip_n']}*{H['vip_p']}+{H['tab_n']}*{H['tab_p']}", USD, "calc", bold=True)
hrow("fee", "Eventbrite fees CCCL paid (% from Inputs)", f"={g['fee']}", '0.0%', "link")
hrow("tix", "Ticket money after fees", f"={H['gross']}*(1-{H['fee']})", USD, "calc")
hrow("spons", "Sponsor money (15 sponsors)", 18000, USD, tip="From Cody. Title sponsor: Dillabough Millwrighting & Welding.")
hrow("bar", "Bar profit to CCCL", 0, USD, tip="The arena or a local group ran the bar (Cody).")
hrow("money_in", "TOTAL MONEY IN", f"={H['tix']}+{H['spons']}+{H['bar']}", USD, "calc", bold=True, fill=fill_goldl)
r += 1
cost_first = r
for k, lab, val in (("c_arena", "Arena rental", 1500), ("c_cage", "Cage", 3000), ("c_med", "Medical", 2000),
                    ("c_off", "Officials", 3000), ("c_ins", "Insurance", 2000), ("c_sec", "Security", 1500),
                    ("c_prod", "Lights, sound, stream", 1000), ("c_else", "Everything else", 15000)):
    hrow(k, lab, val, USD, tip=("From Cody. What's in this $15,000 is the biggest unknown: the Commission permit, ads, "
                                 "belts and medals, hotels, food, seating?") if k == "c_else" else "")
cost_last = r - 1
hrow("cost", "TOTAL COST", f"=SUM(B{cost_first}:B{cost_last})", USD, "calc", bold=True, fill=fill_goldl)
r += 1
hrow("implied", "PROFIT THESE NUMBERS ADD UP TO", f"={H['money_in']}-{H['cost']}", USD, "calc", bold=True, fill=fill_gold)
hrow("said", "Profit left over that Cody remembers ($)", 0, USD,
     tip="Cody said 'broke even, slight profit' because the show's profit went into the gym and other projects (Kyle, Oct 2026).")
hrow("gap", "MONEY TO ACCOUNT FOR", f"={H['implied']}-{H['said']}", USD, "calc", bold=True, fill=fill_gold)
r += 1
section(wsH, r, "PART 2: WHERE THE MONEY WENT  (type what applies, leave the rest at 0)", W); r += 1
exp_first = r
hrow("x_gym", "Show profit put into the gym and other projects ($)", 0, USD,
     tip="Left at 0 until Cody has the real amount (Eventbrite payout report and bank records). Kyle's understanding (Oct 2026): the show made money and it went into the gym and other projects. "
         "This doesn't lower the show's profit; it's where the profit went.")
wsH.cell(r, 1, "Anything below made the show's real profit smaller:").font = f_note; r += 1
hrow("x_tables", "VIP table money already counted inside the $18,000 sponsor total ($)", 0, USD,
     tip="If tables were sold as part of sponsor packages, up to $9,000 got counted twice.")
hrow("x_free_n", "Free tickets counted inside the 400 (how many)", 0, NUM,
     tip="Sponsor tickets, fighters' families, volunteers. Valued at the GA price below.")
hrow("x_free", "... worth (free tickets × GA price)", f"={H['x_free_n']}*{H['ga_p']}", USD, "calc")
hrow("x_unpaid", "Sponsor money promised but never paid, or given as goods/services instead of cash ($)", 0, USD)
hrow("x_hst", "HST CCCL sent to the government ($)", 0, USD,
     tip="Only if CCCL is HST-registered. On tax-included prices it's about 11.5% of ticket and sponsor money.")
hrow("x_refund", "Refunds and chargebacks ($)", 0, USD)
hrow("x_cost", "Bills paid that aren't in the $29,000 ($)", 0, USD,
     tip="Fighter travel, hotels, posters and ads, food, gear CCCL bought, the Commission permit if it wasn't in 'everything else'.")
hrow("explained", "TOTAL ACCOUNTED FOR", f"={H['x_gym']}+{H['x_tables']}+{H['x_free']}+{H['x_unpaid']}+{H['x_hst']}+{H['x_refund']}+{H['x_cost']}",
     USD, "calc", bold=True, fill=fill_goldl)
hrow("unexpl", "STILL UNEXPLAINED", f"={H['gap']}-{H['explained']}", USD, "calc", bold=True, fill=fill_gold,
     tip="Goal: under $1,000 either way.")
wsH.conditional_formatting.add(H["unexpl"], FormulaRule(formula=[f"ABS({H['unexpl']})<=1000"], fill=PatternFill("solid", fgColor="C6EFCE")))
wsH.conditional_formatting.add(H["unexpl"], FormulaRule(formula=[f"ABS({H['unexpl']})>1000"], fill=PatternFill("solid", fgColor="F8CBAD")))
r += 1
section(wsH, r, "PART 3: REAL RESULT vs THE MODEL", W); r += 1
for i, h in enumerate(("", "Real (May 2026)", "Model (Expected)", "Difference", "If it's off, tune this on Inputs")):
    x = wsH.cell(r, 1 + i, h); x.font = f_bold; x.fill = fill_goldl; x.border = box; x.alignment = center
r += 1
checks = [
    ("People in the building", NUM, 500, f"=Compare!B{R['ex_att']}", "Expected crowd %. Cody said 400-600 people; 500 is the middle."),
    ("Paid tickets", NUM, f"={H['sold']}-{H['x_free_n']}", f"=Compare!B{R['ex_paid']}", "'People who didn't buy a ticket' %"),
    ("Ticket money (after fees)", USD, f"={H['tix']}-{H['x_free']}-{H['x_refund']}", f"=Compare!B{R['ex_tix']}", "Ticket prices, tier shares, fee % and HST %"),
    ("Bar + merch profit", USD, f"={H['bar']}", f"=Compare!B{R['ex_bar']}", "Bar and merch profit per person"),
    ("Sponsorship received", USD, f"={H['spons']}-{H['x_tables']}-{H['x_unpaid']}", f"=Inputs!B{v['spons']}", "Sponsorship expected (Haliburton column)"),
    ("Total cost", USD, f"={H['cost']}+{H['x_cost']}", f"=Compare!B{R['fixed']}", "The cost lines. The model adds a contingency buffer on top."),
    ("Profit / (loss)", USD, f"={H['implied']}-({H['explained']}-{H['x_gym']})", f"=Compare!B{R['ex_pl']}", "Real show profit, before the money that went to the gym."),
]
for lab, fmt, real, model, tip in checks:
    wsH.cell(r, 1, lab).font = f_lab
    c2 = wsH.cell(r, 2)
    if isinstance(real, str):
        c2.value = real; c2.font = f_calc; c2.number_format = fmt; c2.border = box
    else:
        inp(c2, real, fmt)
    m = wsH.cell(r, 3, model); m.font = f_link; m.number_format = fmt; m.border = box
    d = wsH.cell(r, 4, f'=IF(B{r}="","",C{r}-B{r})'); d.font = f_calc; d.number_format = fmt; d.border = box
    note(wsH, r, 5, tip)
    r += 1
wsH.freeze_panes = "A4"

# ------------------------------------------------------------------ Town Scout
if TOWNS:
    import json as _json
    from openpyxl.formatting.rule import ColorScaleRule
    td = _json.load(open(TOWNS))["towns"]
    td = sorted(td, key=lambda t: -t["scores"]["overall"])
    wsT = wb.create_sheet("Town Scout")
    TC = [("#", 4), ("Town", 15), ("Overall", 9), ("Crowd", 8), ("Venue", 8), ("Fighters", 10), ("Competition", 13), ("Distance", 10),
          ("In the Compare tab?", 12), ("Drive from Bancroft", 18), ("Population", 20), ("Best venue", 26), ("Capacity", 18),
          ("Rental rate (as published)", 40), ("Dry floor season", 26), ("Combat gyms nearby", 30), ("Combat shows nearby", 30),
          ("Verdict", 46), ("Venue source", 34)]
    banner(wsT, "TOWN SCOUT: WHERE TO RUN CCCL IN 2027",
           "25 towns within about 2 hours of Bancroft, scored 1-5 (5 = best). Distance: 5 = 30 min or less, 1 = over 2 h. "
           "Overall is the simple average. Researched Oct 2026; drive times are estimates. Confirm rates and capacity with each venue.", len(TC))
    for i, (h, w) in enumerate(TC, 1):
        c = wsT.cell(4, i, h); c.font = f_head; c.fill = fill_navy; c.border = box
        c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center" if i <= 9 else "left")
        wsT.column_dimensions[L(i)].width = w
    wsT.row_dimensions[4].height = 32
    INC = {v["town"] for v in VEN if v.get("cap")}
    def first_line(x, n=160):
        x = (x or "").strip().replace("\n", " ")
        return x if len(x) <= n else x[:n - 1].rsplit(" ", 1)[0] + "…"
    for k, t in enumerate(td, 1):
        rr = 4 + k
        v0 = t["venues"][0] if t["venues"] else {}
        sc = t["scores"]
        town = t["town"].replace(" (added)", "")
        vals = [k, town, sc["overall"], sc["crowd"], sc["venue"], sc["fighters"], sc["competition"], sc["distance"],
                "Yes" if any(town.startswith(x) for x in INC) else "",
                first_line(t["drive_from_bancroft"], 60), first_line(t["population"], 60), first_line(v0.get("name", ""), 80),
                first_line(v0.get("capacity", ""), 60), first_line(v0.get("rental_rate", ""), 220),
                first_line(v0.get("dry_floor_season", ""), 140),
                first_line("; ".join(g["name"].split(" (")[0] for g in t.get("gyms_nearby", [])) or "None found", 160),
                first_line("; ".join(f'{e["name"].split(" (")[0]} ({e["date"]})' for e in t.get("combat_events", [])
                                     if e.get("date")) or "None found", 160),
                first_line(t["verdict"], 240), v0.get("source_url", "")]
        for i, val in enumerate(vals, 1):
            c = wsT.cell(rr, i, val); c.font = f_bold if i in (2, 3) else f_lab; c.border = box
            c.alignment = Alignment(wrap_text=True, vertical="top", horizontal="center" if i <= 9 and i != 2 else "left")
            if i == 3: c.number_format = "0.0"
        wsT.row_dimensions[rr].height = 64
    last = 4 + len(td)
    wsT.conditional_formatting.add(f"C5:H{last}", ColorScaleRule(start_type="num", start_value=1, start_color="F8CBAD",
                                   mid_type="num", mid_value=3, mid_color="FFF2CC", end_type="num", end_value=5, end_color="A9D08E"))
    wsT.freeze_panes = "C5"
    wsT.auto_filter.ref = f"A4:{L(len(TC))}{last}"

# ------------------------------------------------------------------ 2027 Venues
if VENUES27:
    import json as _json
    vd = _json.load(open(VENUES27))
    wsV = wb.create_sheet("2027 Venues")
    VC = [("#", 4), ("Town", 14), ("Drive from Bancroft", 16), ("Venue", 28), ("Type", 11), ("Score", 7), ("Winter OK (Jan/Feb/Nov)?", 11),
          ("Event capacity (as published)", 40), ("Rental rate (as published)", 34), ("Combat shows held there", 34),
          ("Who to call", 34), ("Town verdict", 46), ("Clashes to avoid", 40), ("Source", 34)]
    banner(wsV, "2027 VENUES: WHERE A 1,000+ CROWD OR A WINTER SHOW COULD WORK",
           "Cody's 10 towns plus casinos and event centres up to 3 hours from Bancroft, researched October 2026. Score 1-5 (crowd, venue, "
           "fighters, competition, distance). Arenas have ice in winter, so Jan, Feb and Nov need a venue that holds floor events in hockey season.", len(VC))
    for i, (h, w) in enumerate(VC, 1):
        c = wsV.cell(4, i, h); c.font = f_head; c.fill = fill_navy; c.border = box
        c.alignment = Alignment(wrap_text=True, vertical="center")
        wsV.column_dimensions[L(i)].width = w
    wsV.row_dimensions[4].height = 32
    rowsV = []
    for t in vd["towns"]:
        for ve in t["venues"]:
            rowsV.append((t, ve))
    rowsV.sort(key=lambda x: (x[0].get("over_3h", False), -x[1]["scores"]["overall"]))
    def short(x, n):
        x = (x or "").replace("\n", " ").strip()
        return x if len(x) <= n else x[:n - 1].rsplit(" ", 1)[0] + "…"
    win = {"yes": "Yes", "no": "No (ice)", "unclear": "Unclear"}
    for k, (t, ve) in enumerate(rowsV, 1):
        rr = 4 + k
        clashes = "; ".join(f'{c["name"]} ({c["typical_dates"].split(" (from")[0]})' for c in t.get("clashes", []))
        drive = t["drive_from_bancroft"].replace(" (Rome2rio estimate)", "").replace("About ", "~")
        if t.get("over_3h"): drive = "OVER 3 H: " + drive.split(" - ")[0]
        vals = [k, t["town"], short(drive, 40), ve["name"], ve["type"].title(), ve["scores"]["overall"],
                win.get(ve["winter_ok"], ve["winter_ok"]), short(ve["capacity_event"], 200), short(ve["rental_rate"], 170),
                short(ve.get("hosted_combat", ""), 170), short(ve["booking_contact"], 150), short(t["verdict"], 260),
                short(clashes, 220), ve.get("source_url", "")]
        for i, val in enumerate(vals, 1):
            c = wsV.cell(rr, i, val); c.font = f_bold if i in (2, 4) else f_lab; c.border = box
            c.alignment = Alignment(wrap_text=True, vertical="top", horizontal="center" if i in (1, 5, 6, 7) else "left")
            if i == 6: c.number_format = "0.0"
            if t.get("over_3h"): c.fill = fill_grey
        wsV.row_dimensions[rr].height = 92
    lastV = 4 + len(rowsV)
    wsV.conditional_formatting.add(f"F5:F{lastV}", ColorScaleRule(start_type="num", start_value=2, start_color="F8CBAD",
                                   mid_type="num", mid_value=3.2, mid_color="FFF2CC", end_type="num", end_value=4.5, end_color="A9D08E"))
    wsV.conditional_formatting.add(f"G5:G{lastV}", CellIsRule(operator="equal", formula=['"Yes"'], fill=PatternFill("solid", fgColor="C6EFCE")))
    rr = lastV + 2
    x = wsV.cell(rr, 2, "OTHER PROMOTIONS RUNNING IN THESE MARKETS"); x.font = f_sec; rr += 1
    for i, h in enumerate(("Promotion", "", "Where", "", "How often"), 2):
        if h:
            c = wsV.cell(rr, i, h); c.font = f_head; c.fill = fill_navy
    rr += 1
    for pr in vd["promotions"]:
        if pr["name"].startswith("Cottage Country"): continue
        wsV.cell(rr, 2, short(pr["name"], 120)).font = f_bold
        wsV.merge_cells(start_row=rr, start_column=2, end_row=rr, end_column=3)
        wsV.cell(rr, 4, short(pr.get("markets", ""), 160)).font = f_lab
        wsV.merge_cells(start_row=rr, start_column=4, end_row=rr, end_column=5)
        wsV.cell(rr, 6, short(pr.get("frequency", ""), 220)).font = f_lab
        wsV.merge_cells(start_row=rr, start_column=6, end_row=rr, end_column=9)
        for cc in (2, 4, 6): wsV.cell(rr, cc).alignment = wrap
        wsV.row_dimensions[rr].height = 44
        rr += 1
    rr += 1
    note(wsV, rr, 2, "Lead time: the Ontario Athletics Commission needs the event permit at least 30 days ahead (120 days for rule changes), "
                     "so a January 2027 show is very tight and February only works if the venue is booked now. No big venue publishes its rental cost: "
                     "ask each for a quote. Drive times are estimates.")
    wsV.merge_cells(start_row=rr, start_column=2, end_row=rr, end_column=12); wsV.row_dimensions[rr].height = 30
    wsV.freeze_panes = "C5"
    wsV.auto_filter.ref = f"A4:{L(len(VC))}{lastV}"

# ------------------------------------------------------------------ 19+ Show Check
wsA = wb.create_sheet("19+ Show Check")
banner(wsA, "ADULTS-ONLY (19+) SHOW: IS IT WORTH IT?",
       "Compares each venue as a family show, a family show with a 19+ bar area, and an adults-only show. Adults-only means no youth "
       "grappling and no one under 19 in the building, but a full bar and room for an alcohol sponsor. Yellow cells are ESTIMATES: change them.", 9)
for c, w in zip("ABCDEFGHI", (40, 15, 15, 13, 15, 20, 18, 17, 4)):
    wsA.column_dimensions[c].width = w
r = 4
section(wsA, r, "ASSUMPTIONS (ESTIMATES: ask each venue who runs the bar and what CCCL gets)", 8); r += 1
A = {}
for key, lab, val, fmt, nt in (
    ("adult", "Adults (19+) at a family show: % of the crowd", 0.75, PCT, "ESTIMATE. The rest are kids, teens and youth grapplers."),
    ("bar_area", "Bar sales per adult: family show with a 19+ area ($)", 12, '$#,##0.00', "ESTIMATE. Fewer adults drink when the family is with them."),
    ("bar19", "Bar sales per adult: adults-only show ($)", 25, '$#,##0.00', "ESTIMATE."),
    ("share", "CCCL's share of bar sales", 0.50, PCT,
     "Mat McTrash's offer: 50% of total bar sales to CCCL (Cody, Oct 20, 2026). It only applies where Mat McTrash runs the bar; ask every other venue. "
     "A for-profit promoter usually can't hold the public-event liquor permit itself (AGCO: charities, non-profits, or events a council names municipally significant)."),
    ("chg", "Crowd change for an adults-only show", -0.20, '0%;-0%;0%',
     "ESTIMATE. Lose kids, teens and youth grapplers' families. Make it less negative if CCCL's crowd is mostly adults."),
    ("alc", "Alcohol brand sponsorship, adults-only show only ($)", 0, USD,
     "0 until signed. Proper No. Twelve at Gold would be $4,000, a long shot. Alcohol brands need a mostly-adult audience, so not at family shows."),
    ("x_area", "Extra cost: 19+ area at a family show ($)", 750, USD, "ESTIMATE. Fencing, wristbands, ID checks, an extra security guard."),
    ("x_19", "Extra cost: adults-only show ($)", 1000, USD, "ESTIMATE. ID checks at the door and extra security."),
):
    wsA.cell(r, 1, lab).font = f_lab; wsA.cell(r, 1).alignment = Alignment(vertical="top", wrap_text=True)
    inp(wsA.cell(r, 2), val, fmt)
    note(wsA, r, 3, nt); wsA.merge_cells(start_row=r, start_column=3, end_row=r, end_column=8)
    wsA.row_dimensions[r].height = 28 if len(nt) < 120 else 40
    A[key] = f"$B${r}"; r += 1
r += 1
section(wsA, r, "RESULT AT THE EXPECTED CROWD (from the Compare tab)", 8); r += 1
heads = ("Venue", "Family show profit", "Family + 19+ area profit", "Adults-only crowd", "Adults-only profit", "Best option",
         "Bar money CCCL needs per adult for adults-only to beat the best family option", "Biggest crowd drop adults-only can take (vs the best family option)")
for i, h in enumerate(heads, 1):
    x = wsA.cell(r, i, h); x.font = f_head; x.fill = fill_navy; x.border = box; x.alignment = Alignment(wrap_text=True, vertical="center")
wsA.row_dimensions[r].height = 60
r += 1
first_res = r
for i, c in enumerate(cols):
    att, tix, barm, pl, cap = (f"Compare!{c}{R[k]}" for k in ("ex_att", "ex_tix", "ex_bar", "ex_pl", "cap"))
    ok = f"{cap}>0"
    wsA.cell(r, 1, f'=IF({ok},Compare!{c}{R["hdr"]},"")').font = f_link
    fam = f"=IF({ok},{pl},\"\")"
    area = f"=IF({ok},{pl}+{att}*{A['adult']}*{A['bar_area']}*{A['share']}-{A['x_area']},\"\")"
    att19 = f"=IF({ok},ROUND({att}*(1+{A['chg']}),0),\"\")"
    d19 = f"=IF({ok},{pl}+({tix}+{barm})*{A['chg']}+D{r}*{A['bar19']}*{A['share']}+{A['alc']}-{A['x_19']},\"\")"
    best = f'=IF(NOT({ok}),"",IF(MAX(B{r}:C{r},E{r})=E{r},"Adults-only",IF(MAX(B{r}:C{r})=C{r},"Family + 19+ area","Family show")))'
    need = f"=IF(OR(NOT({ok}),D{r}=0),\"\",MAX(0,(MAX(B{r}:C{r})-B{r}+{A['x_19']}-{A['alc']}-({tix}+{barm})*{A['chg']})/D{r}))"
    num = f"(MAX(B{r}:C{r})-B{r}+{A['x_19']}-{A['alc']}-{att}*{A['bar19']}*{A['share']})"
    den = f"({tix}+{barm}+{att}*{A['bar19']}*{A['share']})"
    maxdrop = f'=IF(NOT({ok}),"",IF({num}>=0,"Can\'t match",-{num}/{den}))'
    for j, (val, fmt) in enumerate(((fam, USD), (area, USD), (att19, NUM), (d19, USD), (best, "@"), (need, '$#,##0.00'), (maxdrop, '0%')), 2):
        x = wsA.cell(r, j, val); x.font = f_calc; x.number_format = fmt; x.border = box
        x.alignment = Alignment(horizontal="center", vertical="center")
    wsA.cell(r, 1).border = box; wsA.cell(r, 1).alignment = Alignment(wrap_text=True, vertical="center")
    wsA.row_dimensions[r].height = 30
    r += 1
last_res = r - 1
for col in "BCE":
    rng = f"{col}{first_res}:{col}{last_res}"
    wsA.conditional_formatting.add(rng, CellIsRule(operator="lessThan", formula=["0"], font=red, fill=PatternFill("solid", fgColor="F8CBAD")))
wsA.conditional_formatting.add(f"F{first_res}:F{last_res}", FormulaRule(formula=[f'F{first_res}="Adults-only"'], fill=PatternFill("solid", fgColor="C6EFCE")))
r += 1
for t in ("How to read it: an adults-only show gives up every under-19 ticket (kids, teens, youth grapplers' families) to get a full bar and room for an "
          "alcohol sponsor. Each paid ticket brings in about $73 after fees, and CCCL only gets a slice of the bar, so a big crowd drop is hard to make up. "
          "The second-last column is the number to take to a venue: it's what CCCL's share of the bar must reach per adult for adults-only to beat the better of the two family options. If CCCL can't get that much, keep the show family-friendly.",
          "Rules to check before an adults-only show (a summary, not legal advice): who holds the liquor licence or Special Occasion Permit (AGCO), "
          "the venue's municipal alcohol policy, the Ontario Athletics Commission, and the alcohol advertising rules (no alcohol branding near anything "
          "aimed at under-19s). A family show with a fenced 19+ area keeps the kids and still sells drinks to adults."):
    note(wsA, r, 1, t); wsA.merge_cells(start_row=r, start_column=1, end_row=r, end_column=8); wsA.row_dimensions[r].height = 44; r += 1
wsA.freeze_panes = "A4"

# ------------------------------------------------------------------ Ticket Tracker
wsK = wb.create_sheet("Ticket Tracker")
KC = [("Week ending", 13), ("GA sold (total so far)", 12), ("VIP seats sold (total)", 12), ("VIP table seats sold (total)", 12),
      ("Comps given (total)", 11), ("Paid tickets", 10), ("Ticket money so far ($)", 13), ("% of break-even", 11),
      ("Weeks to show", 9), ("Paid tickets still needed", 12), ("Needed per week", 11), ("Notes", 36)]
banner(wsK, "TICKET TRACKER", "Pick the show, then once a week type the running totals from Eventbrite. It shows how close you are to break-even and how many tickets a week you still need.", len(KC))
for i, (h, w) in enumerate(KC, 1):
    wsK.column_dimensions[L(i)].width = w
labels = [("Show (pick the venue)", None), ("Show date", None), ("Paid tickets to break even", "be"), ("Paid tickets at the expected crowd", "exp"), ("Sellable capacity", "cap")]
hdr_rng = f"Compare!$B${R['hdr']}:${LASTC}${R['hdr']}"
for k, (lab, kind) in enumerate(labels):
    rr = 4 + k
    wsK.cell(rr, 1, lab).font = f_lab
    c = wsK.cell(rr, 3); c.border = box
    if kind is None:
        c.font = f_in; c.fill = fill_in
    else:
        rowkey = {"be": "be_t", "exp": "ex_paid", "cap": "cap"}[kind]
        c.value = f'=IFERROR(INDEX(Compare!$B${R[rowkey]}:${LASTC}${R[rowkey]},MATCH($C$4,{hdr_rng},0)),"")'
        c.font = f_calc; c.number_format = NUM
    wsK.merge_cells(start_row=rr, start_column=1, end_row=rr, end_column=2)
wsK["C4"] = VEN[0]["name"] + " · " + VEN[0]["town"]
wsK["C4"].font = f_in
wsK["C5"].number_format = "mmm d, yyyy"
dvs = DataValidation(type="list", formula1=hdr_rng, allow_blank=False); wsK.add_data_validation(dvs); dvs.add("C4")
dvd2 = DataValidation(type="date", allow_blank=True); wsK.add_data_validation(dvd2); dvd2.add("C5")
wsK.column_dimensions["C"].width = 30
note(wsK, 4, 4, "Pick from the list (it's the venue names on Compare).")
note(wsK, 5, 4, "Type the show date.")
HR2 = 10
for i, (h, w) in enumerate(KC, 1):
    c = wsK.cell(HR2, i, h); c.font = f_head; c.fill = fill_navy; c.border = box; c.alignment = Alignment(wrap_text=True, vertical="center")
wsK.row_dimensions[HR2].height = 42
GA, VS, VT = f"Inputs!$B${tier_start}", f"Inputs!$B${tier_start+1}", f"Inputs!$B${tier_start+2}"
for rr in range(HR2 + 1, HR2 + 31):
    for i in range(1, len(KC) + 1):
        c = wsK.cell(rr, i); c.border = box
        if i in (1, 2, 3, 4, 5, 12):
            c.font = f_in; c.fill = fill_in
        else:
            c.font = f_calc
    wsK.cell(rr, 1).number_format = "mmm d"
    blank = f'A{rr}=""'
    wsK.cell(rr, 6, f'=IF({blank},"",SUM(B{rr}:D{rr}))').number_format = NUM
    wsK.cell(rr, 7, f'=IF({blank},"",B{rr}*{GA}+C{rr}*{VS}+D{rr}*{VT})').number_format = USD
    wsK.cell(rr, 8, f'=IF(OR({blank},$C$6="",$C$6=0),"",F{rr}/$C$6)').number_format = PCT
    wsK.cell(rr, 9, f'=IF(OR({blank},$C$5=""),"",MAX(0,ROUND(($C$5-A{rr})/7,1)))').number_format = "0.0"
    wsK.cell(rr, 10, f'=IF(OR({blank},$C$6=""),"",MAX(0,$C$6-F{rr}))').number_format = NUM
    wsK.cell(rr, 11, f'=IF(OR({blank},J{rr}="",I{rr}=""),"",IF(J{rr}=0,0,IF(I{rr}<=0,J{rr},ROUNDUP(J{rr}/I{rr},0))))').number_format = NUM
for rr in range(HR2 + 1, HR2 + 31):
    for col in "FGHIJK":
        wsK[f"{col}{rr}"].font = f_calc
wsK.conditional_formatting.add(f"H{HR2+1}:H{HR2+30}", CellIsRule(operator="greaterThanOrEqual", formula=["1"], fill=PatternFill("solid", fgColor="C6EFCE")))
wsK.freeze_panes = wsK.cell(HR2 + 1, 2)
note(wsK, 9, 1, "Ticket money uses the prices on Inputs and doesn't take off Eventbrite fees. Comps don't count toward break-even.")
wsK.merge_cells(start_row=9, start_column=1, end_row=9, end_column=8)

# ------------------------------------------------------------------ Start Here
ws0 = wb.create_sheet("Start Here", 0)
banner(ws0, "CCCL EVENT PROFIT PREDICTOR",
       "Compare towns and venues for the next Cottage Country Combat League show before signing anything. "
       "Built for Cody Purtell by Kyle. Opens in Google Sheets or Excel.", 3)
ws0.column_dimensions["A"].width = 4; ws0.column_dimensions["B"].width = 34; ws0.column_dimensions["C"].width = 90
r = 4
w0 = ws0.cell(r, 2, "Good to know: on Cody's numbers the first show made about $18,000. That profit went into the gym and other projects, "
                    "which is why it felt like breaking even. Profits on Compare are the show's profit before any money goes to the gym.")
w0.font = Font(name=F, size=11, bold=True, color="1D5A3A"); w0.alignment = wrap; w0.fill = PatternFill("solid", fgColor="E8F3EC")
ws0.merge_cells(start_row=r, start_column=2, end_row=r, end_column=3); ws0.row_dimensions[r].height = 34
r = 6
steps = [
    ("1", "Haliburton 2026 Check tab", "The first show's real numbers. When Cody has the Eventbrite payout report, put the exact profit that went to the gym in Part 2."),
    ("2", "Inputs tab", "Cody's real 2026 prices and costs are already in. Replace every ESTIMATE with a real quote or invoice."),
    ("3", "Inputs tab, venue columns", "Call each venue (numbers are in the 'Who to call' row) for capacity, the full event price and open dates."),
    ("4", "Compare tab", "Read the verdict, the break-even ticket count and how much sponsorship each option needs. "
                         "Venues whose expected crowd is under Cody's 1,000 minimum get flagged."),
    ("5", "Before committing", "Check the stress-test row: can CCCL survive 30% fewer people and half the sponsors pulling out?"),
    ("6", "2027 Venues tab", "24 venues in Cody's 10 towns plus casinos and event centres: capacity, winter use, rates, who to call, clashes and other promotions."),
    ("7", "Town Scout tab", "25 smaller towns within 2 hours, ranked. To compare any venue that isn't on Compare, type it over a venue column on Inputs."),
    ("8", "19+ Show Check tab", "Is an adults-only show worth it? Compares a family show, a family show with a 19+ bar area, and an adults-only show at every venue."),
    ("9", "Ticket Tracker tab", "Once tickets are on sale: pick the show, type Eventbrite totals once a week, and see how far from break-even you are."),
]
x = ws0.cell(r, 2, "HOW TO USE IT"); x.font = f_sec; r += 1
for n, where, what in steps:
    ws0.cell(r, 1, n).font = Font(name=F, size=12, bold=True, color=GOLD)
    ws0.cell(r, 2, where).font = f_bold
    ws0.cell(r, 3, what).font = f_lab; ws0.cell(r, 3).alignment = wrap
    if len(what) > 95: ws0.row_dimensions[r].height = 28
    r += 1
r += 1
x = ws0.cell(r, 2, "COLOUR KEY"); x.font = f_sec; r += 1
key = [("Yellow cell, blue text", "You type here.", fill_in, f_in),
       ("Green text", "Pulled from another tab. Don't type over it.", None, f_link),
       ("Black text", "Calculated. Don't type over it.", None, f_calc),
       ("ESTIMATE in the notes", "A placeholder. Replace it with a real quote or invoice.", None, f_note)]
for a, b, fill, font in key:
    c = ws0.cell(r, 2, a); c.font = font
    if fill: c.fill = fill
    ws0.cell(r, 3, b).font = f_lab; r += 1
r += 1
x = ws0.cell(r, 2, "WHAT IT ANSWERS"); x.font = f_sec; r += 1
for q in ("How many tickets do we have to sell at each venue just to break even?",
          "How much sponsorship do we need if the crowd is small?",
          "Which town makes the most money at a normal crowd?",
          "What happens if 30% fewer people come and half the sponsors back out?"):
    ws0.cell(r, 3, "•  " + q).font = f_lab; r += 1
r += 1
x = ws0.cell(r, 2, "SOURCES FOR THE NUMBERS"); x.font = f_sec; r += 1
for a, b in (
    ("First show (May 2, 2026)", "The Highlander, Mar 12 2026: thehighlander.ca/2026/03/12/fight-night-coming-to-haliburton-county/  "
                                 "(A.J. LaRue Arena, 8-12 amateur bouts, 800-1,000 expected, sanctioned by the Ontario Athletics Commission)"),
    ("First show real numbers", "Cody's answers, Oct 2026: GA $60, VIP $100, tables $1,500 for 8 (6 sold), 400 tickets, 400-600 people, "
                                "15 sponsors / $18,000, Eventbrite fees paid by CCCL, bar run by others, costs listed at $29,000. Profit went into the gym and other projects (Kyle)."),
    ("Haliburton venue rates", "Dysart et al facility rentals: dysartetal.ca (A.J. LaRue Community Centre): $1,030/day profit-making event, $30/hr setup with staff"),
    ("Bancroft venue", "Town of Bancroft and mybancroftnow.com (Sept 2026): NHCC upgrades complete; rates and capacity not published; call 613-332-2062"),
    ("Amateur event permit", "O. Reg. 507/24 and ontario.ca: $600 per show (one permit should cover the whole card), $3,000 for up to 8 shows in a year; $2M liability insurance; physician on site. Details on the Fight Night Kit's Rules tab."),
    ("Marmora", "marmoraandlake.ca community centre: dry floor 1,000 capacity, $600/day event"),
    ("City arenas", "Peterborough Memorial Centre Tech Pack 2024 (setups 1,722 to 5,501); Visit Kingston (Slush Puppie Place 5,700 cold-season); "
                    "tributecommunitiescentre.com (1,000 to 7,600). None publish rental rates: those columns are ESTIMATES."),
    ("2027 Venues tab", "24 venues and 5 other promotions researched by Claude, October 2026. Sources are in each row."),
    ("Town Scout tab", "25 towns researched by Claude from municipal sites and local news, October 2026. Sources are in each row."),
    ("Everything else", "ESTIMATES by Kyle. Replace with real quotes and invoices."),
):
    ws0.cell(r, 2, a).font = f_bold; ws0.cell(r, 2).alignment = Alignment(vertical="top")
    c = ws0.cell(r, 3, b); c.font = f_lab; c.alignment = wrap
    ws0.row_dimensions[r].height = 30
    r += 1
ws0.sheet_view.showGridLines = False

for ws in wb.worksheets:
    ws.sheet_properties.tabColor = GOLD if ws.title in ("Start Here", "Compare") else NAVY
    ws.page_setup.orientation = "landscape"; ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
wb.active = 0
wb.save(OUT)
import json; json.dump({"R": R, "v": v, "g": g}, open(OUT + ".map.json", "w"))
print("saved", OUT)
