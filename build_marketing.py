"""Build 'CCCL Marketing Kit.xlsx' (Google Sheets compatible).
usage: python3 build_marketing.py OUT [research/marketing_tools.json]"""
import sys, json, os, datetime as dt
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.utils import get_column_letter as L

OUT = sys.argv[1]
MT = json.load(open(sys.argv[2])) if len(sys.argv) > 2 and os.path.exists(sys.argv[2]) else None

INK, PINE, PINE_L, CREAM, RED = "151C19", "354940", "E6EEE9", "F2ECDC", "AE2221"
F = "Arial"
f_title = Font(name=F, size=16, bold=True, color=CREAM)
f_sub = Font(name=F, size=10, color="D9D5C8")
f_head = Font(name=F, size=10, bold=True, color="FFFFFF")
f_sec = Font(name=F, size=11, bold=True, color=PINE)
f_lab = Font(name=F, size=10)
f_bold = Font(name=F, size=10, bold=True)
f_in = Font(name=F, size=10, color="0000FF")
f_calc = Font(name=F, size=10, color="000000")
f_link = Font(name=F, size=10, color="008000")
f_note = Font(name=F, size=9, italic=True, color="555555")
f_flag = Font(name=F, size=10, bold=True, color="C00000")
fill_ink = PatternFill("solid", fgColor=INK)
fill_pine = PatternFill("solid", fgColor=PINE)
fill_pl = PatternFill("solid", fgColor=PINE_L)
fill_in = PatternFill("solid", fgColor="FFFF00")
fill_ok = PatternFill("solid", fgColor="C6EFCE")
fill_bad = PatternFill("solid", fgColor="F8CBAD")
fill_warn = PatternFill("solid", fgColor="FFEB9C")
thin = Side(style="thin", color="BFBFBF")
box = Border(top=thin, bottom=thin, left=thin, right=thin)
wrap = Alignment(wrap_text=True, vertical="top")
top = Alignment(vertical="top")
USD = '$#,##0;($#,##0);"-"'
USD2 = '$#,##0.00;($#,##0.00);"-"'
NUM = '#,##0;(#,##0);"-"'
DATE = 'ddd mmm d'

wb = Workbook()

def banner(ws, title, sub, width):
    for r in (1, 2):
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=width)
        for c in range(1, width + 1):
            ws.cell(r, c).fill = fill_ink
    ws["A1"] = title; ws["A1"].font = f_title; ws["A1"].alignment = Alignment(vertical="center")
    ws["A2"] = sub; ws["A2"].font = f_sub; ws["A2"].alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[1].height = 30; ws.row_dimensions[2].height = 34

def head(ws, r, labels, widths=None, height=34):
    for i, h in enumerate(labels, 1):
        c = ws.cell(r, i, h); c.font = f_head; c.fill = fill_pine; c.border = box
        c.alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[r].height = height
    if widths:
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[L(i)].width = w

def section(ws, r, text, width):
    ws.cell(r, 1, text).font = f_head
    for c in range(1, width + 1):
        ws.cell(r, c).fill = fill_pine
    ws.row_dimensions[r].height = 20

def inp(c, v=None, fmt=None):
    if v is not None: c.value = v
    c.font = f_in; c.fill = fill_in; c.border = box
    if fmt: c.number_format = fmt

def note(ws, r, c, text, merge_to=None):
    x = ws.cell(r, c, text); x.font = f_note; x.alignment = wrap
    if merge_to: ws.merge_cells(start_row=r, start_column=c, end_row=r, end_column=merge_to)

def dv_list(ws, f1, rng):
    dv = DataValidation(type="list", formula1=f1, allow_blank=True); ws.add_data_validation(dv); dv.add(rng)

def q(items): return '"' + ",".join(items) + '"'
YN = ["Yes", "No"]

# ------------------------------------------------------------------ Show Setup
wsS = wb.active; wsS.title = "Show Setup"
banner(wsS, "SHOW SETUP", "Fill this in once per show. The campaign dates and the 'worth it' checks use these numbers.", 4)
for c, w in zip("ABCD", (46, 34, 4, 70)): wsS.column_dimensions[c].width = w
S = {}
r = 4
section(wsS, r, "THE SHOW", 4); r += 1
for key, lab, val, fmt, nt in (
    ("name", "Show name", "CCCL 2", None, ""),
    ("town", "Town", None, None, "Fill in once the show is booked; use the same town and date as the Fight Night Kit."),
    ("date", "Show date", None, 'ddd mmm d, yyyy', "Fill in once the show is booked: the campaign dates follow it. Until then they stay blank."),
    ("link", "Eventbrite ticket link", "", None, "The main link. Every source on the Ticket Sources tab gets its own tracking link or code."),
    ("fb", "CCCL Facebook page", "", None, "CCCL's own page (Kyle, Oct 2026)."),
    ("ig", "CCCL Instagram", "", None, "CCCL's own account."),
    ("tags", "Hashtags", "#CCCL #CottageCountryCombatLeague #AmateurMMA", None, "Add the town: #Haliburton, #Peterborough, #Kingston..."),
):
    wsS.cell(r, 1, lab).font = f_lab
    inp(wsS.cell(r, 2), val if val != "" else None, fmt)
    if nt: note(wsS, r, 4, nt)
    S[key] = f"'Show Setup'!$B${r}"; r += 1
r += 1
section(wsS, r, "IS A PAID AD WORTH IT?", 4); r += 1
for key, lab, val, fmt, nt in (
    ("net", "Money CCCL keeps per paid ticket ($)", 73, USD2,
     "From the Profit Predictor (Compare tab, 'Total per paid ticket'). About $73 at 2026 prices after Eventbrite's Canadian fees (3.5% + $1.29 per ticket, plus 2.9% processing)."),
    ("target", "Most you'll pay in ads to sell one ticket ($)", 20, USD2,
     "STARTING POINT. Below $73 the ad pays for itself on paper, but some of those buyers would have come anyway, so aim well under it. Raise it if ads clearly bring new people."),
    ("test", "First test budget per ad ($)", 50, USD,
     "Spend a little, check the tickets on that ad's tracking link after a week, then keep or stop."),
):
    wsS.cell(r, 1, lab).font = f_lab; wsS.cell(r, 1).alignment = wrap
    inp(wsS.cell(r, 2), val, fmt)
    note(wsS, r, 4, nt)
    wsS.row_dimensions[r].height = 30 if len(nt) > 90 else 22
    S[key] = f"'Show Setup'!$B${r}"; r += 1
wsS.freeze_panes = "A4"

# ------------------------------------------------------------------ Campaign
wsC = wb.create_sheet("Campaign")
CC = [("Weeks before (minus = days after)", 12), ("Date", 12), ("Channel", 20), ("What to do", 58), ("Graphic / asset", 22), ("Owner", 10), ("Status", 11), ("Notes", 34)]
banner(wsC, "SHOW CAMPAIGN",
       "Everything that sells tickets, week by week. Dates follow Show Setup. Every post needs Cody's OK (Approvals tab in Constrictor HQ). "
       "Matches the Countdown tab in the Fight Night Kit.", len(CC))
head(wsC, 4, [c[0] for c in CC], [c[1] for c in CC])
CAMP = [
    (16, "CCCL Facebook + Instagram", "Save the date: town, venue, date. Share to BBG's pages.", "01 (as Save the date) + 02 story", "Kyle"),
    (16, "Facebook event", "Create the Facebook event on the CCCL page; ask BBG, the venue and sponsors to co-host or share.", "09 Facebook event cover", "Kyle"),
    (14, "Eventbrite", "Tickets on sale. Make a tracking link for every source on the Ticket Sources tab (Marketing > Tracking links) before sharing anything. Add the Meta Pixel (Marketing > Tracking pixels) if you'll run ads.", "", "Kyle"),
    (14, "CCCL Facebook + Instagram", "Tickets on sale post with prices and the link. Pin it.", "01 Tickets on sale + 02 story", "Kyle"),
    (14, "Press", "Release 1: show announced (Fight Night Kit, Press tab).", "", "Cody"),
    (14, "Community calendars", "List the show on kawarthaNOW, Moose FM and the town's event calendar.", "", "Kyle"),
    (14, "Email", "Email last show's ticket buyers from Eventbrite (Marketing > Email your audience; up to 250 emails a day).", "", "Kyle"),
    (13, "CCCL Facebook + Instagram", "Sponsors wanted: 'put your business cage-side' with the sponsor sheet.", "Sponsors wanted", "Kyle"),
    (12, "CCCL Facebook + Instagram", "Fighters wanted: Ontario gyms, link to the sign-up form.", "Fighters wanted", "Kyle"),
    (12, "Gym (BBG)", "Members and families: group-ticket push at the gym and in the members' chat.", "", "Cody"),
    (10, "Posters", "Posters up around town with a QR code to the posters' tracking link (Poster Drop tab).", "03 Poster (11x17 and letter)", "Volunteer"),
    (10, "Shorts / Reels / TikTok", "Behind the scenes: Cody training, the venue, why this town. 30-60 seconds.", "", "Kyle"),
    (8, "Paid ads (optional)", "Ad test 1: small budget, its own tracking link. Check after a week (Ads tab).", "Ad image", "Kyle"),
    (8, "CCCL Facebook + Instagram", "Thank-you post for sponsors signed so far (only what their package includes).", "07 Sponsor thank-you", "Kyle"),
    (6, "Fighters", "Send each confirmed fighter their Fighter Package (through the coach) and adults their promo pack: personal ticket link, ready-made posts, graphics.", "11 + 12 I'm fighting (each adult)", "Cody"),
    (6, "CCCL Facebook + Instagram", "Fighter spotlights start: one or two a week, adults who agree only.", "06 Fighter spotlight", "Kyle"),
    (5, "Radio", "Pitch an interview or event listing to local radio (Moose FM, Quinte News...).", "", "Cody"),
    (4, "CCCL Facebook + Instagram", "Fight card announced: official bouts only, adults only by name. Swap the early posters for the main-event poster.", "04 + 16 card, 05 vs, 10 poster", "Kyle"),
    (4, "Press", "Release 2: the card (Fight Night Kit, Press tab).", "", "Cody"),
    (4, "Sponsors", "Individual posts for Title, Gold and Silver; one group post for Bronze and Community. Send each sponsor a share kit.", "15 Proud sponsor", "Kyle"),
    (3, "Paid ads (optional)", "Scale only ads under the target cost per ticket; stop the rest.", "", "Kyle"),
    (3, "Stories", "Countdown stories: 21, 14, 7 days.", "02 Countdown story", "Kyle"),
    (2, "CCCL Facebook + Instagram", "VIP tables or tickets running low, only if it's true (Ticket Tracker).", "", "Kyle"),
    (2, "CCCL Facebook + Instagram", "Livestream link announced for people who can't come.", "Livestream", "Kyle"),
    (1, "CCCL Facebook + Instagram", "Fight week: final card graphic, doors time, parking, what to bring.", "16 Fight card story", "Kyle"),
    (1, "Fighters + gyms", "Fight week reminder to every fighter and coach to share their link one last time.", "", "Cody"),
    (0, "Stories / live", "Show day: weigh-in results post once weigh-ins are done, 'doors at 6', clips between bouts. No youth without a parent's OK.", "13 Weigh-in results", "Kyle"),
    (-1, "CCCL Facebook + Instagram", "The day after: official results and thank-you to fighters, officials, medical staff, volunteers, venue.", "08 Results", "Kyle"),
    (-1, "Fighters", "The day after: each adult fighter's result post: WIN, or THANK YOU for a loss. Official results only.", "14 Win (each adult)", "Kyle"),
    (-1, "Press", "The day after: Release 3, results (official only).", "", "Cody"),
    (-3, "Sponsors", "Thank-you posts with photos of their banners and logos (within 3 days).", "07 Sponsor thank-you", "Kyle"),
    (-7, "YouTube + Shorts / Reels", "Highlight video and clips; 'next show' teaser.", "Thumbnail", "Kyle"),
    (-7, "Report", "Which sources sold tickets (Ticket Sources tab) into the Fight Night Kit's Results & Report.", "", "Kyle"),
]
CHANNELS = sorted({c[1] for c in CAMP})
for k, (w, ch, what, asset, owner) in enumerate(CAMP):
    rr = 5 + k
    inp(wsC.cell(rr, 1), w, "0")
    c = wsC.cell(rr, 2, f'=IF({S["date"]}="","",IF(A{rr}>=0,{S["date"]}-A{rr}*7,{S["date"]}-A{rr}))'); c.number_format = DATE; c.font = f_calc; c.border = box
    for i, v in ((3, ch), (4, what), (5, asset)):
        x = wsC.cell(rr, i, v); x.font = f_lab; x.border = box; x.alignment = wrap
    inp(wsC.cell(rr, 6), owner); inp(wsC.cell(rr, 7), "To do"); inp(wsC.cell(rr, 8))
    wsC.row_dimensions[rr].height = max(30, 14 * (-(-len(what) // 60)) + 4)
lastC = 5 + len(CAMP) - 1
for rr in range(lastC + 1, lastC + 16):
    for i in range(1, len(CC) + 1): inp(wsC.cell(rr, i))
    c = wsC.cell(rr, 2, f'=IF(OR(A{rr}="",{S["date"]}=""),"",IF(A{rr}>=0,{S["date"]}-A{rr}*7,{S["date"]}-A{rr}))'); c.number_format = DATE; c.font = f_calc; c.fill = PatternFill(fill_type=None)
dv_list(wsC, q(["To do", "Drafted", "Approved", "Done", "Skipped"]), f"G5:G{lastC + 15}")
for s_, fl in (("Done", fill_ok), ("Approved", fill_pl), ("Drafted", fill_warn)):
    wsC.conditional_formatting.add(f"G5:G{lastC + 15}", CellIsRule(operator="equal", formula=[f'"{s_}"'], fill=fl))
wsC.conditional_formatting.add(f"B5:B{lastC + 15}", FormulaRule(formula=[f'AND(B5<>"",B5<TODAY(),G5<>"Done",G5<>"Skipped")'], fill=fill_bad))
wsC.freeze_panes = "D5"

# ------------------------------------------------------------------ Ticket Sources
wsT = wb.create_sheet("Ticket Sources")
TC = [("Source", 30), ("Type", 14), ("Tracking link or code", 32), ("Cost ($)", 9), ("Tickets sold", 9), ("Ticket money ($)", 11),
      ("Cost per ticket", 10), ("Worth it?", 13), ("Notes", 32)]
banner(wsT, "WHERE THE TICKETS CAME FROM",
       "Give every source its own Eventbrite tracking link (Marketing > Tracking links), then copy the sales per link in each week (Reporting > Tracking links). "
       "That's how you find out which posts, posters, fighters and ads actually sell tickets.", len(TC))
HT = 7
wsT.cell(4, 1, "Total tickets sold (from Eventbrite)").font = f_bold; inp(wsT.cell(4, 2), None, NUM)
wsT.cell(5, 1, "Tickets with a known source").font = f_bold
wsT.cell(5, 2, f"=SUM(E{HT+1}:E{HT+80})").font = f_calc
wsT.cell(4, 4, "Share with a known source").font = f_bold
c = wsT.cell(4, 6, f'=IF(N(B4)=0,"",B5/B4)'); c.number_format = "0%"; c.font = f_calc
note(wsT, 5, 4, "Use tracking links, not promo codes: an Eventbrite promo code has to give a discount or unlock a hidden ticket. "
     "Sponsor tickets: add them as Complimentary attendees (Manage Attendees > Add Attendees), which has no fees.", 9)
wsT.row_dimensions[5].height = 30
head(wsT, HT, [c[0] for c in TC], [c[1] for c in TC])
TYPES = ["Social (free)", "Fighter", "Gym", "Sponsor", "Ad", "Poster", "Press / radio", "Calendar", "Email", "Walk-up", "Other"]
seed = [("CCCL Facebook (posts and event)", "Social (free)"), ("CCCL Instagram", "Social (free)"), ("BBG Facebook + Instagram", "Social (free)"), ("Cody's page (Facebook + Instagram)", "Social (free)"),
        ("Posters (QR code)", "Poster"), ("Press releases", "Press / radio"), ("Radio", "Press / radio"), ("Community calendars", "Calendar"),
        ("Email to past buyers", "Email"), ("BBG members", "Gym")] + \
       [(f"Fighter: [name]", "Fighter") for _ in range(16)] + [(f"Gym: [name]", "Gym") for _ in range(4)] + \
       [(f"Sponsor: [name]", "Sponsor") for _ in range(6)] + [(f"Ad test {i}", "Ad") for i in range(1, 6)] + [("Walk-up at the door", "Walk-up")]
for k, (src, typ) in enumerate(seed):
    rr = HT + 1 + k
    inp(wsT.cell(rr, 1), src); inp(wsT.cell(rr, 2), typ); inp(wsT.cell(rr, 3)); inp(wsT.cell(rr, 4), None, USD); inp(wsT.cell(rr, 5), None, NUM)
    inp(wsT.cell(rr, 6), None, USD); inp(wsT.cell(rr, 9))
for rr in range(HT + 1, HT + 81):
    if wsT.cell(rr, 1).value is None:
        for i in (1, 2, 3, 4, 5, 6, 9): inp(wsT.cell(rr, i))
        wsT.cell(rr, 4).number_format = USD; wsT.cell(rr, 5).number_format = NUM; wsT.cell(rr, 6).number_format = USD
    c = wsT.cell(rr, 7, f'=IF(OR(N(D{rr})=0,N(E{rr})=0),"",D{rr}/E{rr})'); c.number_format = USD2; c.font = f_calc; c.border = box
    c = wsT.cell(rr, 8, f'=IF(A{rr}="","",IF(N(D{rr})=0,IF(N(E{rr})>0,"Free","")'
                         f',IF(N(E{rr})=0,"No sales yet",IF(G{rr}<={S["target"]},"Worth it",IF(G{rr}<={S["net"]},"Pays, barely","Not worth it")))))')
    c.font = f_bold; c.border = box
dv_list(wsT, q(TYPES), f"B{HT+1}:B{HT+80}")
for s_, fl in (("Worth it", fill_ok), ("Free", fill_ok), ("Pays, barely", fill_warn), ("Not worth it", fill_bad)):
    wsT.conditional_formatting.add(f"H{HT+1}:H{HT+80}", CellIsRule(operator="equal", formula=[f'"{s_}"'], fill=fl))
wsT.freeze_panes = wsT.cell(HT + 1, 2)
# by type summary
sr = HT + 83
section(wsT, sr, "BY TYPE", 9); sr += 1
for i, h in enumerate(("Type", "Tickets", "Cost ($)", "Cost per ticket"), 1):
    c = wsT.cell(sr, i, h); c.font = f_bold; c.fill = fill_pl; c.border = box
sr += 1
for t in TYPES:
    wsT.cell(sr, 1, t).font = f_lab
    wsT.cell(sr, 2, f'=SUMIF($B${HT+1}:$B${HT+80},A{sr},$E${HT+1}:$E${HT+80})').number_format = NUM
    wsT.cell(sr, 3, f'=SUMIF($B${HT+1}:$B${HT+80},A{sr},$D${HT+1}:$D${HT+80})').number_format = USD
    c = wsT.cell(sr, 4, f'=IF(OR(C{sr}=0,B{sr}=0),"",C{sr}/B{sr})'); c.number_format = USD2
    for i in range(1, 5): wsT.cell(sr, i).border = box
    sr += 1

# ------------------------------------------------------------------ Ads
wsA = wb.create_sheet("Ads")
AC = [("Start date", 11), ("Platform", 16), ("Ad or boosted post", 30), ("Who it targets", 30), ("Tracking link name", 20),
      ("Spent ($)", 9), ("Link clicks", 9), ("Tickets (from its link)", 10), ("Cost per ticket", 10), ("Keep or stop?", 14), ("Notes", 30)]
banner(wsA, "AD TESTS",
       "Only spend where it pays. Each ad gets its own tracking link. Start with the test budget on Show Setup, check its tickets after a week, "
       "keep ads under the target cost per ticket and stop the rest.", len(AC))
head(wsA, 4, [c[0] for c in AC], [c[1] for c in AC])
for rr in range(5, 35):
    for i in (1, 2, 3, 4, 5, 6, 7, 8, 11): inp(wsA.cell(rr, i))
    wsA.cell(rr, 1).number_format = "mmm d"; wsA.cell(rr, 6).number_format = USD; wsA.cell(rr, 8).number_format = NUM
    c = wsA.cell(rr, 9, f'=IF(OR(N(F{rr})=0,N(H{rr})=0),"",F{rr}/H{rr})'); c.number_format = USD2; c.border = box
    c = wsA.cell(rr, 10, f'=IF(N(F{rr})=0,"",IF(N(H{rr})=0,IF(F{rr}>={S["test"]},"STOP: no sales","Wait"),IF(I{rr}<={S["target"]},"KEEP",IF(I{rr}<={S["net"]},"Maybe: pays, barely","STOP"))))')
    c.font = f_bold; c.border = box
dv_list(wsA, q(["Facebook + Instagram", "Facebook", "Instagram", "Eventbrite", "Other"]), "B5:B34")
for s_, fl in (("KEEP", fill_ok), ("STOP", fill_bad), ("STOP: no sales", fill_bad)):
    wsA.conditional_formatting.add("J5:J34", CellIsRule(operator="equal", formula=[f'"{s_}"'], fill=fl))
wsA.conditional_formatting.add("J5:J34", FormulaRule(formula=['LEFT(J5,5)="Maybe"'], fill=fill_warn))
note(wsA, 36, 1, "How to: use the ad's own tracking link as its website link. Target a radius around the show town (Meta allows 1-80 km around a pin) and adults 18+ "
     "(19+ if the ad mentions a bar). Ads Manager shows the minimum daily budget in dollars when you set it; a few dollars a day is enough to test. "
     "Optional: add the Meta Pixel in Eventbrite (Marketing > Tracking pixels) so Meta sees purchases.", 11)
note(wsA, 37, 1, "Ad rules: Meta bans shocking or excessively violent images, so use face-offs, posed fighter photos and event graphics, never blood or a knockout. "
     "No youth athletes in ads. No alcohol in any ad that could reach under-19s. Real numbers only ('400 tickets sold at our first show').", 11)
wsA.row_dimensions[36].height = 42; wsA.row_dimensions[37].height = 30
wsA.freeze_panes = "C5"

# ------------------------------------------------------------------ Livestream
wsL = wb.create_sheet("Livestream")
banner(wsL, "LIVESTREAM: CAN IT MAKE MONEY?",
       "The 2026 stream made nothing. Two simple ways to charge for it, worked out for your numbers, plus what else is out there. "
       "Whatever you choose, the Athletics Commission gets a recording if the show is filmed.", 6)
for c, w in zip("ABCDEF", (40, 16, 16, 16, 16, 44)): wsL.column_dimensions[c].width = w
r = 4
section(wsL, r, "YOUR NUMBERS (ESTIMATES: change them)", 6); r += 1
LV = {}
for key, lab, val, fmt, nt in (("price", "Stream ticket price ($)", 15, USD2, "ESTIMATE. What a viewer pays to watch."),
                               ("buyers", "People who buy the stream", 50, NUM, "ESTIMATE. Fighters' families out of town are the likely buyers."),
                               ("fx", "US dollar in Canadian dollars", 1.38, '0.00', "ESTIMATE. Check today's rate; Dacast charges in US dollars."),
                               ("dacast_plan", "Dacast plan for one month (US$)", 39, USD, "Dacast lists plans from about US$39 to US$63 a month; the paywall may be an add-on. Check before signing up.")):
    wsL.cell(r, 1, lab).font = f_lab
    inp(wsL.cell(r, 2), val, fmt); note(wsL, r, 3, nt, 6)
    LV[key] = f"$B${r}"; r += 1
r += 1
section(wsL, r, "WHAT CCCL WOULD KEEP", 6); r += 1
for i, h in enumerate(("Option", "Fees per buyer ($)", "CCCL keeps per buyer ($)", "Monthly or fixed cost ($)", "CCCL keeps in total ($)", "How it works"), 1):
    c = wsL.cell(r, i, h); c.font = f_bold; c.fill = fill_pl; c.border = box; c.alignment = wrap
wsL.row_dimensions[r].height = 30
r += 1
P_, B_, FX = LV["price"], LV["buyers"], LV["fx"]
rows_ = [("Eventbrite online ticket", f"={P_}*0.035+1.29+{P_}*0.029", 0,
          "A separate online event on Eventbrite. Buyers get a private page with the stream link. Same Canadian fees as tickets (3.5% + $1.29, plus 2.9%). Links can be shared, so some people will watch for free."),
         ("Dacast pay-per-view", f"={P_}*0.099+0.40*{FX}", f"={LV['dacast_plan']}*{FX}",
          "A paywall around the stream: 9.9% + US$0.40 per sale plus a monthly plan. Harder to share than a link.")]
first = r
for name, fee, fixed, how in rows_:
    wsL.cell(r, 1, name).font = f_bold
    c = wsL.cell(r, 2, fee); c.number_format = USD2
    c = wsL.cell(r, 3, f"={P_}-B{r}"); c.number_format = USD2
    c = wsL.cell(r, 4, fixed); c.number_format = USD
    c = wsL.cell(r, 5, f"=C{r}*{B_}-D{r}"); c.number_format = USD; c.font = f_bold
    c = wsL.cell(r, 6, how); c.font = f_lab; c.alignment = wrap
    for i in range(1, 7): wsL.cell(r, i).border = box
    wsL.row_dimensions[r].height = 44; r += 1
r += 1
section(wsL, r, "OTHER OPTIONS (researched Oct 2026)", 6); r += 1
others = []
if MT and MT.get("livestream"):
    keep = ("YouTube", "Facebook Stars", "TrillerTV", "Vimeo OTT", "UFC Fight Pass")
    for l in MT["livestream"]:
        if any(l["option"].startswith(k) for k in keep):
            others.append((l["option"], l.get("fees", ""), l.get("source_url", "")))
short = {"YouTube": "YouTube pays creators 70% of channel memberships and Super Chat. Members-only streams need a channel with 500+ subscribers and 3,000 watch hours.",
         "Facebook Stars": "Viewers send Stars; Meta pays 1 cent per Star. Facebook's paid online events ended in 2023.",
         "TrillerTV": "Pay-per-view platform, but risky for 2027: AEW sued it in May 2026 for unpaid revenue and its operating company told a court it can't pay its debts.",
         "Vimeo OTT": "Live pay-per-view only on a custom-priced Enterprise plan.",
         "UFC Fight Pass": "Carries bigger Canadian promotions under multi-year deals. A long-term goal, not a 2027 option."}
for name, fees, src in others:
    k = next(k for k in short if name.startswith(k))
    wsL.cell(r, 1, name.split(" (")[0].split(" - ")[0]).font = f_bold; wsL.cell(r, 1).alignment = wrap
    c = wsL.cell(r, 2, short[k]); c.font = f_lab; c.alignment = wrap
    wsL.merge_cells(start_row=r, start_column=2, end_row=r, end_column=5)
    c = wsL.cell(r, 6, src); c.font = f_note; c.alignment = wrap
    wsL.row_dimensions[r].height = 32; r += 1
r += 1
note(wsL, r, 1, "The safest money from the stream is sponsors: every package already includes the livestream (logo or name sized by level). "
     "Charging for the stream only makes sense once enough out-of-town family and fans want to watch. Put any stream money on the Predictor's Inputs tab.", 6)
wsL.row_dimensions[r].height = 30

# ------------------------------------------------------------------ Poster Drop
wsP = wb.create_sheet("Poster Drop")
PC = [("Place", 28), ("Kind of place", 22), ("Area / address", 26), ("Asked permission?", 10), ("Put up (date)", 11), ("Taken down (date)", 11), ("Notes", 34)]
banner(wsP, "POSTER DROP", "Where the posters go in the show town. Ask first, put them up about 10 weeks out, take them down after the show. Every poster has the QR code for the posters' tracking link.", len(PC))
head(wsP, 4, [c[0] for c in PC], [c[1] for c in PC])
KINDS = [("Every sponsor's business", "Sponsor"), ("The venue's lobby and front doors", "Venue"), ("Grocery store community boards", "Community board"),
         ("Gas stations and convenience stores", "Store"), ("Coffee shops (Tim Hortons, local cafes)", "Restaurant / cafe"),
         ("Every gym and fitness studio in town", "Gym"), ("Public library board", "Community board"), ("Legion, Lions or community hall", "Community board"),
         ("Hardware and farm supply stores", "Store"), ("Restaurants and pubs (adults)", "Restaurant / cafe"), ("Barber shops and salons", "Store"),
         ("Motels, cottage resorts and campgrounds (summer)", "Tourism"), ("Arena bulletin boards (other arenas nearby)", "Community board")]
for k, (place, kind) in enumerate(KINDS):
    rr = 5 + k
    inp(wsP.cell(rr, 1), place); inp(wsP.cell(rr, 2), kind)
    for i in (3, 4, 5, 6, 7): inp(wsP.cell(rr, i))
for rr in range(5 + len(KINDS), 65):
    for i in range(1, 8): inp(wsP.cell(rr, i))
for rr in range(5, 65):
    wsP.cell(rr, 5).number_format = "mmm d"; wsP.cell(rr, 6).number_format = "mmm d"
dv_list(wsP, q(YN), "D5:D64")
note(wsP, 66, 1, "The first rows are reminders of the kinds of places that work in small towns. Replace them with real places. No posters at schools.", 7)
wsP.freeze_panes = "B5"

# ------------------------------------------------------------------ Shot List
wsV = wb.create_sheet("Shot List")
VC = [("Shot", 40), ("When", 16), ("Used for", 28), ("Consent needed", 26), ("Who films", 12), ("Got it?", 8)]
banner(wsV, "FIGHT NIGHT SHOT LIST", "Photos and video to grab on the night, for the highlight video, sponsor thank-yous, press and next year's sponsor pitch.", len(VC))
head(wsV, 4, [c[0] for c in VC], [c[1] for c in VC])
SHOTS = [
    ("Empty venue, wide, before doors", "Afternoon", "Highlight intro, venue thank-you", "None"),
    ("Cage build time-lapse", "Load-in", "Reels / Shorts", "Crew OK"),
    ("Every sponsor's banner with the crowd behind it", "During the show", "Sponsor thank-you emails and posts", "None (crowd shot)"),
    ("Weigh-ins on the scale", "Weigh-ins", "Stories, fighter posts", "Adult fighters who agree only"),
    ("Fighters' walkouts", "Each bout", "Highlight video", "Fighter agrees (sign-up form)"),
    ("Crowd reactions, packed seats", "Main card", "Next year's sponsor pitch, ads", "None (crowd shot)"),
    ("Each finish and each hand raised", "Each bout", "Results posts, highlight video", "Fighter agrees"),
    ("Cody's welcome and the Title Sponsor read", "Opening", "Sponsor thank-you, highlight", "None"),
    ("Short winner interviews (15 seconds)", "After each bout", "Reels / Shorts, press", "Adult winners who agree"),
    ("VIP tables and sponsor guests", "During the show", "Next year's VIP table sales", "Ask the table first"),
    ("Volunteers and officials at work", "During the show", "Thank-you post", "Ask first"),
    ("Merch table and event shirts", "Intermission", "Merch posts", "None"),
    ("Livestream screen with the sponsor overlay", "During the show", "Proof for sponsors", "None"),
    ("Youth grappling", "Youth matches", "Only if a parent has said yes", "Parent's written OK for every child shown"),
]
for k, (shot, when, use, consent) in enumerate(SHOTS):
    rr = 5 + k
    for i, v in enumerate((shot, when, use, consent), 1):
        x = wsV.cell(rr, i, v); x.font = f_lab; x.border = box; x.alignment = wrap
    inp(wsV.cell(rr, 5), "Kyle"); inp(wsV.cell(rr, 6))
    wsV.row_dimensions[rr].height = 28
dv_list(wsV, q(YN), f"F5:F{4 + len(SHOTS)}")
wsV.conditional_formatting.add(f"F5:F{4 + len(SHOTS)}", CellIsRule(operator="equal", formula=['"Yes"'], fill=fill_ok))
wsV.freeze_panes = "A5"

# ------------------------------------------------------------------ Canva Kit
wsK = wb.create_sheet("Canva Kit")
banner(wsK, "CANVA KIT", "Sizes, colours and fonts for CCCL graphics in Canva, and the starter designs in the Marketing folder.", 5)
for c, w in zip("ABCDE", (30, 18, 12, 44, 34)): wsK.column_dimensions[c].width = w
r = 4
section(wsK, r, "COLOURS (from the CCCL logo)", 5); r += 1
for name, hexv, use in (("Near-black", "#151C19", "Backgrounds"), ("Forest green", "#354940", "Panels, headers"), ("Dark green", "#26362F", "Cards, footers"),
                        ("Cream", "#F2ECDC", "Text on dark, badges"), ("Red (maple leaf)", "#AE2221", "One accent per design: the date, VS, a price"),
                        ("Bright red (big text on dark)", "#D9483B", "Headlines on near-black")):
    wsK.cell(r, 1, name).font = f_bold; wsK.cell(r, 2, hexv).font = f_lab; wsK.cell(r, 4, use).font = f_lab
    wsK.cell(r, 3).fill = PatternFill("solid", fgColor=hexv[1:])
    r += 1
r += 1
section(wsK, r, "FONTS (all free in Canva)", 5); r += 1
for name, use in (("Anton", "Headlines: tall, condensed, all caps, like the logo (the starter designs use it)"), ("Oswald (SemiBold)", "Subheads, fighter names, dates"),
                  ("Inter", "Small text: venue, tickets, sponsor lines")):
    wsK.cell(r, 1, name).font = f_bold; wsK.cell(r, 4, use).font = f_lab; r += 1
r += 1
section(wsK, r, "SIZES", 5); r += 1
for i, h in enumerate(("Format", "Size (px)", "Ratio", "Use", "Source"), 1):
    c = wsK.cell(r, i, h); c.font = f_bold; c.fill = fill_pl; c.border = box
r += 1
SIZES = [("Instagram + Facebook feed post (use this)", "1080 x 1350", "4:5", "Most posts. The starter designs use it.", "sproutsocial.com social image sizes guide (Oct 2026)"),
         ("Instagram 3:4 post", "1080 x 1440", "3:4", "Supported since May 2025; the profile grid crops to 3:4, so keep text away from the top and bottom edges", "9to5mac.com, May 29 2025"),
         ("Stories and Reels", "1080 x 1920", "9:16", "Countdowns, show day. Keep text inside the middle 1080 x 1610", "sproutsocial.com"),
         ("Facebook event cover", "1920 x 1005", "about 1.91:1", "The Facebook event. Keep the title and date in the middle 60%", "postfa.st (Meta no longer publishes a spec)"),
         ("Facebook link / landscape", "1200 x 628", "1.91:1", "Shared links, some ads", "blog.hootsuite.com"),
         ("Poster", "11 x 17 in (also letter 8.5 x 11)", "", "Printed posters with the QR code", "")]
for row in SIZES:
    for i, v in enumerate(row, 1):
        c = wsK.cell(r, i, v); c.font = f_lab; c.border = box; c.alignment = wrap
    wsK.row_dimensions[r].height = 28; r += 1
r += 1
section(wsK, r, "STARTER DESIGNS (Marketing / Canva starter designs folder)", 5); r += 1
for name, use in (("01 Tickets on sale (post)", "Weeks 16 and 14"), ("02 Countdown story", "21, 14 and 7 days out"), ("03 Poster, early (11x17)", "Week 10, printed with the QR code"),
                  ("04 Fight card (post)", "Card announced, fight week"), ("05 Fighter vs fighter", "One per adult bout, weeks 4 to 1"), ("06 Fighter spotlight", "Weeks 6 to 2, adults who agree"),
                  ("07 Sponsor thank-you", "Week 8 and after the show"), ("08 Results", "Day after, official only"), ("09 Facebook event cover", "Week 16"),
                  ("10 Poster with main event (11x17)", "Once the main event is signed"), ("11 I'm fighting (post)", "Each adult fighter, week 6"), ("12 I'm fighting (story)", "Each adult fighter, with the link sticker"),
                  ("13 Weigh-in results", "Fight day"), ("14 Win (or Thank you)", "Day after, each adult fighter"), ("15 Proud sponsor", "Each sponsor, week 4"), ("16 Fight card (story)", "Card announced, fight week")):
    wsK.cell(r, 1, name).font = f_bold; wsK.cell(r, 4, use).font = f_lab; r += 1
r += 1
note(wsK, r, 1, "Logo files: Brand/CCCL logo.png and Brand/CCCL logo (transparent).png. Keep the logo round and uncropped; put it on dark backgrounds or the cream badge. "
     "Never put alcohol logos on youth content.", 5)
wsK.row_dimensions[r].height = 30
r += 2
section(wsK, r, "CANVA: GOOD TO KNOW (canva.com help, Oct 2026)", 5); r += 1
for a_, b_ in (("Importing the starter PDFs", "Create a design > Import file > pick the PDF. Every word is its own text box; the background is one image (lock it). If a font gets swapped, re-pick Anton, Oswald or Inter. Photos: delete the grey photo box and drop the photo into a Frame (Elements > Frames)."),
               ("Free vs Pro", "Free Canva works for all of this. Pro (about $142 CAD a year) adds a full Brand Kit with all the CCCL colours, saved brand templates and one-click resizing (post to story)."),
               ("Claude + Canva", "Canva has a connector for Claude that works on free Canva: Claude can make, edit and export designs. Filling a saved template automatically (Autofill) needs Pro."),
               ("Canva for Nonprofits", "Not an option for CCCL: Canva excludes sports leagues and for-profit companies.")):
    wsK.cell(r, 1, a_).font = f_bold; wsK.cell(r, 1).alignment = wrap
    c = wsK.cell(r, 2, b_); c.font = f_lab; c.alignment = wrap
    wsK.merge_cells(start_row=r, start_column=2, end_row=r, end_column=5); wsK.row_dimensions[r].height = 30; r += 1

# ------------------------------------------------------------------ Start Here
ws0 = wb.create_sheet("Start Here", 0)
banner(ws0, "CCCL MARKETING KIT",
       "How each show sells out: the campaign week by week, where tickets come from, whether ads are worth it, posters, the shot list and the Canva kit. "
       "Built for Cody Purtell by Kyle. Opens in Google Sheets or Excel.", 3)
ws0.column_dimensions["A"].width = 4; ws0.column_dimensions["B"].width = 24; ws0.column_dimensions["C"].width = 98
r = 4
def sec0(t):
    global r
    ws0.cell(r, 2, t).font = f_sec; r += 1
def line0(a, b, n=None):
    global r
    if n: ws0.cell(r, 1, n).font = Font(name=F, size=12, bold=True, color=RED)
    ws0.cell(r, 2, a).font = f_bold; ws0.cell(r, 2).alignment = top
    c = ws0.cell(r, 3, b); c.font = f_lab; c.alignment = wrap
    if len(b) > 105: ws0.row_dimensions[r].height = 28
    r += 1
sec0("IN ORDER")
for n, (a, b) in enumerate((
        ("Show Setup", "Show date, ticket link, CCCL's pages, and the numbers that decide if an ad is worth it."),
        ("Campaign", "Every post, poster, press release and push from 16 weeks out to the highlight video after. Dates follow the show date."),
        ("Ticket Sources", "One Eventbrite tracking link or code per source (each fighter, posters, each ad...). Copy the sales in weekly to see what works."),
        ("Ads", "Optional. Test small, keep only ads that sell tickets under the target cost, stop the rest."),
        ("Livestream", "What a paid stream could bring in (Eventbrite online ticket or Dacast), and the other options with their risks."),
        ("Poster Drop", "Where the posters go in the show town, and when they come down."),
        ("Shot List", "Photos and video to get on fight night, and whose OK each one needs."),
        ("Canva Kit", "CCCL colours, fonts and sizes, and the starter designs to import into Canva."),
), 1):
    line0(a, b, str(n))
r += 1
sec0("THE BIG IDEAS")
for a, b in (("Fighters sell tickets", "In 2026 fighters told friends and family informally. Give each one a personal link and a promo pack, and track it: it's usually the biggest source at amateur shows."),
             ("Measure everything", "If a post, poster or ad doesn't have its own link or code, you can't tell if it worked."),
             ("Ads only if they pay", "Every paid ticket is worth about $73 to CCCL. An ad is only worth running if it sells tickets for well under that."),
             ("Real numbers only", "400 tickets and 15 sponsors at the first show are real. Never post made-up crowd or viewer numbers."),
             ("Kids", "No youth athletes in ads. Photos or video of kids only with a parent's written OK. Never name anyone under 18 in a press release.")):
    line0(a, b)
r += 1
sec0("WORKS WITH")
for a, b in (("Fight Night Kit", "Same show date. Press list, fighters and the Results & Report tab."),
             ("Constrictor HQ", "Drafted posts go on the Content Calendar and wait on the Approvals tab for Cody's OK."),
             ("Claude", "The bbg-cccl-social skill drafts the whole campaign, captions and fighter promo packs from this sheet.")):
    line0(a, b)
ws0.sheet_view.showGridLines = False

for ws in wb.worksheets:
    ws.sheet_properties.tabColor = RED if ws.title in ("Start Here", "Campaign") else PINE
    ws.page_setup.orientation = "landscape"; ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
wb.active = 0
wb.save(OUT)
print("saved", OUT)
