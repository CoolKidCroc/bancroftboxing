"""Build 'Second Chance Tracker.xlsx' (its own workbook, kept apart from BBG and CCCL) and the second-chance skill reference.
Google Sheets compatible (COUNTIF/COUNTIFS/SUMPRODUCT; inline dropdowns without commas, under 255 characters).
usage: python3 build_second_chance.py OUT research/second_chance_plan.json"""
import sys, os, json, datetime as dt
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule, CellIsRule
from openpyxl.utils import get_column_letter as L

OUT = sys.argv[1]; D = json.load(open(sys.argv[2]))
# Second Chance has its own colours (deep teal, warm sand) so nobody mistakes it for a gym or CCCL sheet
TEAL, SAND, SAND_L = "12343B", "F0C38E", "F6E9D8"
F = "Arial"
f_title = Font(name=F, size=16, bold=True, color=SAND); f_sub = Font(name=F, size=10, color="FFFFFF")
f_head = Font(name=F, size=10, bold=True, color="FFFFFF"); f_lab = Font(name=F, size=10); f_bold = Font(name=F, size=10, bold=True)
f_in = Font(name=F, size=10, color="0000FF"); f_note = Font(name=F, size=9, italic=True, color="555555")
f_sec = Font(name=F, size=11, bold=True, color=TEAL); f_calc = Font(name=F, size=10, bold=True, color="000000"); f_link = Font(name=F, size=9, color="1F5FA8")
f_warn = Font(name=F, size=10, bold=True, color="9C0006")
fill_teal = PatternFill("solid", fgColor=TEAL); fill_in = PatternFill("solid", fgColor="FFFF00"); fill_sand = PatternFill("solid", fgColor=SAND_L)
fill_warn = PatternFill("solid", fgColor="FFC7CE")
thin = Side(style="thin", color="BFBFBF"); box = Border(top=thin, bottom=thin, left=thin, right=thin)
wrap = Alignment(wrap_text=True, vertical="top"); center = Alignment(horizontal="center", vertical="top", wrap_text=True)
DATE = "mmm d, yyyy"

wb = Workbook()
def banner(ws, title, sub, width):
    for r in (1, 2):
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=width)
        for c in range(1, width + 1): ws.cell(r, c).fill = fill_teal
    ws["A1"] = title; ws["A1"].font = f_title; ws["A1"].alignment = Alignment(vertical="center")
    ws["A2"] = sub; ws["A2"].font = f_sub; ws["A2"].alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[1].height = 30; ws.row_dimensions[2].height = 34
def head(ws, r, cols):
    for i, (t, w) in enumerate(cols, 1):
        c = ws.cell(r, i, t); c.font = f_head; c.fill = fill_teal; c.border = box; c.alignment = Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[L(i)].width = w
    ws.row_dimensions[r].height = 32
def inp(c, v=None, fmt=None):
    c.value = v; c.font = f_in; c.fill = fill_in; c.border = box; c.alignment = wrap
    if fmt: c.number_format = fmt
def txt(c, v, font=None):
    c.value = v; c.font = font or f_lab; c.border = box; c.alignment = wrap
def calc(c, v, fmt=None):
    c.value = v; c.font = f_calc; c.border = box; c.alignment = center
    if fmt: c.number_format = fmt
def dv(ws, items, rng):
    s = '"' + ",".join(items) + '"'; assert len(s) <= 257 and not any("," in i for i in items), items
    d = DataValidation(type="list", formula1=s, allow_blank=True); ws.add_data_validation(d); d.add(rng)
def dv_date(ws, rng):
    d = DataValidation(type="date", allow_blank=True); ws.add_data_validation(d); d.add(rng)

STEP_STATUS = ["Not started", "In progress", "Done", "Stuck"]
YNP = ["Yes", "No", "Pending"]
STAGES = ["Referral", "Welcome", "Training", "Work experience", "Mentoring", "Job connection", "Working", "Paused", "Left"]
CHECK = ["Not needed", "Pending", "Cleared", "Role-limited"]
ROLES = ["Setup and teardown", "Merch table with staff", "Cleanup", "Production runner", "Ticket scanning with staff", "Placement host role", "Training only"]
PARTNER_STATUS = ["Not contacted", "Contacted", "Meeting booked", "Partner", "Not now"]
FUND_STATUS = ["Not started", "Researching", "Talking to a partner", "Drafting", "Submitted", "Approved", "Declined", "Not yet eligible"]

# ---------------- Start Here
ws = wb.active; ws.title = "Start Here"
banner(ws, "SECOND CHANCE TRACKER", "Setup, board, partners, placements, participants and funding for Second Chance. Its own workbook, in its own Drive folder: only Cody, the board, and Claude (through Cody's CCCL account) can open it.", 3)
ws.column_dimensions["A"].width = 4; ws.column_dimensions["B"].width = 24; ws.column_dimensions["C"].width = 96
r = 4
ws.cell(r, 2, "PRIVACY FIRST").font = f_warn; ws.cell(r, 2).fill = fill_warn
c = ws.cell(r, 3, "Never write down what anyone was charged with or convicted of, health details or addresses. Participants appear by first name and a code (SC-001) only. "
                  "This workbook never goes in the BBG or CCCL folders and is never shared with gym staff."); c.font = f_bold; c.fill = fill_warn; c.alignment = wrap
ws.row_dimensions[r].height = 46; r += 2
for n, a, b in [("1", "Setup", "The first 90 days, in order. Mark each step Done as you go; the count at the top says how far along Second Chance is."),
                ("2", "Board", "The directors. CRA needs more than half at arm's length (not related to each other or to Cody); the checks at the top say if the board is ready."),
                ("3", "Partners", "Organizations near Bancroft that could refer people, partner on a grant or lead an application while Second Chance gets established."),
                ("4", "Placements", "Where participants get work experience. Several hosts, never CCCL alone. Note whether young people are present."),
                ("5", "Participants", "First name and code only. Stage, consent, check status, roles allowed, mentor and the next check-in. Overdue check-ins turn red."),
                ("6", "Funding", "Grants that fit Second Chance and what each needs (incorporation, charity status or a partner)."),
                ("7", "Money", "Second Chance's money goes in its own bank account and its own copy of the Money Summary, never BBG's or CCCL's.")]:
    ws.cell(r, 1, n).font = Font(name=F, size=12, bold=True, color=TEAL); ws.cell(r, 2, a).font = f_bold
    c = ws.cell(r, 3, b); c.font = f_lab; c.alignment = wrap; ws.row_dimensions[r].height = 30; r += 1
r += 1
ws.cell(r, 2, "About Second Chance").font = f_sec; r += 1
for a, b in (("Mission", D["about"]["mission"]), ("For", D["about"]["who"]), ("What it does", D["about"]["does"]), ("Where it stands", D["about"]["status"] + " (" + D["about"]["source"] + ")")):
    ws.cell(r, 2, a).font = f_bold; c = ws.cell(r, 3, b); c.font = f_lab; c.alignment = wrap; ws.row_dimensions[r].height = 30; r += 1
r += 1
ws.cell(r, 2, "Colours").font = f_sec; r += 1
c = ws.cell(r, 2, "Yellow, blue text"); c.font = f_in; c.fill = fill_in; ws.cell(r, 3, "You type or pick here").font = f_lab; r += 1
ws.cell(r, 2, "Black, bold").font = f_calc; ws.cell(r, 3, "Calculated: don't type over it").font = f_lab; r += 2
c = ws.cell(r, 2, "Research, not legal or accounting advice. A lawyer reviews the purposes, by-law and policies; an accountant sets up the books. See the Second Chance Starter Plan."); c.font = f_note
ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=3)

# ---------------- Setup
wsS = wb.create_sheet("Setup")
SC = [("Area", 11), ("Step", 70), ("Who", 13), ("When", 18), ("Status", 13), ("Done on", 12), ("Notes", 36)]
banner(wsS, "SETUP: THE FIRST 90 DAYS", "In order. Mark each step Done. The research behind each step is in the Second Chance Starter Plan.", len(SC))
HR = 5; SMAX = HR + len(D["checklist"]) + 15
wsS.cell(4, 1, "Done").font = f_bold
calc(wsS.cell(4, 2), f'=COUNTIF(E{HR+1}:E{SMAX},"Done")&" of "&COUNTA(B{HR+1}:B{SMAX})&" steps done"')
wsS.cell(4, 2).alignment = Alignment(horizontal="left")
head(wsS, HR, SC)
for k in range(SMAX - HR):
    rr = HR + 1 + k
    row = D["checklist"][k] if k < len(D["checklist"]) else None
    if row:
        for i, v in enumerate(row, 1): txt(wsS.cell(rr, i), v, f_bold if i == 1 else None)
        wsS.row_dimensions[rr].height = 44
    else:
        for i in range(1, 5): inp(wsS.cell(rr, i))
    inp(wsS.cell(rr, 5), "Not started" if row else None); inp(wsS.cell(rr, 6), None, DATE); inp(wsS.cell(rr, 7))
dv(wsS, STEP_STATUS, f"E{HR+1}:E{SMAX}"); dv_date(wsS, f"F{HR+1}:F{SMAX}")
wsS.conditional_formatting.add(f"E{HR+1}:E{SMAX}", CellIsRule(operator="equal", formula=['"Done"'], fill=PatternFill("solid", fgColor="C6EFCE"), font=Font(name=F, color="006100", bold=True)))
wsS.conditional_formatting.add(f"E{HR+1}:E{SMAX}", CellIsRule(operator="equal", formula=['"Stuck"'], fill=fill_warn, font=Font(name=F, color="9C0006", bold=True)))
wsS.freeze_panes = f"C{HR+1}"

# ---------------- Board
wsB = wb.create_sheet("Board")
BC = [("First name", 26), ("Role", 14), ("Agreed to serve (in writing)", 14), ("At arm's length?", 13), ("Vulnerable sector check", 14), ("Conflict policy signed", 13), ("Notes", 50)]
banner(wsB, "THE BOARD", "At least 3 directors; more than half at arm's length (not related to each other or to Cody by blood, marriage, common-law partnership or adoption); no more than one-third employees.", len(BC))
HB = 7; BMAX = HB + 12
labels = [("Directors who've agreed", f'=COUNTIF(C{HB+1}:C{BMAX},"Yes")'),
          ("At arm's length", f'=COUNTIF(D{HB+1}:D{BMAX},"Yes")'),
          ("Board ready?", f'=IF(AND(COUNTIF(C{HB+1}:C{BMAX},"Yes")>=3,COUNTIFS(C{HB+1}:C{BMAX},"Yes",D{HB+1}:D{BMAX},"Yes")*2>COUNTIF(C{HB+1}:C{BMAX},"Yes")),"Yes: 3+ directors, more than half at arm\'s length","Not yet")')]
for i, (a, f) in enumerate(labels):
    wsB.cell(4 + i, 1, a).font = f_bold; calc(wsB.cell(4 + i, 2), f); wsB.cell(4 + i, 2).alignment = Alignment(horizontal="left")
    wsB.merge_cells(start_row=4 + i, start_column=2, end_row=4 + i, end_column=4)
head(wsB, HB, BC)
for k in range(BMAX - HB):
    rr = HB + 1 + k
    name = D["board"][k] if k < len(D["board"]) else None
    inp(wsB.cell(rr, 1), name); inp(wsB.cell(rr, 2), "Director" if name else None)
    for c in range(3, 7): inp(wsB.cell(rr, c), "Pending" if name else None)
    inp(wsB.cell(rr, 7), "Named by Cody (Oct 20, 2026)" if name else None)
dv(wsB, ["Chair", "Treasurer", "Secretary", "Director"], f"B{HB+1}:B{BMAX}")
for col in "CEF": dv(wsB, YNP, f"{col}{HB+1}:{col}{BMAX}")
dv(wsB, ["Yes", "No", "Check"], f"D{HB+1}:D{BMAX}")
wsB.conditional_formatting.add("B6", FormulaRule(formula=['LEFT(B6,3)="Yes"'], fill=PatternFill("solid", fgColor="C6EFCE"), font=Font(name=F, color="006100", bold=True)))
wsB.conditional_formatting.add("B6", FormulaRule(formula=['B6="Not yet"'], fill=fill_warn, font=Font(name=F, color="9C0006", bold=True)))
wsB.freeze_panes = f"B{HB+1}"

# ---------------- Partners
wsP = wb.create_sheet("Partners")
PC = [("Organization", 32), ("What they do", 50), ("Where", 16), ("Website", 30), ("Status", 15), ("Last contact", 12), ("Next step", 30), ("Notes", 30)]
banner(wsP, "PARTNERS", "Referral partners, co-applicants, or a registered charity that could lead a grant while Second Chance gets established. Public websites only; log contact names here, not personal numbers.", len(PC))
HP = 4; PMAX = HP + len(D["partners"]) + 15
head(wsP, HP, PC)
for k in range(PMAX - HP):
    rr = HP + 1 + k
    p = D["partners"][k] if k < len(D["partners"]) else None
    if p:
        txt(wsP.cell(rr, 1), p[0], f_bold); txt(wsP.cell(rr, 2), p[1]); txt(wsP.cell(rr, 3), p[2])
        c = wsP.cell(rr, 4, p[3]); c.hyperlink = p[3]; c.font = f_link; c.border = box; c.alignment = wrap
        wsP.row_dimensions[rr].height = 44
    else:
        for i in range(1, 5): inp(wsP.cell(rr, i))
    inp(wsP.cell(rr, 5), "Not contacted" if p else None); inp(wsP.cell(rr, 6), None, DATE); inp(wsP.cell(rr, 7)); inp(wsP.cell(rr, 8))
dv(wsP, PARTNER_STATUS, f"E{HP+1}:E{PMAX}"); dv_date(wsP, f"F{HP+1}:F{PMAX}")
wsP.freeze_panes = f"B{HP+1}"

# ---------------- Placements
wsL = wb.create_sheet("Placements")
LC = [("Host", 30), ("Type", 14), ("Roles", 46), ("Young people present?", 12), ("Agreed?", 11), ("Contact (first name)", 14), ("Notes", 46)]
banner(wsL, "PLACEMENTS", "Work experience with several hosts. CCCL fight nights are one placement among several, never the only one (CRA: training, not jobs, and not for one employer).", len(LC))
HL = 4; LMAX = HL + 30
head(wsL, HL, LC)
good = dict(D["fight_night_roles"])
seed = [("CCCL fight nights", "Fight night", good["Good fits"], "Yes", "Yes", "Cody",
         "Youth grapplers are at the shows: no roles with authority over young athletes, never alone with one, Rule of Two. Cody declares his interest in CCCL.")]
for k in range(LMAX - HL):
    rr = HL + 1 + k
    s = seed[k] if k < len(seed) else None
    for i in range(7): inp(wsL.cell(rr, i + 1), s[i] if s else None)
    if s: wsL.row_dimensions[rr].height = 58
dv(wsL, ["Business", "Fight night", "Partner", "Training provider", "Other"], f"B{HL+1}:B{LMAX}")
dv(wsL, ["Yes", "No"], f"D{HL+1}:D{LMAX}"); dv(wsL, ["Yes", "Asked", "No"], f"E{HL+1}:E{LMAX}")
wsL.cell(LMAX + 2, 1, "Hosts to ask (none confirmed): local businesses that know Cody, such as the 2026 CCCL sponsors, plus the partners. Ask first; add them here when they say yes.").font = f_note

# ---------------- Participants
wsQ = wb.create_sheet("Participants")
QC = [("Code", 9), ("First name", 13), ("Referred by", 20), ("Start date", 12), ("Stage", 15), ("Written consent on file?", 12), ("Police check", 13), ("Check date", 12),
      ("Roles allowed", 24), ("Placement", 22), ("Mentor (first name)", 13), ("Last check-in", 12), ("Next check-in", 12), ("Due?", 11), ("Notes (no offences, health or addresses)", 40)]
banner(wsQ, "PARTICIPANTS", "First name and code only. NEVER write charges, convictions, health details or addresses. Police checks: record only the date and Cleared / Role-limited; the result belongs to the person.", len(QC))
HQ = 6; QMAX = HQ + 120
wsQ.cell(4, 1, "Active").font = f_bold
calc(wsQ.cell(4, 2), f'=SUMPRODUCT((B{HQ+1}:B{QMAX}<>"")*(E{HQ+1}:E{QMAX}<>"Left")*(E{HQ+1}:E{QMAX}<>"Paused")*(LEFT(B{HQ+1}:B{QMAX},9)<>"(Example)"))')
wsQ.cell(4, 4, "Check-ins due in 7 days or overdue").font = f_bold; wsQ.merge_cells(start_row=4, start_column=4, end_row=4, end_column=6)
calc(wsQ.cell(4, 7), f'=COUNTIF(N{HQ+1}:N{QMAX},"Overdue")+COUNTIF(N{HQ+1}:N{QMAX},"This week")')
head(wsQ, HQ, QC)
for k in range(QMAX - HQ):
    rr = HQ + 1 + k
    ex = k == 0
    vals = ["SC-000", "(Example) Alex", "John Howard Society", dt.date(2026, 11, 2), "Welcome", "Yes", "Cleared", dt.date(2026, 10, 28), "Setup and teardown", "CCCL fight nights", "(Example) Sam",
            dt.date(2026, 11, 2), dt.date(2026, 12, 2)] if ex else [None] * 13
    for i, v in enumerate(vals, 1):
        inp(wsQ.cell(rr, i), v, DATE if i in (4, 8, 12, 13) else None)
    calc(wsQ.cell(rr, 14), f'=IF(OR(B{rr}="",M{rr}="",E{rr}="Left",E{rr}="Paused"),"",IF(M{rr}<TODAY(),"Overdue",IF(M{rr}-TODAY()<=7,"This week","")))')
    inp(wsQ.cell(rr, 15), "Example row: delete it once the first real participant is in." if ex else None)
dv(wsQ, STAGES, f"E{HQ+1}:E{QMAX}"); dv(wsQ, ["Yes", "No"], f"F{HQ+1}:F{QMAX}"); dv(wsQ, CHECK, f"G{HQ+1}:G{QMAX}"); dv(wsQ, ROLES, f"I{HQ+1}:I{QMAX}")
for col in "DHLM": dv_date(wsQ, f"{col}{HQ+1}:{col}{QMAX}")
wsQ.conditional_formatting.add(f"N{HQ+1}:N{QMAX}", CellIsRule(operator="equal", formula=['"Overdue"'], fill=fill_warn, font=Font(name=F, color="9C0006", bold=True)))
wsQ.conditional_formatting.add(f"N{HQ+1}:N{QMAX}", CellIsRule(operator="equal", formula=['"This week"'], fill=PatternFill("solid", fgColor="FFEB9C"), font=Font(name=F, color="9C5700", bold=True)))
wsQ.conditional_formatting.add(f"F{HQ+1}:F{QMAX}", CellIsRule(operator="equal", formula=['"No"'], fill=fill_warn, font=Font(name=F, color="9C0006", bold=True)))
wsQ.freeze_panes = f"C{HQ+1}"

# ---------------- Funding
wsF = wb.create_sheet("Funding")
FC = [("Program", 32), ("What and when", 64), ("Needs", 22), ("Status", 18), ("Next step", 34), ("Notes", 30)]
banner(wsF, "FUNDING", "Programs that fit Second Chance. Most need incorporation first; some need charity status, or a registered charity partner to lead. Full list: The Constrictor Funding Tracker.", len(FC))
HF = 4; FMAX = HF + len(D["funding"]) + 10
head(wsF, HF, FC)
for k in range(FMAX - HF):
    rr = HF + 1 + k
    f = D["funding"][k] if k < len(D["funding"]) else None
    if f:
        txt(wsF.cell(rr, 1), f[0], f_bold); txt(wsF.cell(rr, 2), f[1]); txt(wsF.cell(rr, 3), f[2]); wsF.row_dimensions[rr].height = 44
    else:
        for i in range(1, 4): inp(wsF.cell(rr, i))
    inp(wsF.cell(rr, 4), ("Not yet eligible" if f and "Incorporated" in f[2] or (f and "year" in f[2]) else ("Not started" if f else None)))
    inp(wsF.cell(rr, 5)); inp(wsF.cell(rr, 6))
dv(wsF, FUND_STATUS, f"D{HF+1}:D{FMAX}")
wsF.freeze_panes = f"B{HF+1}"

for w in wb.worksheets:
    w.page_setup.orientation = "landscape"; w.sheet_properties.pageSetUpPr.fitToPage = True; w.page_setup.fitToWidth = 1; w.page_setup.fitToHeight = 0
wb.save(OUT); print("saved", OUT)

# ---------------- skill reference
Dd = os.path.join(os.path.dirname(os.path.abspath(__file__)), "skills/second-chance/references"); os.makedirs(Dd, exist_ok=True)
M = ["# Second Chance: the plan", "", f"From the Second Chance Starter Plan and the research behind it ({D['generated']}). Research, not legal or accounting advice.", "",
     "## About", "", f"- **Mission**: {D['about']['mission']}", f"- **For**: {D['about']['who']}", f"- **What it does**: {D['about']['does']}",
     f"- **Where it stands**: {D['about']['status']} ({D['about']['source']})", f"- **Board (first names)**: {', '.join(D['board'])}", "",
     "## Setup steps (the Setup tab)", "", "| Area | Step | Who | When |", "|---|---|---|---|"] + [f"| {a} | {b} | {c} | {d} |" for a, b, c, d in D["checklist"]] + [
     "", "## The pathway", ""] + [f"- **{a}**: {b}" for a, b in D["pathway"]] + [
     "", "## Fight-night roles", ""] + [f"- **{a}**: {b}" for a, b in D["fight_night_roles"]] + [
     "", "## Privacy", ""] + [f"- {x}" for x in D["privacy"]] + [
     "", "## Charity-status cautions", ""] + [f"- {x}" for x in D["charity_cautions"]] + [
     "", "## Partners", ""] + [f"- **{a}** ({c}): {b} {d}" for a, b, c, d in D["partners"]] + [
     "", "## Funding", ""] + [f"- **{a}** (needs: {c}): {b}" for a, b, c in D["funding"]] + [
     "", "## First-year costs", ""] + [f"- {a}: {b}" for a, b in D["budget"]] + [
     "", "## Sources", ""] + [f"- {a}: {b}" for a, b in D["sources"]] + [""]
open(os.path.join(Dd, "plan.md"), "w").write("\n".join(M)); print("wrote", os.path.join(Dd, "plan.md"))
