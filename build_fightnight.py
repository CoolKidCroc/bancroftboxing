"""Build 'CCCL Fight Night Kit.xlsx' (Google Sheets compatible: no post-2007 functions).
usage: python3 build_fightnight.py OUT [research/oac_fight_night.json] [research/media_gyms.json]"""
import sys, json, os, datetime as dt
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.utils import get_column_letter as L

OUT = sys.argv[1]
OAC = json.load(open(sys.argv[2])) if len(sys.argv) > 2 and os.path.exists(sys.argv[2]) else None
MG = json.load(open(sys.argv[3])) if len(sys.argv) > 3 and os.path.exists(sys.argv[3]) else None

# CCCL palette (from the logo)
INK, PINE, PINE_L, CREAM, RED = "151C19", "354940", "E6EEE9", "F2ECDC", "AE2221"
F = "Arial"
f_title = Font(name=F, size=16, bold=True, color=CREAM)
f_sub = Font(name=F, size=10, color="D9D5C8")
f_head = Font(name=F, size=10, bold=True, color="FFFFFF")
f_sec = Font(name=F, size=11, bold=True, color=PINE)
f_lab = Font(name=F, size=10)
f_bold = Font(name=F, size=10, bold=True)
f_in = Font(name=F, size=10, color="0000FF")
f_calc = Font(name=F, size=10, color="000000")
f_link = Font(name=F, size=10, color="008000")
f_note = Font(name=F, size=9, italic=True, color="555555")
f_flag = Font(name=F, size=10, bold=True, color="C00000")
fill_ink = PatternFill("solid", fgColor=INK)
fill_pine = PatternFill("solid", fgColor=PINE)
fill_pl = PatternFill("solid", fgColor=PINE_L)
fill_in = PatternFill("solid", fgColor="FFFF00")
fill_ok = PatternFill("solid", fgColor="C6EFCE")
fill_bad = PatternFill("solid", fgColor="F8CBAD")
fill_warn = PatternFill("solid", fgColor="FFEB9C")
thin = Side(style="thin", color="BFBFBF")
box = Border(top=thin, bottom=thin, left=thin, right=thin)
wrap = Alignment(wrap_text=True, vertical="top")
top = Alignment(vertical="top")
USD = '$#,##0;($#,##0);"-"'
NUM = '#,##0;(#,##0);"-"'
DATE = 'ddd mmm d, yyyy'
TIME = 'h:mm AM/PM'

wb = Workbook()

def banner(ws, title, sub, width):
    for r in (1, 2):
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=width)
        for c in range(1, width + 1):
            ws.cell(r, c).fill = fill_ink
    ws["A1"] = title; ws["A1"].font = f_title; ws["A1"].alignment = Alignment(vertical="center")
    ws["A2"] = sub; ws["A2"].font = f_sub; ws["A2"].alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[1].height = 30; ws.row_dimensions[2].height = 34

def head(ws, r, labels, widths=None, height=34):
    for i, h in enumerate(labels, 1):
        c = ws.cell(r, i, h); c.font = f_head; c.fill = fill_pine; c.border = box
        c.alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[r].height = height
    if widths:
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[L(i)].width = w

def section(ws, r, text, width):
    ws.cell(r, 1, text).font = f_head
    for c in range(1, width + 1):
        ws.cell(r, c).fill = fill_pine
    ws.row_dimensions[r].height = 20

def inp(c, v=None, fmt=None):
    if v is not None: c.value = v
    c.font = f_in; c.fill = fill_in; c.border = box
    if fmt: c.number_format = fmt

def note(ws, r, c, text, merge_to=None):
    x = ws.cell(r, c, text); x.font = f_note; x.alignment = wrap
    if merge_to: ws.merge_cells(start_row=r, start_column=c, end_row=r, end_column=merge_to)

def dv_list(ws, items_or_ref, rng):
    f1 = items_or_ref if items_or_ref.startswith("=") else items_or_ref
    dv = DataValidation(type="list", formula1=f1, allow_blank=True); ws.add_data_validation(dv); dv.add(rng)

def q(items): return '"' + ",".join(items) + '"'

YN = ["Yes", "No"]
YNP = ["Yes", "No", "Pending"]
YNA = ["Yes", "No", "N/A"]

# ------------------------------------------------------------------ Lists (helper tab)
wsL = wb.active; wsL.title = "Lists"
lists = {
    "Discipline": ["Amateur MMA", "Youth grappling", "Adult grappling"],
    "Age division": ["Senior (19+)", "Youth A (16-18)", "Youth B (14-15)", "Youth C (12-13)", "Under 12 (grappling only)"],
    "Grappling level": ["Beginner (under 1 year)", "Intermediate (1-3 years)", "Advanced (3+ years)"],
    "Status": ["Applied", "Contacted", "Available", "Matched", "Confirmed", "Withdrawn"],
    "Card position": ["Main event", "Co-main", "Main card", "Prelim", "Youth"],
    "Rules": ["Amateur MMA (Senior)", "Amateur MMA (Youth)", "Youth grappling", "Adult grappling"],
    "Gym stage": ["Not contacted", "Contacted", "Interested", "Sending fighters", "Not this time"],
    "Method": ["Decision (unanimous)", "Decision (split)", "Decision (majority)", "KO", "TKO", "Submission", "Doctor stoppage",
               "Corner stoppage", "Disqualification", "Draw", "No contest"],
}
LREF = {}
for j, (k, items) in enumerate(lists.items(), 1):
    c = wsL.cell(1, j, k); c.font = f_bold
    for i, it in enumerate(items, 2):
        wsL.cell(i, j, it)
    LREF[k] = f"=Lists!${L(j)}$2:${L(j)}${len(items) + 1}"
    wsL.column_dimensions[L(j)].width = 26
# Ontario single-bout weight-difference table (amateur MMA club shows)
c0 = len(lists) + 2
WD = [(0, 4), (148, 9), (165.01, 11), (200.01, 24), (265.01, 999)]
if OAC and OAC.get("weight_difference_single_bout"):
    src_note = "Ontario amateur MMA ruleset v2.0: maximum weight difference for single bouts. Uses the lighter fighter's weight; at exactly 165 and 200 lb the published bands overlap, so the smaller allowance is used (ask the Commission)."
else:
    src_note = "Weight-difference table (check against the Ontario amateur MMA rules)."
wsL.cell(1, c0, "From (lb, lighter fighter)").font = f_bold; wsL.cell(1, c0 + 1, "Max difference (lb)").font = f_bold
for i, (lo, mx) in enumerate(WD, 2):
    wsL.cell(i, c0, lo); wsL.cell(i, c0 + 1, mx)
WD_LO = f"Lists!${L(c0)}$2:${L(c0)}${len(WD) + 1}"; WD_MX = f"Lists!${L(c0 + 1)}$2:${L(c0 + 1)}${len(WD) + 1}"
wsL.cell(len(WD) + 3, c0, src_note).font = f_note
wsL.cell(len(WD) + 4, c0, "999 = no limit (over 265 lb).").font = f_note
wsL.column_dimensions[L(c0)].width = 24; wsL.column_dimensions[L(c0 + 1)].width = 20

# ------------------------------------------------------------------ Show Setup
wsS = wb.create_sheet("Show Setup", 0)
banner(wsS, "SHOW SETUP", "Fill this in first. Every other tab uses these answers: the countdown dates, the run of show times and the bout checks.", 4)
for c, w in zip("ABCD", (44, 30, 4, 70)): wsS.column_dimensions[c].width = w
S = {}
r = 4
section(wsS, r, "THE SHOW", 4); r += 1
rows = [
    ("name", "Show name", "CCCL 2", None, "Whatever goes on the poster."),
    ("town", "Town", None, None, "Fill in once the show is booked (no 2027 show is booked yet)."),
    ("venue", "Venue", None, None, ""),
    ("date", "Show date", None, DATE, "Fill in once the show is booked: the Countdown and Run of Show follow it. Until then they stay blank."),
    ("load", "Load-in and cage build start (show day)", dt.time(10, 0), TIME, "The cage has to sit on a platform 3-4 ft off the floor (Ontario rule)."),
    ("weigh", "Weigh-ins start (show day)", dt.time(16, 15), TIME, "Ontario rule: on the day, within 3 hours of the first bout. No day-before weigh-ins."),
    ("doors", "Doors open", dt.time(18, 0), TIME, ""),
    ("first", "First bout", dt.time(19, 0), TIME, ""),
]
for key, lab, val, fmt, nt in rows:
    wsS.cell(r, 1, lab).font = f_lab
    inp(wsS.cell(r, 2), val, fmt)
    if nt: note(wsS, r, 4, nt)
    S[key] = f"'Show Setup'!$B${r}"; r += 1
wsS.cell(r, 1, "Weigh-in timing check").font = f_bold
c = wsS.cell(r, 2, f'=IF(AND({S["first"]}-{S["weigh"]}<=3/24+0.00001,{S["first"]}>{S["weigh"]}),"OK","TOO EARLY: within 3 hours of the first bout")')
c.font = f_bold; c.border = box
wsS.conditional_formatting.add(f"B{r}", CellIsRule(operator="equal", formula=['"OK"'], fill=fill_ok))
wsS.conditional_formatting.add(f"B{r}", CellIsRule(operator="notEqual", formula=['"OK"'], fill=fill_bad))
r += 2
section(wsS, r, "BOUT TIMING (for the Run of Show)", 4); r += 1
for key, lab, val, nt in (
    ("mma_min", "Minutes per amateur MMA bout (walkouts, rounds, breaks, decision)", 20, "ESTIMATE. Senior bouts are 3 x 3-minute rounds with 1-minute breaks; youth bouts are shorter."),
    ("grap_min", "Minutes per grappling match", 10, "ESTIMATE."),
    ("inter_after", "Intermission after bout number", 5, "Sponsor reads and the merch push happen here."),
    ("inter_min", "Intermission length (minutes)", 20, ""),
):
    wsS.cell(r, 1, lab).font = f_lab; wsS.cell(r, 1).alignment = wrap
    inp(wsS.cell(r, 2), val, "0")
    note(wsS, r, 4, nt); S[key] = f"'Show Setup'!$B${r}"; r += 1
r += 1
section(wsS, r, "MATCHMAKING CHECKS (flags only: Cody and the officials decide every bout)", 4); r += 1
wsS.cell(r, 1, "Amateur MMA weight gap").font = f_lab
wsS.cell(r, 2, "Ontario table (built in)").font = f_calc
note(wsS, r, 4, "Max difference by weight: 4 lb under 148, 9 lb at 148-165, 11 lb at 165-200, 24 lb at 200-265, no limit over 265. See the Lists and Rules tabs.")
r += 1
wsS.cell(r, 1, "Grappling: flag a weight gap over (lb)").font = f_lab
inp(wsS.cell(r, 2), 10, "0"); S["ggap"] = f"'Show Setup'!$B${r}"
note(wsS, r, 4, "STARTING POINT. Use the grappling ruleset's weight classes for the age group (Rules tab).")
r += 1
for rr in range(5, r):
    wsS.row_dimensions[rr].height = 24
    for cc in (1, 2, 4):
        x = wsS.cell(rr, cc); x.alignment = Alignment(wrap_text=True, vertical="center", horizontal=x.alignment.horizontal)
wsS.freeze_panes = "A4"

# ------------------------------------------------------------------ Fighters
wsF = wb.create_sheet("Fighters")
FC = [("Fighter (name on the card)", 24), ("Under 18?", 8), ("Gym / team", 22), ("Coach (first name)", 13), ("Contact via (gym email or page)", 24),
      ("Discipline", 15), ("Age division", 15), ("Fight weight (lb)", 9), ("MMA wins", 7), ("MMA losses", 7), ("Other amateur contests (any combat sport)", 11),
      ("Contests", 8), ("Experience class", 11), ("Grappling level", 15),
      ("Annual medical done", 10), ("Codes & rules declaration signed", 11), ("CCCL waiver signed", 9), ("Parent consent (under 18)", 10),
      ("Concussion confirmation (Rowan's Law)", 12), ("Status", 12), ("Show", 14), ("Ready to fight?", 13), ("Notes", 30)]
banner(wsF, "FIGHTERS",
       "One row per fighter who applies. Medical: Yes/No only, never details. Under 18: first name and last initial, a parent's consent, and the parent signs the day-of forms. "
       "No personal phone numbers: contact fighters through their gym or coach. Anyone with pro fights or a prize worth $300+ isn't an amateur.", len(FC))
HR = 4
head(wsF, HR, [c[0] for c in FC], [c[1] for c in FC], 56)
FCOL = {c[0]: L(i + 1) for i, c in enumerate(FC)}
MAXF = HR + 200
calc_cols = {"Contests", "Experience class", "Ready to fight?"}
for rr in range(HR + 1, MAXF + 1):
    for i, (h, _) in enumerate(FC, 1):
        c = wsF.cell(rr, i); c.border = box
        if h in calc_cols:
            c.font = f_calc
        else:
            c.font = f_in; c.fill = fill_in
    A = f"A{rr}"
    W_, L_, O_ = (f"{FCOL[k]}{rr}" for k in ("MMA wins", "MMA losses", "Other amateur contests (any combat sport)"))
    wsF[f"{FCOL['Contests']}{rr}"] = f'=IF({A}="","",N({W_})+N({L_})+N({O_}))'
    n = f"{FCOL['Contests']}{rr}"
    wsF[f"{FCOL['Experience class']}{rr}"] = f'=IF({A}="","",IF({n}<=5,"Pre-Novice",IF({n}<=10,"Novice","Open")))'
    med, cod, wai, par, u18, ro, dis = (f"{FCOL[k]}{rr}" for k in ("Annual medical done", "Codes & rules declaration signed", "CCCL waiver signed",
                                        "Parent consent (under 18)", "Under 18?", "Concussion confirmation (Rowan's Law)", "Discipline"))
    missing = (f'(({med}<>"Yes")+({cod}<>"Yes")+({wai}<>"Yes")+AND({u18}="Yes",{par}<>"Yes")'
               f'+AND(OR({u18}="Yes",ISNUMBER(SEARCH("grappling",{dis}))),{ro}<>"Yes"))')
    wsF[f"{FCOL['Ready to fight?']}{rr}"] = f'=IF({A}="","",IF({missing}=0,"READY",{missing}&" to do"))'
FR = lambda h: f"Fighters!${FCOL[h]}${HR + 1}:${FCOL[h]}${MAXF}"
rng = lambda h: f"{FCOL[h]}{HR + 1}:{FCOL[h]}{MAXF}"
dv_list(wsF, q(YN), rng("Under 18?"))
dv_list(wsF, LREF["Discipline"], rng("Discipline"))
dv_list(wsF, LREF["Age division"], rng("Age division"))
dv_list(wsF, LREF["Grappling level"], rng("Grappling level"))
dv_list(wsF, q(YNP), rng("Annual medical done"))
dv_list(wsF, q(YN), rng("Codes & rules declaration signed"))
dv_list(wsF, q(YN), rng("CCCL waiver signed"))
dv_list(wsF, q(YNA), rng("Parent consent (under 18)"))
dv_list(wsF, q(YNA), rng("Concussion confirmation (Rowan's Law)"))
dv_list(wsF, LREF["Status"], rng("Status"))
wsF.conditional_formatting.add(rng("Ready to fight?"), CellIsRule(operator="equal", formula=['"READY"'], fill=fill_ok))
wsF.conditional_formatting.add(rng("Ready to fight?"), FormulaRule(formula=[f'RIGHT({FCOL["Ready to fight?"]}{HR+1},5)="to do"'], fill=fill_warn))
ex = ["(Example) Jordan Smith", "No", "Big Country Lindsay", "Mike", "Through the gym's page", "Amateur MMA", "Senior (19+)", 154, 2, 1, 2, None, None, "",
      "Pending", "No", "Yes", "N/A", "N/A", "Available", "", None, "Example row. Replace it with a real fighter."]
for i, v in enumerate(ex, 1):
    if v is not None: wsF.cell(HR + 1, i, v)
wsF.freeze_panes = wsF.cell(HR + 1, 2)
wsF.auto_filter.ref = f"A{HR}:{L(len(FC))}{MAXF}"

# ------------------------------------------------------------------ Bout Card
wsB = wb.create_sheet("Bout Card")
BC = [("Bout #", 6), ("Card position", 12), ("Rules", 17), ("Red corner", 22), ("Blue corner", 22),
      ("Red weight", 8), ("Blue weight", 8), ("Weight gap (lb)", 8), ("Allowed gap (lb)", 8), ("Red class", 10), ("Blue class", 10),
      ("Both ready?", 9), ("Checks", 30), ("Cody approved", 9), ("Officials OK", 9),
      ("Winner", 18), ("Method", 16), ("Round", 6), ("Time", 7), ("Notes", 26)]
banner(wsB, "BOUT CARD",
       "Pick both corners from the Fighters list, in the order the bouts will happen (main event last). It checks the Ontario weight-gap table, "
       "experience classes, age divisions, double bookings and paperwork. Flags are a prompt to look again: Cody and the officials decide. Results only once official.", len(BC))
BH = 7
sumr = 4
labels = [("Bouts on the card", f'=COUNTIF(D{BH+1}:D{BH+30},"?*")'),
          ("Approved by Cody and the officials", f'=COUNTIFS(D{BH+1}:D{BH+30},"?*",N{BH+1}:N{BH+30},"Yes",O{BH+1}:O{BH+30},"Yes")'),
          ("Bouts with a flag to look at", f'=COUNTIF(M{BH+1}:M{BH+30},"?*")-COUNTIF(M{BH+1}:M{BH+30},"OK")')]
for i, (lab, fml) in enumerate(labels):
    col = (1, 6, 13)[i]
    wsB.cell(sumr, col, lab).font = f_bold
    wsB.merge_cells(start_row=sumr, start_column=col, end_row=sumr, end_column=col + 4)
    c = wsB.cell(sumr + 1, col, fml); c.font = Font(name=F, size=14, bold=True, color=PINE); c.alignment = Alignment(horizontal="left")
head(wsB, BH, [c[0] for c in BC], [c[1] for c in BC], 44)
BCOL = {c[0]: L(i + 1) for i, c in enumerate(BC)}
FA = FR("Fighter (name on the card)")
def look(name_cell, col):
    return f'IFERROR(INDEX({FR(col)},MATCH({name_cell},{FA},0)),"")'
for k, rr in enumerate(range(BH + 1, BH + 31), 1):
    for i, (h, _) in enumerate(BC, 1):
        c = wsB.cell(rr, i); c.border = box
        if h in ("Card position", "Rules", "Red corner", "Blue corner", "Cody approved", "Officials OK", "Winner", "Method", "Round", "Time", "Notes"):
            c.font = f_in; c.fill = fill_in
        else:
            c.font = f_calc
    wsB.cell(rr, 1, k)
    R_, B_ = f"D{rr}", f"E{rr}"
    blank = f'OR({R_}="",{B_}="")'
    mma = f'LEFT(C{rr},11)="Amateur MMA"'
    wsB[f"F{rr}"] = f'=IF({R_}="","",{look(R_, "Fight weight (lb)")})'
    wsB[f"G{rr}"] = f'=IF({B_}="","",{look(B_, "Fight weight (lb)")})'
    wsB[f"H{rr}"] = f'=IF(OR(F{rr}="",G{rr}=""),"",ABS(F{rr}-G{rr}))'
    wsB[f"I{rr}"] = (f'=IF(OR(F{rr}="",G{rr}=""),"",IF({mma},INDEX({WD_MX},MATCH(MIN(F{rr},G{rr}),{WD_LO},1)),{S["ggap"]}))')
    wsB[f"J{rr}"] = f'=IF({R_}="","",{look(R_, "Experience class")})'
    wsB[f"K{rr}"] = f'=IF({B_}="","",{look(B_, "Experience class")})'
    wsB[f"L{rr}"] = f'=IF({blank},"",IF(AND({look(R_, "Ready to fight?")}="READY",{look(B_, "Ready to fight?")}="READY"),"Yes","Not yet"))'
    ra, ba = look(R_, "Age division"), look(B_, "Age division")
    flags = (f'IF(OR(H{rr}="",I{rr}=""),"",IF(H{rr}>I{rr},"Weight gap over the limit. ",""))'
             f'&IF(AND({mma},J{rr}<>"",K{rr}<>"",J{rr}<>K{rr}),"Different experience classes (needs the Commissioner\'s OK). ","")'
             f'&IF({ra}<>{ba},"Different age divisions. ","")'
             f'&IF(AND({mma},OR(LEFT({ra},5)="Under",LEFT({ba},5)="Under")),"Under 12 can\'t do MMA. ","")'
             f'&IF(AND(C{rr}="Amateur MMA (Senior)",OR(LEFT({ra},5)="Youth",LEFT({ba},5)="Youth")),"Youth on Senior rules. ","")'
             f'&IF({R_}={B_},"Same fighter twice. ","")'
             f'&IF(COUNTIF($D${BH+1}:$E${BH+30},{R_})+COUNTIF($D${BH+1}:$E${BH+30},{B_})>2,"A fighter is in two bouts. ","")'
             f'&IF(L{rr}="Not yet","Paperwork not done. ","")'
             f'&IF(C{rr}="","Pick the rules. ","")')
    wsB[f"M{rr}"] = f'=IF({blank},"",IF(({flags})="","OK",{flags}))'
    for col in "FGHI": wsB[f"{col}{rr}"].number_format = "0"
    wsB.row_dimensions[rr].height = 30
    wsB[f"M{rr}"].alignment = wrap
for h, items in (("Card position", LREF["Card position"]), ("Rules", LREF["Rules"]), ("Cody approved", q(YN)), ("Officials OK", q(YN)), ("Method", LREF["Method"])):
    dv_list(wsB, items, f"{BCOL[h]}{BH+1}:{BCOL[h]}{BH+30}")
for h in ("Red corner", "Blue corner", "Winner"):
    dv_list(wsB, "=" + FA, f"{BCOL[h]}{BH+1}:{BCOL[h]}{BH+30}")
wsB.conditional_formatting.add(f"M{BH+1}:M{BH+30}", CellIsRule(operator="equal", formula=['"OK"'], fill=fill_ok))
wsB.conditional_formatting.add(f"M{BH+1}:M{BH+30}", FormulaRule(formula=[f'AND(M{BH+1}<>"",M{BH+1}<>"OK")'], fill=fill_bad, font=f_flag))
wsB.conditional_formatting.add(f"L{BH+1}:L{BH+30}", CellIsRule(operator="equal", formula=['"Yes"'], fill=fill_ok))
note(wsB, BH + 31, 1, "Ontario amateur MMA: Senior (19+) 3 x 3-minute rounds; Youth A (16-18) 3 x 2; Youth B and C (12-15) 2 x 2; no head or face strikes for any youth level; "
     "no elbows, knees to the head, heel hooks, toe holds, neck cranks or stomps for any amateur. Opponents in different experience classes need the Commissioner's approval. Full list on the Rules tab.", 20)
wsB.row_dimensions[BH + 31].height = 30
wsB.freeze_panes = wsB.cell(BH + 1, 4)

# ------------------------------------------------------------------ Recruitment
wsR = wb.create_sheet("Recruitment")
RC = [("Gym", 26), ("Town", 18), ("Disciplines", 26), ("Competition team?", 11), ("Contact (public)", 30), ("Website", 28),
      ("Stage", 14), ("Fighters offered", 9), ("Last contact", 11), ("Next follow-up", 11), ("Follow-up due?", 9), ("Notes", 50), ("Source", 28)]
banner(wsR, "FIGHTER RECRUITMENT: GYMS",
       "Gyms that could send amateur fighters. Contact the head coach through the gym, never the athletes directly. "
       "Track who's been asked and how many fighters they can send.", len(RC))
head(wsR, 4, [c[0] for c in RC], [c[1] for c in RC])
gyms = []
if MG and MG.get("gyms"):
    gyms = MG["gyms"]
MAXR = 4 + max(60, len(gyms) + 30)
for rr in range(5, MAXR + 1):
    for i, (h, _) in enumerate(RC, 1):
        c = wsR.cell(rr, i); c.border = box
        c.font = f_in if h in ("Stage", "Fighters offered", "Last contact", "Next follow-up", "Notes") else f_lab
        if h in ("Stage", "Fighters offered", "Last contact", "Next follow-up", "Notes"): c.fill = fill_in
        c.alignment = wrap
    wsR[f"I{rr}"].number_format = "mmm d"; wsR[f"J{rr}"].number_format = "mmm d"
    wsR[f"K{rr}"] = f'=IF(OR(A{rr}="",J{rr}="",G{rr}="Sending fighters",G{rr}="Not this time"),"",IF(J{rr}<=TODAY(),"DUE",""))'
    wsR[f"K{rr}"].font = f_flag
CONF = {"high": 0, "medium": 1, "low": 2}
gyms = sorted(gyms, key=lambda g: (0 if str(g.get("has_competition_team", "")).startswith("yes") else 1, CONF.get(g.get("confidence"), 3)))
for k, g in enumerate(gyms):
    rr = 5 + k
    town = g.get("town", "") + (f" ({g['approx_drive_from_bancroft'].replace(' (est.)', '')} drive)" if g.get("approx_drive_from_bancroft") else "")
    promo = str(g.get("runs_own_promotion", "no"))
    nts = " ".join(x for x in [
        ("RUNS ITS OWN PROMOTION: " + promo[4:] + ".") if promo.lower().startswith("yes") else "",
        g.get("notes", "") or "",
        ("Competition: " + g["competition_evidence"]) if g.get("competition_evidence") else "",
        "Low confidence: check the contact first." if g.get("confidence") == "low" else ""] if x)
    vals = [g.get("name", ""), town, g.get("disciplines", ""), g.get("has_competition_team", ""),
            g.get("contact", ""), g.get("website", ""), "Not contacted", None, None, None, None, nts, g.get("source_url", "")]
    for i, v in enumerate(vals, 1):
        if v is not None and i != 11: wsR.cell(rr, i, v)
    wsR.row_dimensions[rr].height = max(42, 13 * (len(nts) // 55 + 1))
if not gyms:
    wsR.cell(5, 1, "(Gym list being researched)").font = f_note
dv_list(wsR, LREF["Gym stage"], f"G5:G{MAXR}")
wsR.freeze_panes = "B5"; wsR.auto_filter.ref = f"A4:{L(len(RC))}{MAXR}"

# ------------------------------------------------------------------ Countdown
wsC = wb.create_sheet("Countdown")
CC = [("Weeks before (minus = days after)", 12), ("Task", 58), ("Area", 12), ("Owner", 10), ("Due", 15), ("Done?", 8), ("Status", 11), ("Notes", 34)]
banner(wsC, "SHOW COUNTDOWN",
       "Every job from booking the venue to the report after the show. Due dates follow the show date on Show Setup. "
       "Mark Done = Yes as things finish. A working list, not legal advice: confirm requirements with the Athletics Commission and the insurer.", len(CC))
head(wsC, 4, [c[0] for c in CC], [c[1] for c in CC])
TASKS = [
    (26, "Pick town and date; check clashes (other combat shows, fairs, long weekends, big UFC cards)", "Venue", "Cody"),
    (26, "Venue quote: rent for setup + show day, capacity, chairs/tables, staff, who runs the bar and CCCL's share", "Venue", "Cody"),
    (24, "Run the Profit Predictor with the real quote; go / no-go", "Money", "Cody"),
    (22, "Hold the date with a deposit; council approval if the venue needs it", "Venue", "Cody"),
    (20, "Sponsor push starts: 2026 sponsors first (Sponsorship Pipeline)", "Sponsors", "Cody"),
    (18, "Apply for the Athletics Commission event permit: $600 for one show (one permit should cover the whole card; confirm). At least 30 days ahead; 120 days for any rule change", "Commission", "Cody"),
    (18, "Book officials for every bout: referee, 3 judges, timekeeper, scorekeeper, plus a neutral cut-person", "Commission", "Cody"),
    (18, "Book the physician (Ontario-licensed, on site all night) and a second medical person. Ask the Commission if it wants an ambulance", "Medical", "Cody"),
    (16, "$2M liability insurance certificate that also covers referees, judges, officials and the Commissioner", "Commission", "Cody"),
    (16, "Book the cage on a platform raised 3-4 ft (Ontario rule) and delivery", "Production", "Cody"),
    (16, "Youth grappling: confirm it's covered by the CCCL permit or sanctioned by the Ontario Grappling Association", "Commission", "Cody"),
    (14, "Tickets on sale on Eventbrite (GA, VIP seats, VIP tables)", "Tickets", "Kyle"),
    (14, "Announce the show: press release to the media list, first social post", "Press", "Kyle"),
    (12, "Fighter recruitment: message the gyms on the Recruitment tab; open the sign-up form", "Fighters", "Cody"),
    (12, "Book security", "Production", "Cody"),
    (10, "Book lights, sound, livestream crew and the announcer", "Production", "Cody"),
    (10, "Poster with sponsor logos; social countdown starts", "Press", "Kyle"),
    (8, "Volunteers: fill every role on the Volunteers tab", "Volunteers", "Cody"),
    (8, "Order event shirts with sponsor logos (sized by level)", "Sponsors", "Cody"),
    (8, "Concussion confirmations (Rowan's Law) from youth and grappling athletes and their parents; concussion code and removal/return-to-sport steps ready", "Fighters", "Cody"),
    (6, "Draft bout card; Cody and officials review (same experience class unless the Commissioner approves)", "Fighters", "Cody"),
    (5, "Send each confirmed fighter their Fighter Package through the coach (Fight Night folder; fill the yellow fields)", "Fighters", "Cody"),
    (4, "Bout card final; give the Commission what it asks for", "Commission", "Cody"),
    (4, "Announce the card: press release and fighter posts (adults who agree; nobody under 18 named)", "Press", "Kyle"),
    (4, "Banners printed; sponsor deliverables checked on the Fulfilment tab", "Sponsors", "Cody"),
    (3, "Day-of forms ready: codes of conduct, rules declaration, no-suspension form (a parent signs for under-18s), permit sign", "Commission", "Cody"),
    (3, "Every fighter READY on the Fighters tab", "Fighters", "Cody"),
    (2, "Gloves (open-hand, 6 oz or more), hand wraps and tape ready: CCCL supplies them, fighters can't bring their own", "Production", "Cody"),
    (2, "Run of show final; announcer script with sponsor reads by level", "Production", "Kyle"),
    (2, "Livestream test from the venue (internet, cameras, overlay with sponsor logos)", "Production", "Kyle"),
    (1, "Volunteer briefing sent; radios, wristbands, scale and supplies on the Production tab ready", "Volunteers", "Cody"),
    (0, "SHOW DAY: weigh-ins within 3 hours of the first bout, then follow the Run of Show", "Show", "Everyone"),
    (-1, "The day after: official results posted; thank-you post", "Press", "Kyle"),
    (-1, "The day after: results press release to the media list", "Press", "Kyle"),
    (-3, "Sponsor thank-yous with photos (within 3 days)", "Sponsors", "Cody"),
    (-3, "Amateur books updated (results, suspensions). Injury details stay in the private amateur books, never in these sheets", "Commission", "Cody"),
    (-7, "Results & Report tab filled in; real numbers into the Profit Predictor", "Money", "Kyle"),
    (-14, "Send the Commission the event recording (required if the show was filmed)", "Commission", "Kyle"),
    (-14, "Post-event report to the Commission within 14 days: contestants, bouts, injuries without names, officials", "Commission", "Cody"),
]
for k, (w, t, area, owner) in enumerate(TASKS):
    rr = 5 + k
    for i in range(1, len(CC) + 1):
        c = wsC.cell(rr, i); c.border = box; c.alignment = wrap
    inp(wsC.cell(rr, 1), w, "0"); wsC.cell(rr, 2, t).font = f_lab; wsC.cell(rr, 3, area).font = f_lab
    inp(wsC.cell(rr, 4), owner)
    c = wsC.cell(rr, 5, f'=IF({S["date"]}="","",IF(A{rr}>=0,{S["date"]}-A{rr}*7,{S["date"]}-A{rr}))'); c.number_format = DATE; c.font = f_calc
    inp(wsC.cell(rr, 6)); inp(wsC.cell(rr, 8))
    wsC.cell(rr, 7, f'=IF(F{rr}="Yes","DONE",IF(E{rr}="","No date yet",IF(E{rr}<TODAY(),"OVERDUE",IF(E{rr}<=TODAY()+7,"This week","Upcoming"))))').font = f_bold
    wsC.row_dimensions[rr].height = 30
lastC = 5 + len(TASKS) - 1
for rr in range(lastC + 1, lastC + 16):
    for i in range(1, len(CC) + 1):
        c = wsC.cell(rr, i); c.border = box
    for i in (1, 2, 3, 4, 6, 8): inp(wsC.cell(rr, i))
    c = wsC.cell(rr, 5, f'=IF(OR(B{rr}="",{S["date"]}=""),"",IF(A{rr}>=0,{S["date"]}-A{rr}*7,{S["date"]}-A{rr}))'); c.number_format = DATE
    wsC.cell(rr, 7, f'=IF(B{rr}="","",IF(F{rr}="Yes","DONE",IF(E{rr}="","No date yet",IF(E{rr}<TODAY(),"OVERDUE",IF(E{rr}<=TODAY()+7,"This week","Upcoming")))))').font = f_bold
dv_list(wsC, q(YN), f"F5:F{lastC + 15}")
for s_, fl in (("DONE", fill_ok), ("OVERDUE", fill_bad), ("This week", fill_warn)):
    wsC.conditional_formatting.add(f"G5:G{lastC + 15}", CellIsRule(operator="equal", formula=[f'"{s_}"'], fill=fl))
wsC.freeze_panes = "C5"

# ------------------------------------------------------------------ Run of Show
wsO = wb.create_sheet("Run of Show")
banner(wsO, "RUN OF SHOW",
       "Times follow Show Setup. Bout names come from the Bout Card, in bout-number order. Print it for the announcer, the timekeeper and the crew lead.", 5)
for c, w in zip("ABCDE", (11, 44, 22, 20, 36)): wsO.column_dimensions[c].width = w
head(wsO, 4, ["Time", "What", "Who", "Where", "Notes"])
r = 5
def orow(time_fml, what, who="", where="", notes="", bold=False, fill=None):
    global r
    c = wsO.cell(r, 1, time_fml); c.number_format = TIME; c.font = f_calc
    wsO.cell(r, 2, what).font = f_bold if bold else f_lab
    for i, v in ((3, who), (4, where), (5, notes)):
        x = wsO.cell(r, i, v); x.font = f_in; x.fill = fill_in
    for i in range(1, 6):
        wsO.cell(r, i).border = box; wsO.cell(r, i).alignment = wrap
        if fill: wsO.cell(r, i).fill = fill
    r += 1
def sec_o(text):
    global r
    section(wsO, r, text, 5); r += 1
sec_o("SHOW DAY")
D = S["doors"]
orow(f"={S['load']}", "Load-in: cage platform (3-4 ft high) and cage build, chairs, tables, VIP setup, permit sign posted", "Crew lead + volunteers", "Arena floor")
orow(f"={S['load']}+120/1440", "Lights, sound and livestream setup; internet test", "Production crew", "")
orow(f"={D}-150/1440", "Officials and medical team arrive; walk the venue and the medical room", "Cody", "")
orow(f"={D}-120/1440", "Sponsor banners hung where each package promised; photo of each", "Volunteers", "", "Printed banner list (from the Sponsorship Pipeline's Fulfilment tab).")
orow(f"={S['weigh']}", "OFFICIAL WEIGH-INS (within 3 hours of the first bout)", "Commission / officials", "Weigh-in room",
     "Overweight: 1 hour to make weight, and out if they lose more than 1% in that hour. Weights onto the Fighters tab.", True, fill_pl)
orow(f"={S['weigh']}+15/1440", "Gloves, wraps and gear checked in the dressing rooms. The referee briefs each fighter there before their own bout (no group meeting)", "Referee + Cody", "Dressing rooms", "CCCL supplies gloves and wraps. Up to 3 corners per fighter.")
orow(f"={S['weigh']}+30/1440", "Fighter check-in after weigh-ins: ID, signed codes and rules declaration, no-suspension form (a parent signs for under-18s). Grappling weigh-ins close (at least 60 minutes before the event)", "Cody + officials", "Weigh-in room", "Medical: confirm Yes/No only. Never write details. Parents may come with junior athletes.")
orow(f"={D}-60/1440", "Pre-bout medicals start (each fighter no later than 1 hour before their bout)", "Physician", "Medical room")
orow(f"={D}-45/1440", "Volunteer briefing: roles, radios, ID and wristband rules, emergency plan", "Cody", "Front entrance")
orow(f"={D}-15/1440", "Ticket scanners and VIP check-in ready", "Door team", "Front entrance")
orow(f"={D}", "DOORS OPEN", "Door team", "", "", True, fill_pl)
orow(f"={S['first']}-10/1440", "Livestream goes live (sponsor slide loop); recording for the Commission running", "Stream crew", "")
orow(f"={S['first']}-5/1440", "Welcome, rules for the crowd, Title Sponsor read", "Announcer", "Cage", "Sponsor reads by level: see the announcer script template.")
sec_o("THE CARD (from the Bout Card tab)")
BR = f"'Bout Card'!$A${BH+1}:$A${BH+30}"
def bout_name(k):
    red = f"INDEX('Bout Card'!$D${BH+1}:$D${BH+30},{k})"
    blue = f"INDEX('Bout Card'!$E${BH+1}:$E${BH+30},{k})"
    rules = f"INDEX('Bout Card'!$C${BH+1}:$C${BH+30},{k})"
    pos = f"INDEX('Bout Card'!$B${BH+1}:$B${BH+30},{k})"
    return red, blue, rules, pos
first_bout_row = r
for k in range(1, 13):
    red, blue, rules, pos = bout_name(k)
    if k == 1:
        t = f"={S['first']}"
    else:
        prev = r - 1 if k - 1 != 0 else r
        # previous bout length depends on its rules; intermission inserted after bout N via extra minutes
        prules = f"INDEX('Bout Card'!$C${BH+1}:$C${BH+30},{k-1})"
        plen = f"IF(OR(LEFT({prules},11)=\"Amateur MMA\",{prules}=\"\"),{S['mma_min']},{S['grap_min']})"
        t = f"=A{r-1}+({plen}+IF({k-1}={S['inter_after']},{S['inter_min']},0))/1440"
    c = wsO.cell(r, 1, t); c.number_format = TIME; c.font = f_calc
    what = f'=IF({red}="","Bout {k}: (not set)","Bout {k}"&IF({pos}="",""," ("&{pos}&")")&": "&{red}&" vs "&{blue}&IF({rules}="",""," · "&{rules}))'
    wsO.cell(r, 2, what).font = f_bold
    wsO.cell(r, 3, "Referee, judges, timekeeper").font = f_lab
    wsO.cell(r, 4, "Cage").font = f_lab
    nt = f'=IF({k}={S["inter_after"]},"INTERMISSION after this bout: sponsor reads, merch, washrooms","")'
    wsO.cell(r, 5, nt).font = f_note
    for i in range(1, 6): wsO.cell(r, i).border = box; wsO.cell(r, i).alignment = wrap
    r += 1
last_bout_row = r - 1
note(wsO, r, 2, "Unused bout rows just say '(not set)'. Hide or delete them when you print.", 5); r += 1
sec_o("AFTER THE MAIN EVENT")
end_fml = (f'=IF(SUMPRODUCT(MAX(ISERROR(SEARCH("(not set)",B{first_bout_row}:B{last_bout_row}))*A{first_bout_row}:A{last_bout_row}))=0,A{first_bout_row},'
           f'SUMPRODUCT(MAX(ISERROR(SEARCH("(not set)",B{first_bout_row}:B{last_bout_row}))*A{first_bout_row}:A{last_bout_row})))+{S["mma_min"]}/1440')
orow(end_fml, "Thank-you from Cody; sponsor thank-you read (all levels); next show date", "Announcer + Cody", "Cage")
orow(f"=A{r-1}+10/1440", "Livestream ends (sponsor slide)", "Stream crew", "")
orow(f"=A{r-1}+15/1440", "Crowd out; post-bout medicals as the physician requires; record any suspensions in the amateur books", "Physician + Cody", "")
orow(f"=A{r-1}+20/1440", "Teardown: cage, banners (return or store), chairs; venue walk-through", "Crew lead + volunteers", "")
wsO.freeze_panes = "A5"

# ------------------------------------------------------------------ Volunteers
wsV = wb.create_sheet("Volunteers")
VC = [("Role", 28), ("What they do", 52), ("Needed", 8), ("Names (first names)", 26), ("Filled", 8), ("Shift", 14), ("Briefed?", 9), ("Notes", 26)]
banner(wsV, "VOLUNTEERS & ROLES", "Who does what on fight night. First names only. Everyone gets the briefing sheet the week before.", len(VC))
head(wsV, 4, [c[0] for c in VC], [c[1] for c in VC])
ROLES = [
    ("Door: tickets and scanning", "Scan Eventbrite tickets, sell walk-ups if allowed, count comps.", 2),
    ("Door: ID and wristbands", "19+ wristbands if there's a bar area. Check ID. No wristband, no bar.", 1),
    ("VIP host", "Seat VIP tables and sponsors, hand out sponsor tickets.", 1),
    ("Ushers and crowd", "Help people find seats, keep aisles and the cage area clear.", 2),
    ("Fighter runner", "Bring fighters and corners to the cage on time; keep the bout order.", 2),
    ("Dressing room monitors", "One per side. Only fighters, corners and officials in the room.", 2),
    ("Youth grappling parent liaison", "Meets parents, checks consent forms, keeps youth athletes with their coach.", 1),
    ("Cage crew", "Wipe the canvas between bouts, check padding, stools and buckets.", 2),
    ("Sponsor banners and photos", "Hang banners where promised, photograph each one with the crowd.", 1),
    ("Merch and 50/50 table", "Sell shirts; run any raffle only if properly licensed.", 2),
    ("Photo and video", "Clips for social and Kyle's highlight video. No youth without parent OK.", 1),
    ("Setup and teardown", "Chairs, tables, banners, cage help.", 6),
]
for k, (role, what, n) in enumerate(ROLES):
    rr = 5 + k
    wsV.cell(rr, 1, role).font = f_bold; wsV.cell(rr, 2, what).font = f_lab
    inp(wsV.cell(rr, 3), n, "0"); inp(wsV.cell(rr, 4)); inp(wsV.cell(rr, 6)); inp(wsV.cell(rr, 7)); inp(wsV.cell(rr, 8))
    wsV.cell(rr, 5, f'=IF(D{rr}="",0,LEN(D{rr})-LEN(SUBSTITUTE(D{rr},",",""))+1)').font = f_calc
    for i in range(1, len(VC) + 1): wsV.cell(rr, i).border = box; wsV.cell(rr, i).alignment = wrap
    wsV.row_dimensions[rr].height = 30
lastV = 5 + len(ROLES) - 1
dv_list(wsV, q(YN), f"G5:G{lastV}")
wsV.cell(lastV + 1, 1, "Total").font = f_bold
wsV.cell(lastV + 1, 3, f"=SUM(C5:C{lastV})").font = f_bold
wsV.cell(lastV + 1, 5, f"=SUM(E5:E{lastV})").font = f_bold
note(wsV, lastV + 2, 1, "Type names separated by commas; 'Filled' counts them. A role turns red until it's filled.", 8)
wsV.conditional_formatting.add(f"E5:E{lastV}", FormulaRule(formula=["E5<C5"], fill=fill_bad))
wsV.conditional_formatting.add(f"E5:E{lastV}", FormulaRule(formula=["E5>=C5"], fill=fill_ok))
wsV.freeze_panes = "B5"

# ------------------------------------------------------------------ Production
wsP = wb.create_sheet("Production")
PC = [("Item", 34), ("Area", 14), ("Supplier / who", 22), ("Cost ($)", 10), ("Booked?", 9), ("Confirmed show week?", 11), ("Notes", 40)]
banner(wsP, "PRODUCTION CHECKLIST", "Everything that has to physically be there. Costs here feed the Results & Report tab.", len(PC))
head(wsP, 4, [c[0] for c in PC], [c[1] for c in PC])
ITEMS = [
    ("Cage on a platform raised 3-4 ft, padded, with canvas (Ontario rule): rental, delivery, build, teardown", "Cage"),
    ("Canvas cleaning supplies", "Cage"), ("Corner stools, buckets, towels, ice", "Cage"),
    ("Timekeeper bell or horn, clock", "Officials"), ("Judges' tables (3 sides), scorecards, pens; scorekeeper table", "Officials"),
    ("Scale for weigh-ins", "Officials"), ("Permit sign posted at the venue", "Officials"),
    ("Open-hand MMA gloves, 6 oz or more, in each size (CCCL supplies; fighters can't bring their own)", "Fighters"),
    ("Hand wraps, tape and someone to wrap", "Fighters"), ("Spare soft neoprene shin guards (no straps or Velcro)", "Fighters"),
    ("Neutral cut-person (third second)", "Fighters"), ("Dressing rooms: two sides, chairs, water", "Fighters"),
    ("Medical room, physician and second medical person", "Medical"), ("First-aid supplies; clear path for an ambulance if one is called", "Medical"),
    ("Lights over the cage", "Production"), ("Sound system and microphones", "Production"),
    ("Announcer", "Production"), ("Walkout music (fighters send songs in advance)", "Production"),
    ("Livestream crew, cameras, internet, overlay with sponsor logos", "Livestream"), ("Recording of the show for the Commission (if filmed)", "Livestream"),
    ("Big screen or TV for the crowd (optional)", "Production"),
    ("Sponsor banners printed", "Sponsors"), ("Event shirts with sponsor logos", "Sponsors"),
    ("Ticket scanners (phones) and walk-up payment", "Front door"), ("Wristbands (19+ area)", "Front door"),
    ("Radios for the crew", "Crew"), ("Security company", "Security"), ("Chairs, tables, VIP table setup", "Venue"),
    ("Merch stock and card reader", "Merch"),
]
for k, (it, area) in enumerate(ITEMS):
    rr = 5 + k
    wsP.cell(rr, 1, it).font = f_lab; wsP.cell(rr, 2, area).font = f_lab
    for i in (3, 4, 5, 6, 7): inp(wsP.cell(rr, i))
    wsP.cell(rr, 4).number_format = USD
    for i in range(1, len(PC) + 1): wsP.cell(rr, i).border = box; wsP.cell(rr, i).alignment = wrap
lastP = 5 + len(ITEMS) - 1
for rr in range(lastP + 1, lastP + 11):
    for i in range(1, len(PC) + 1): inp(wsP.cell(rr, i))
    wsP.cell(rr, 4).number_format = USD
dv_list(wsP, q(YN), f"E5:F{lastP + 10}")
wsP.cell(lastP + 11, 1, "Total production cost").font = f_bold
c = wsP.cell(lastP + 11, 4, f"=SUM(D5:D{lastP + 10})"); c.font = f_bold; c.number_format = USD
wsP.conditional_formatting.add(f"E5:F{lastP + 10}", CellIsRule(operator="equal", formula=['"Yes"'], fill=fill_ok))
wsP.freeze_panes = "B5"
PROD_TOTAL = f"Production!$D${lastP + 11}"

# ------------------------------------------------------------------ Results & Report
wsX = wb.create_sheet("Results & Report")
banner(wsX, "RESULTS & REPORT",
       "Fill in within a week of the show. Budget = the venue's column on the Profit Predictor's Compare tab (expected crowd). "
       "The real numbers then go into the Predictor so the next forecast is better.", 5)
for c, w in zip("ABCDE", (44, 16, 16, 16, 52)): wsX.column_dimensions[c].width = w
r = 4
head(wsX, r, ["", "Real", "Budget (Predictor)", "Difference", "Where the real number comes from"])
r += 1
def xrow(lab, fmt=USD, src="", real=None, bold=False):
    global r
    wsX.cell(r, 1, lab).font = f_bold if bold else f_lab
    c = wsX.cell(r, 2)
    if real is None: inp(c, None, fmt)
    else: c.value = real; c.font = f_calc; c.number_format = fmt; c.border = box
    inp(wsX.cell(r, 3), None, fmt)
    d = wsX.cell(r, 4, f'=IF(OR(B{r}="",C{r}=""),"",B{r}-C{r})'); d.number_format = fmt; d.font = f_calc; d.border = box
    note(wsX, r, 5, src)
    k = r; r += 1; return k
section(wsX, r, "PEOPLE", 5); r += 1
x_att = xrow("People in the building", NUM, "Door count or scans plus comps")
x_paid = xrow("Paid tickets", NUM, "Eventbrite sales report")
x_comp = xrow("Free tickets (sponsors, fighters' teams, volunteers)", NUM, "Comp list")
x_views = xrow("Livestream viewers (peak or total)", NUM, "Streaming platform stats")
section(wsX, r, "MONEY IN", 5); r += 1
m1 = xrow("Ticket money after fees", USD, "Eventbrite payout report")
m2 = xrow("Sponsorship actually paid", USD, "Sponsorship Pipeline: 'Amount paid' for this show")
m3 = xrow("Bar share to CCCL", USD, "Venue settlement")
m4 = xrow("Merch profit", USD, "Merch table count")
m5 = xrow("Livestream income", USD, "Platform payout")
mt = xrow("TOTAL MONEY IN", USD, "", real=f"=SUM(B{m1}:B{m5})", bold=True)
wsX[f"C{mt}"] = f'=IF(COUNT(C{m1}:C{m5})=0,"",SUM(C{m1}:C{m5}))'; wsX[f"C{mt}"].font = f_calc
section(wsX, r, "COSTS (same lines as the Predictor)", 5); r += 1
cost_rows = []
for lab, src in (("Venue rent and staff", "Venue invoice"), ("Athletics Commission permit", "Commission receipt"), ("Cage", "Invoice"),
                 ("Officials", "Invoice"), ("Medical", "Invoices"), ("Insurance", "Policy"), ("Security", "Invoice"),
                 ("Lights, sound, livestream", f"Production tab total is a starting point"), ("Marketing, posters, ads", "Receipts"),
                 ("Athlete costs: belts, medals, travel", "Receipts"), ("Chairs, tables, VIP setup", "Invoice"), ("Crew and volunteer food", "Receipts"),
                 ("Shirts and banners", "Invoices"), ("Anything else", "Receipts")):
    cost_rows.append(xrow(lab, USD, src))
ct = xrow("TOTAL COST", USD, "", real=f"=SUM(B{cost_rows[0]}:B{cost_rows[-1]})", bold=True)
wsX[f"C{ct}"] = f'=IF(COUNT(C{cost_rows[0]}:C{cost_rows[-1]})=0,"",SUM(C{cost_rows[0]}:C{cost_rows[-1]}))'; wsX[f"C{ct}"].font = f_calc
pl = xrow("PROFIT / (LOSS)", USD, "Before any money goes to the gym", real=f"=B{mt}-B{ct}", bold=True)
wsX[f"C{pl}"] = f'=IF(OR(C{mt}="",C{ct}=""),"",C{mt}-C{ct})'; wsX[f"C{pl}"].font = f_calc
for rr in (mt, ct, pl):
    for cc in "ABCD": wsX[f"{cc}{rr}"].fill = fill_pl
wsX.conditional_formatting.add(f"B{pl}:D{pl}", CellIsRule(operator="lessThan", formula=["0"], font=f_flag))
r += 1
section(wsX, r, "THE FIGHTS (from the Bout Card)", 5); r += 1
for lab, fml in (("Bouts that happened", f"=COUNTIF('Bout Card'!$P${BH+1}:$P${BH+30},\"?*\")"),
                 ("Finishes (KO, TKO, submission)", f"=COUNTIF('Bout Card'!$Q${BH+1}:$Q${BH+30},\"KO\")+COUNTIF('Bout Card'!$Q${BH+1}:$Q${BH+30},\"TKO\")+COUNTIF('Bout Card'!$Q${BH+1}:$Q${BH+30},\"Submission\")"),
                 ("Decisions", f"=COUNTIF('Bout Card'!$Q${BH+1}:$Q${BH+30},\"Decision*\")")):
    wsX.cell(r, 1, lab).font = f_lab; c = wsX.cell(r, 2, fml); c.font = f_link; c.border = box; r += 1
r += 1
section(wsX, r, "WRAP-UP", 5); r += 1
for lab in ("Post-event report sent to the Commission (date)", "Sponsor thank-yous sent (date)", "Real numbers copied into the Predictor (date)"):
    wsX.cell(r, 1, lab).font = f_lab; inp(wsX.cell(r, 2), None, "mmm d, yyyy"); r += 1
for lab in ("What worked", "What to change next time", "Fighter or coach feedback", "Sponsor feedback"):
    wsX.cell(r, 1, lab).font = f_bold; inp(wsX.cell(r, 2))
    wsX.merge_cells(start_row=r, start_column=2, end_row=r, end_column=5); wsX.row_dimensions[r].height = 48
    wsX.cell(r, 2).alignment = wrap; r += 1
wsX.freeze_panes = "A5"

# ------------------------------------------------------------------ Press
wsM = wb.create_sheet("Press")
MC = [("Outlet", 28), ("Type", 14), ("Area", 18), ("How to send news", 40), ("Show announced (date)", 11), ("Card sent (date)", 11),
      ("Results sent (date)", 11), ("Ran a story?", 10), ("Notes", 30), ("Source", 30)]
banner(wsM, "PRESS LIST",
       "Local papers, news sites, radio, TV and combat-sports media, with each one's own way to send news. "
       "Three releases per show: announcement (about 14 weeks out), the card (4 weeks), results (the day after).", len(MC))
head(wsM, 4, [c[0] for c in MC], [c[1] for c in MC])
media = MG.get("media", []) if MG else []
MAXM = 4 + max(40, len(media) + 15)
for rr in range(5, MAXM + 1):
    for i, (h, _) in enumerate(MC, 1):
        c = wsM.cell(rr, i); c.border = box; c.alignment = wrap
        if i in (5, 6, 7, 8, 9): c.font = f_in; c.fill = fill_in
        else: c.font = f_lab
    for col in "EFG": wsM[f"{col}{rr}"].number_format = "mmm d"
for k, m in enumerate(media):
    rr = 5 + k
    for i, v in enumerate([m.get("name", ""), m.get("type", ""), m.get("area", ""), m.get("how_to_submit", ""), None, None, None, None,
                           (m.get("notes", "") or "") + (" Low confidence: check first." if m.get("confidence") == "low" else ""), m.get("source_url", "")], 1):
        if v is not None: wsM.cell(rr, i, v)
    wsM.row_dimensions[rr].height = 42
if not media:
    wsM.cell(5, 1, "(Media list being researched)").font = f_note
dv_list(wsM, q(["Yes", "No"]), f"H5:H{MAXM}")
wsM.freeze_panes = "B5"; wsM.auto_filter.ref = f"A4:{L(len(MC))}{MAXM}"

# ------------------------------------------------------------------ Rules (from research)
wsQ = wb.create_sheet("Rules")
banner(wsQ, "RULES THAT SHAPE FIGHT NIGHT (Ontario)",
       "A plain-English summary from official sources, for planning only. It isn't legal advice, and rules change: "
       "confirm everything with the Ontario Athletics Commission (AthleticsComm@ontario.ca, 416-326-0416) before each show.", 4)
for c, w in zip("ABCD", (26, 74, 46, 11)): wsQ.column_dimensions[c].width = w
r = 4
if OAC:
    if OAC.get("scope_note"):
        note(wsQ, r, 1, OAC["scope_note"], 4); wsQ.row_dimensions[r].height = 30; r += 2
    section(wsQ, r, "THE BIG ONES FOR CCCL", 4); r += 1
    BIG = [("One permit per show", "One $600 permit should cover the whole card (a 'contest' is one bout or several in a row). Confirm with the Commission."),
           ("Weigh-ins", "Show day only, within 3 hours of the first bout. Overweight: 1 hour to make weight, out if they lose more than 1% in that hour."),
           ("Weight gap (MMA)", "Max difference: 4 lb under 148, 9 lb at 148-165, 11 lb at 165-200, 24 lb at 200-265, no limit over 265."),
           ("Experience classes", "Pre-Novice 0-5 amateur contests (any combat sport), Novice 6-10, Open 11+. Different classes need the Commissioner's OK."),
           ("Ages and rounds", "Senior 19+: 3 x 3 min. Youth A 16-18: 3 x 2. Youth B 14-15 and Youth C 12-13: 2 x 2. No MMA under 12. No head or face strikes for any youth."),
           ("Banned for all amateurs", "Elbows, knees to the head, heel hooks, toe holds, neck cranks, stomps. Pre-Novice also: ground strikes to the head, slams, up-kicks."),
           ("Gear", "CCCL supplies open-hand gloves (6 oz+) and hand wraps; fighters can't bring their own gloves. Soft neoprene shin guards for everyone. Cup for men. No headgear or shoes."),
           ("Officials and medical", "Every bout: referee, 3 judges, timekeeper, scorekeeper, ringside medical. A CPSO-licensed physician on site all night plus one more medical person."),
           ("Cage", "On a platform raised 3-4 ft. A floor-level cage would need an exemption asked for 120 days ahead."),
           ("Amateur status", "A purse or prize worth $300 or more makes a fighter a pro. Belts, medals and trophies are fine."),
           ("Youth grappling", "Regulated, not exempt: covered by the OAC permit or sanctioned by the Ontario Grappling Association. Rowan's Law concussion steps likely apply."),
           ("After the show", "Post-event report within 14 days. Give the Commission the recording if the show was filmed. Keep the amateur books.")]
    for t, b_ in BIG:
        c = wsQ.cell(r, 1, t); c.font = f_bold; c.alignment = wrap; c.border = box
        c = wsQ.cell(r, 2, b_); c.font = f_lab; c.alignment = wrap; c.border = box
        wsQ.merge_cells(start_row=r, start_column=2, end_row=r, end_column=4)
        wsQ.row_dimensions[r].height = 30; r += 1
    r += 1
    section(wsQ, r, "EVERY RULE FOUND, WITH SOURCES", 4); r += 1
    head(wsQ, r, ["Topic", "What it means for CCCL", "Source", "Confidence"]); r += 1
    for it in OAC.get("items", []):
        vals = [it.get("topic", ""), it.get("rule", ""), it.get("source_url", "") + (" (secondary source)" if it.get("source_type") == "secondary" else ""), it.get("confidence", "")]
        for i, v in enumerate(vals, 1):
            c = wsQ.cell(r, i, v); c.font = f_bold if i == 1 else f_lab; c.alignment = wrap; c.border = box
        wsQ.row_dimensions[r].height = max(28, 13 * (len(vals[1]) // 80 + 1))
        r += 1
    if OAC.get("open_questions"):
        r += 1; section(wsQ, r, "ASK THE COMMISSION", 4); r += 1
        for qq in OAC["open_questions"]:
            c = wsQ.cell(r, 1, "•  " + qq); c.font = f_lab; c.alignment = wrap
            wsQ.merge_cells(start_row=r, start_column=1, end_row=r, end_column=4)
            wsQ.row_dimensions[r].height = max(18, 13 * (len(qq) // 140 + 1)); r += 1
else:
    wsQ.cell(r, 1, "(Being researched)").font = f_note
wsQ.freeze_panes = "A4"

# ------------------------------------------------------------------ Start Here
ws0 = wb.create_sheet("Start Here", 0)
banner(ws0, "CCCL FIGHT NIGHT KIT",
       "Everything to run a Cottage Country Combat League show: fighters, the bout card, the countdown, the night itself and the report after. "
       "Built for Cody Purtell by Kyle. Opens in Google Sheets or Excel.", 3)
ws0.column_dimensions["A"].width = 4; ws0.column_dimensions["B"].width = 26; ws0.column_dimensions["C"].width = 96
r = 4
def sec0(t):
    global r
    ws0.cell(r, 2, t).font = f_sec; r += 1
def line0(a, b, n=None):
    global r
    if n: ws0.cell(r, 1, n).font = Font(name=F, size=12, bold=True, color=RED)
    ws0.cell(r, 2, a).font = f_bold; ws0.cell(r, 2).alignment = top
    c = ws0.cell(r, 3, b); c.font = f_lab; c.alignment = wrap
    if len(b) > 100: ws0.row_dimensions[r].height = 28
    r += 1
sec0("IN ORDER")
for n, (a, b) in enumerate((
        ("Show Setup", "Show name, venue, date and times. Everything else follows these."),
        ("Countdown", "Every job from booking the venue to the Commission report after, with due dates worked out from the show date."),
        ("Recruitment", "Gyms that could send fighters. Message the head coach through the gym and track who's sending fighters."),
        ("Fighters", "Everyone who applies. Works out each fighter's experience class. 'Ready to fight?' turns green once the medical, signed forms, waiver, and parent consent and concussion confirmation where needed are done."),
        ("Bout Card", "Pick both corners. It checks the Ontario weight-gap table, experience classes, age divisions and paperwork. Cody and the officials decide."),
        ("Volunteers", "Every fight-night role, how many people, who's doing it."),
        ("Production", "Cage, lights, sound, livestream, supplies: booked and confirmed."),
        ("Run of Show", "The night minute by minute, built from Show Setup and the Bout Card. Print it."),
        ("Press", "Media list with each outlet's own way to send news. Three releases per show."),
        ("Results & Report", "Real numbers vs the Predictor, the fight results and what to change. Feeds the next forecast."),
        ("Rules", "Ontario rules that shape matchmaking and the night, with sources. Confirm with the Commission."),
), 1):
    line0(a, b, str(n))
r += 1
sec0("GROUND RULES")
for a, b in (("Who decides", "Cody and the officials decide every bout. The sheet only flags things to look at again. Results go in only once official."),
             ("Medical", "Medical cleared: Yes, No or Pending. Never write a diagnosis, test result or injury in this sheet."),
             ("Youth athletes", "First name and last initial only, a parent's consent first, and no photos or video without a parent's OK."),
             ("Contact", "Reach fighters through their gym or coach. No personal phone numbers in the sheet."),
             ("Colour key", "Yellow cells with blue text are yours to fill in. Black text is calculated. Green text is pulled from another tab.")):
    line0(a, b)
r += 1
sec0("WORKS WITH")
for a, b in (("Profit Predictor", "Budget numbers for the Results tab; real numbers go back into it after the show."),
             ("Sponsorship Pipeline", "Sponsor banners, shirt logos and announcer reads by level (Fulfilment tab)."),
             ("Claude", "The cccl-fight-night skill drafts gym outreach, fighter packages, the announcer script, press releases and the post-show report from this sheet.")):
    line0(a, b)
ws0.sheet_view.showGridLines = False

wsL.sheet_state = "hidden"
order = ["Start Here", "Show Setup", "Countdown", "Recruitment", "Fighters", "Bout Card", "Volunteers", "Production", "Run of Show",
         "Press", "Results & Report", "Rules", "Lists"]
wb._sheets = [wb[n] for n in order]
for ws in wb.worksheets:
    ws.sheet_properties.tabColor = RED if ws.title in ("Start Here", "Bout Card", "Run of Show") else PINE
    ws.page_setup.orientation = "landscape"; ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
wb.active = 0
wb.save(OUT)
print("saved", OUT)
