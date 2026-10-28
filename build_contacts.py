"""Build 'BBG & CCCL Contacts.xlsx': one permanent address book for every venue, gym, media outlet, supplier, official and partner.
Seeded from the research files. Google Sheets compatible (COUNTIFS, SUMPRODUCT; inline dropdowns without commas).
usage: python3 build_contacts.py OUT research/venues_2027.json research/town_scout.json research/media_gyms.json"""
import sys, json, re
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.utils import get_column_letter as L

OUT = sys.argv[1]
V27 = json.load(open(sys.argv[2])); TS = json.load(open(sys.argv[3])); MG = json.load(open(sys.argv[4]))
NAVY, GOLD, GOLD_L = "081C2B", "EAA959", "FBEBD3"
F = "Arial"
f_title = Font(name=F, size=16, bold=True, color=GOLD); f_sub = Font(name=F, size=10, color="FFFFFF")
f_head = Font(name=F, size=10, bold=True, color="FFFFFF"); f_lab = Font(name=F, size=10); f_bold = Font(name=F, size=10, bold=True)
f_in = Font(name=F, size=10, color="0000FF"); f_note = Font(name=F, size=9, italic=True, color="555555")
f_sec = Font(name=F, size=11, bold=True, color=NAVY); f_calc = Font(name=F, size=10, bold=True, color="000000")
f_link = Font(name=F, size=9, color="1F5FA8")
fill_navy = PatternFill("solid", fgColor=NAVY); fill_in = PatternFill("solid", fgColor="FFFF00"); fill_goldl = PatternFill("solid", fgColor=GOLD_L)
thin = Side(style="thin", color="BFBFBF"); box = Border(top=thin, bottom=thin, left=thin, right=thin)
wrap = Alignment(wrap_text=True, vertical="top"); DATE = "mmm d, yyyy"

TYPES = ["Venue", "Gym", "Media", "Supplier", "Official", "Sponsor", "Partner", "Other"]
RELS = ["Not contacted", "Cold", "Warm", "Strong", "Partner", "Not a fit"]
USED = ["CCCL", "BBG", "Both"]
OWNERS = ["Cody", "Kyle", "Sean", "Jordan", "Front desk"]
CHANNELS = ["Email", "Phone", "Text", "In person", "Facebook", "Instagram", "Other"]

def short(s, n):
    s = re.sub(r"\s+", " ", str(s or "")).strip()
    return s if len(s) <= n else s[: n - 1].rstrip() + "…"
def norm(s): return re.sub(r"\(.*?\)|[^a-z0-9]", "", s.lower())

rows = []  # (org, type, town, contact, website, used, rel, owner, notes, source)
seen = set()
for t in V27["towns"]:
    for v in t["venues"]:
        k = norm(v["name"]); seen.add(k)
        rel = "Strong" if "LaRue" in v["name"] else "Not contacted"
        note = "Hosted CCCL 1 (May 2026). " if "LaRue" in v["name"] else ""
        note += f"Capacity: {short(v.get('capacity_event'), 110)} Rate: {short(v.get('rental_rate'), 110)}"
        rows.append((v["name"], "Venue", t["town"], short(v.get("booking_contact"), 160), "", "CCCL", rel, "Cody", note, v.get("source_url", "")))
for t in TS["towns"]:
    for v in t["venues"]:
        if norm(v["name"]) in seen or any(norm(v["name"])[:18] == s[:18] for s in seen): continue
        seen.add(norm(v["name"]))
        note = f"Capacity: {short(v.get('capacity'), 90)} Rate: {short(v.get('rental_rate'), 110)}"
        rows.append((v["name"], "Venue", t["town"].replace(" (added)", ""), short(v.get("contact"), 160), "", "CCCL", "Not contacted", "Cody", note, v.get("source_url", "")))
for g in MG["gyms"]:
    ev = (g.get("competition_evidence") or "") + " " + (g.get("notes") or "")
    worked = ("CCCL" in ev) and any(w in ev for w in ("CCCL 1", "already", "Already", "sent a fighter", "fought on CCCL"))
    note = short(g.get("notes"), 190)
    rows.append((g["name"], "Gym", g["town"], short(g.get("contact"), 160), g.get("website", ""), "Both" if "Bancroft" in g["town"] else "CCCL",
                 "Warm" if worked else "Not contacted", "Cody", note, g.get("source_url", "")))
for m in MG["media"]:
    worked = any(w in (m.get("notes") or "") for w in ("Already covers CCCL", "covered CCCL", "Best proven", "wrote"))
    rows.append((m["name"], "Media", m.get("area", ""), short(m.get("how_to_submit"), 170), "", "Both",
                 "Warm" if worked else "Not contacted", "Cody", short(m.get("notes"), 170), m.get("source_url", "")))
rows.append(("Ontario Athletics Commission", "Official", "Toronto (province-wide)", "AthleticsComm@ontario.ca; 416-326-0416", "https://www.ontario.ca/page/amateur-combative-sports",
             "CCCL", "Warm", "Cody", "Event permit ($600 per show), officials, rules questions, post-event report within 14 days.", "https://www.ontario.ca/page/amateur-combative-sports"))

wb = Workbook()
def banner(ws, title, sub, width):
    for r in (1, 2):
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=width)
        for c in range(1, width + 1): ws.cell(r, c).fill = fill_navy
    ws["A1"] = title; ws["A1"].font = f_title; ws["A1"].alignment = Alignment(vertical="center")
    ws["A2"] = sub; ws["A2"].font = f_sub; ws["A2"].alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[1].height = 30; ws.row_dimensions[2].height = 32
def head(ws, r, cols):
    for i, (h, w) in enumerate(cols, 1):
        c = ws.cell(r, i, h); c.font = f_head; c.fill = fill_navy; c.border = box; c.alignment = Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[L(i)].width = w
    ws.row_dimensions[r].height = 30
def inp(c, v=None, fmt=None):
    c.value = v; c.font = f_in; c.fill = fill_in; c.border = box; c.alignment = wrap
    if fmt: c.number_format = fmt
def dv_list(ws, items, rng):
    d = DataValidation(type="list", formula1='"' + ",".join(items) + '"', allow_blank=True); ws.add_data_validation(d); d.add(rng)

# ---------------- Start Here
ws0 = wb.active; ws0.title = "Start Here"
banner(ws0, "BBG & CCCL CONTACTS", "Everyone we work with, in one place: venues, gyms, media, suppliers, officials and partners. It stays the same from show to show.", 3)
ws0.column_dimensions["A"].width = 4; ws0.column_dimensions["B"].width = 28; ws0.column_dimensions["C"].width = 92
steps = [
 ("1", "What it's for", "The permanent address book. The Fight Night Kit's Recruitment and Press tabs are per-show working lists (you copy the Kit for each show); this sheet keeps the relationship from one show to the next."),
 ("2", "Sponsors live elsewhere", "Sponsors and prospects stay in the CCCL Sponsorship Pipeline (stages, money, follow-ups). Only add a sponsor here if they're also a supplier or partner."),
 ("3", "Log every contact", "On the Log tab, add one row each time someone calls, emails or meets a contact. 'Last contact' on the Contacts tab fills in by itself."),
 ("4", "Next step and date", "Put the next step and its date on the Contacts tab. 'Due?' turns red when it's overdue, and Claude's morning report lists what's due."),
 ("5", "Starting data", f"{len(rows)} contacts from Kyle's research (Oct 2026): venues, gyms, media and the Athletics Commission. Public business contact details only. Check a number before relying on it."),
 ("6", "Privacy", "Public business contacts only: no personal cell numbers, home addresses or anything about health. Members, kids and Second Chance participants never go in this sheet."),
]
r = 4
for n, a, b in steps:
    ws0.cell(r, 1, n).font = Font(name=F, size=12, bold=True, color=GOLD); ws0.cell(r, 2, a).font = f_bold
    ws0.cell(r, 1).alignment = Alignment(vertical="top"); ws0.cell(r, 2).alignment = Alignment(vertical="top", wrap_text=True)
    c = ws0.cell(r, 3, b); c.font = f_lab; c.alignment = wrap; ws0.row_dimensions[r].height = 32; r += 1
r += 1
ws0.cell(r, 2, "Relationship").font = f_sec; r += 1
for a, b in (("Not contacted", "From research; nobody has reached out yet"), ("Cold", "Contacted, no real conversation yet"), ("Warm", "Talking, or worked together once"),
             ("Strong", "Worked together more than once; picks up when Cody calls"), ("Partner", "Formal or ongoing arrangement (e.g. a venue or gym we always use)"), ("Not a fit", "Keep for reference; don't chase")):
    ws0.cell(r, 2, a).font = f_bold; ws0.cell(r, 3, b).font = f_lab; r += 1
r += 1
ws0.cell(r, 2, "Colours").font = f_sec; r += 1
for lab, fill, txt in (("Yellow, blue text", fill_in, "You type here"), ("Black, bold", None, "Calculated: don't type over it")):
    c = ws0.cell(r, 2, lab); c.font = f_in if fill else f_lab
    if fill: c.fill = fill
    ws0.cell(r, 3, txt).font = f_lab; r += 1

# ---------------- Contacts
ws = wb.create_sheet("Contacts")
CC = [("Organization", 34), ("Type", 10), ("Town / area", 16), ("Contact (public business details)", 44), ("Website", 22), ("For", 7), ("Relationship", 13),
      ("Last contact", 12), ("Next step", 30), ("Next date", 12), ("Owner", 10), ("Due?", 11), ("Notes", 46), ("Source", 22), ("Days since", 8)]
banner(ws, "CONTACTS", "One row per organization. Type the next step and date; 'Last contact' comes from the Log tab. Filter by Type or Due?.", len(CC))
HR = 4; head(ws, HR, CC)
MAXR = HR + len(rows) + 150
order = {"Venue": 0, "Official": 1, "Gym": 2, "Media": 3}
rows.sort(key=lambda x: (order.get(x[1], 9), x[2], x[0]))
for k in range(MAXR - HR):
    rr = HR + 1 + k
    row = rows[k] if k < len(rows) else None
    if row:
        org, typ, town, contact, web, used, rel, owner, notes, src = row
        for i, v in ((1, org), (2, typ), (3, town), (4, contact), (5, web), (6, used), (7, rel)):
            inp(ws.cell(rr, i), v)
        ws.cell(rr, 1).font = Font(name=F, size=10, bold=True, color="0000FF")
        inp(ws.cell(rr, 9)); inp(ws.cell(rr, 10), None, DATE); inp(ws.cell(rr, 11), owner); inp(ws.cell(rr, 13), notes)
        c = ws.cell(rr, 14, src); c.font = f_link; c.border = box; c.alignment = wrap
        ws.row_dimensions[rr].height = 54
    else:
        for i in list(range(1, 8)) + [9, 10, 11, 13, 14]: inp(ws.cell(rr, i))
        ws.cell(rr, 10).number_format = DATE
    c = ws.cell(rr, 8, f'=IF(A{rr}="","",IF(COUNTIFS(Log!$B$5:$B$2000,A{rr})=0,"",SUMPRODUCT(MAX((Log!$B$5:$B$2000=A{rr})*Log!$A$5:$A$2000))))')
    c.number_format = DATE; c.font = f_calc; c.border = box; c.alignment = wrap
    c = ws.cell(rr, 12, f'=IF(OR(A{rr}="",J{rr}=""),"",IF(J{rr}<TODAY(),"OVERDUE",IF(J{rr}=TODAY(),"TODAY",IF(J{rr}<=TODAY()+7,"THIS WEEK",""))))')
    c.font = f_calc; c.border = box; c.alignment = wrap
    c = ws.cell(rr, 15, f'=IF(H{rr}="","",TODAY()-H{rr})'); c.font = f_calc; c.border = box; c.number_format = "0"
dv_list(ws, TYPES, f"B5:B{MAXR}"); dv_list(ws, USED, f"F5:F{MAXR}"); dv_list(ws, RELS, f"G5:G{MAXR}"); dv_list(ws, OWNERS, f"K5:K{MAXR}")
d = DataValidation(type="date", allow_blank=True); ws.add_data_validation(d); d.add(f"J5:J{MAXR}")
Lr = f"L5:L{MAXR}"
ws.conditional_formatting.add(Lr, CellIsRule(operator="equal", formula=['"OVERDUE"'], fill=PatternFill("solid", fgColor="FFC7CE"), font=Font(name=F, color="9C0006", bold=True)))
ws.conditional_formatting.add(Lr, CellIsRule(operator="equal", formula=['"TODAY"'], fill=PatternFill("solid", fgColor="FFEB9C"), font=Font(name=F, color="9C5700", bold=True)))
ws.conditional_formatting.add(Lr, CellIsRule(operator="equal", formula=['"THIS WEEK"'], fill=PatternFill("solid", fgColor="DDEBF7"), font=Font(name=F, color="1F4E78", bold=True)))
ws.conditional_formatting.add(f"G5:G{MAXR}", CellIsRule(operator="equal", formula=['"Partner"'], fill=PatternFill("solid", fgColor="C6EFCE")))
ws.conditional_formatting.add(f"G5:G{MAXR}", CellIsRule(operator="equal", formula=['"Strong"'], fill=PatternFill("solid", fgColor="E2F0D9")))
ws.freeze_panes = "B5"; ws.auto_filter.ref = f"A{HR}:O{MAXR}"
ws.column_dimensions["O"].hidden = False

# ---------------- Log
wl = wb.create_sheet("Log")
LC = [("Date", 12), ("Organization (pick from Contacts)", 36), ("Who (staff)", 12), ("How", 11), ("What happened", 56), ("Next step", 34), ("Next date", 12)]
banner(wl, "CONTACT LOG", "One row every time someone calls, emails, texts or meets a contact. Newest at the bottom. Copy the next step and date onto the Contacts tab.", len(LC))
head(wl, 4, LC)
for rr in range(5, 2001):
    for i in range(1, 8): inp(wl.cell(rr, i))
    wl.cell(rr, 1).number_format = DATE; wl.cell(rr, 7).number_format = DATE
d = DataValidation(type="list", formula1=f"=Contacts!$A$5:$A${MAXR}", allow_blank=True); wl.add_data_validation(d); d.add("B5:B2000")
dv_list(wl, OWNERS, "C5:C2000"); dv_list(wl, CHANNELS, "D5:D2000")
d = DataValidation(type="date", allow_blank=True); wl.add_data_validation(d); d.add("A5:A2000"); d.add("G5:G2000")
wl.freeze_panes = "A5"

# ---------------- Summary
wsS = wb.create_sheet("Summary", 1)
cols = ["Type"] + RELS + ["Total", "Due this week"]
banner(wsS, "SUMMARY", "Who we know and how well, and what's due. Updates by itself.", len(cols))
wsS.column_dimensions["A"].width = 16
for i in range(2, len(cols) + 1): wsS.column_dimensions[L(i)].width = 13
for i, h in enumerate(cols, 1):
    c = wsS.cell(4, i, h); c.font = f_head; c.fill = fill_navy; c.border = box; c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
wsS.row_dimensions[4].height = 30
B_ = f"Contacts!$B$5:$B${MAXR}"; G_ = f"Contacts!$G$5:$G${MAXR}"; L_ = f"Contacts!$L$5:$L${MAXR}"; A_ = f"Contacts!$A$5:$A${MAXR}"
for k, t in enumerate(TYPES):
    r = 5 + k
    wsS.cell(r, 1, t).font = f_bold; wsS.cell(r, 1).border = box
    for j, rel in enumerate(RELS):
        c = wsS.cell(r, 2 + j, f'=COUNTIFS({B_},$A{r},{G_},"{rel}")'); c.font = f_calc; c.border = box; c.alignment = Alignment(horizontal="center")
    c = wsS.cell(r, 2 + len(RELS), f'=COUNTIFS({B_},$A{r})'); c.font = f_calc; c.border = box; c.alignment = Alignment(horizontal="center")
    c = wsS.cell(r, 3 + len(RELS), f'=COUNTIFS({B_},$A{r},{L_},"OVERDUE")+COUNTIFS({B_},$A{r},{L_},"TODAY")+COUNTIFS({B_},$A{r},{L_},"THIS WEEK")')
    c.font = f_calc; c.border = box; c.alignment = Alignment(horizontal="center")
tr = 5 + len(TYPES)
wsS.cell(tr, 1, "All").font = f_bold
for i in range(2, len(cols) + 1):
    c = wsS.cell(tr, i, f"=SUM({L(i)}5:{L(i)}{tr-1})"); c.font = f_bold; c.border = box; c.alignment = Alignment(horizontal="center")
for i in range(1, len(cols) + 1): wsS.cell(tr, i).fill = fill_goldl
r = tr + 2
wsS.cell(r, 1, "Follow-ups").font = f_sec; r += 1
for lab, crit in (("Overdue", "OVERDUE"), ("Due today", "TODAY"), ("Due in the next 7 days", "THIS WEEK")):
    wsS.cell(r, 1, lab).font = f_lab; c = wsS.cell(r, 2, f'=COUNTIFS({L_},"{crit}")'); c.font = f_calc; c.border = box; r += 1
wsS.cell(r, 1, "Gone quiet").font = f_lab
c = wsS.cell(r, 2, f'=COUNTIFS({G_},"Strong",Contacts!$O$5:$O${MAXR},">90")+COUNTIFS({G_},"Partner",Contacts!$O$5:$O${MAXR},">90")'); c.font = f_calc; c.border = box
wsS.cell(r, 3, "Strong or Partner contacts with no logged contact in 90+ days: worth a quick call.").font = f_note
r += 1
wsS.cell(r, 1, "Contacts logged").font = f_lab; c = wsS.cell(r, 2, '=COUNTA(Log!$B$5:$B$2000)'); c.font = f_calc; c.border = box

for w in wb.worksheets:
    w.page_setup.orientation = "landscape"; w.sheet_properties.pageSetUpPr.fitToPage = True; w.page_setup.fitToWidth = 1; w.page_setup.fitToHeight = 0
wb.active = 0
wb.save(OUT)
from collections import Counter
print("saved", OUT, len(rows), "contacts", dict(Counter(r[1] for r in rows)), "rels", dict(Counter(r[6] for r in rows)))
